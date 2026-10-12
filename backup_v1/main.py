from __future__ import annotations

import json
import os
import queue
import re
import sys
import tempfile
import threading
import time
import traceback
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter

APP_NAME = "WPS 表格固定流程处理器"
CONFIG_PATH = Path.home() / ".wps_sheet_automation.json"
DEFAULT_SETTINGS: dict[str, Any] = {
    "template_path": "",
    "sources": ["", "", ""],
    "output_dir": "",
    "output_filename": "处理结果.xlsx",
    "source_start_row": 2,
    "copy_start_col": "A",
    "copy_end_col": "",
    "jobs": [
        {"template_sheet": "1当天派单", "source_sheet": "", "header_rows": 2, "start_row": 3, "formula_range": "EB:EN", "result_col": "EJ", "target_col": "DG"},
        {"template_sheet": "2当天归档", "source_sheet": "", "header_rows": 2, "start_row": 3, "formula_range": "EB:EU", "result_col": "EQ", "target_col": "DG"},
        {"template_sheet": "3装移全业务在途", "source_sheet": "", "header_rows": 3, "start_row": 4, "formula_range": "EB:FV", "result_col": "FQ", "target_col": "DG"},
    ],
}


class ProcessingError(RuntimeError):
    """An actionable error safe to show to the user."""


def column_number(value: str) -> int:
    value = str(value).strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", value):
        raise ProcessingError(f"列号“{value or '(空)'}”无效；请使用 A、DG、XFD 这样的 Excel 列名。")
    try:
        number = column_index_from_string(value)
    except ValueError as exc:
        raise ProcessingError(f"列号“{value}”超出 Excel 列范围。") from exc
    if number > 16384:
        raise ProcessingError(f"列号“{value}”超出 Excel 最大列 XFD。")
    return number


def parse_column_range(value: str) -> tuple[int, int]:
    parts = str(value).strip().upper().split(":")
    if len(parts) != 2:
        raise ProcessingError(f"列范围“{value}”格式错误；请按 A:C 或 EB:EN 填写。")
    start, end = (column_number(part) for part in parts)
    if start > end:
        raise ProcessingError(f"列范围“{value}”的起始列不能在末列之后。")
    return start, end


def ranges_overlap(first: tuple[int, int], second: tuple[int, int]) -> bool:
    return first[0] <= second[1] and second[0] <= first[1]


def validate_static_settings(settings: dict[str, Any]) -> None:
    jobs = settings.get("jobs")
    if not isinstance(jobs, list) or len(jobs) != 3:
        raise ProcessingError("必须配置恰好 3 个目标工作表。")
    for key in ("source_start_row",):
        try:
            val = int(settings[key])
        except (KeyError, TypeError, ValueError) as exc:
            raise ProcessingError(f"“{key}”必须是正整数。") from exc
        if val < 1 or val > 1_048_576:
            raise ProcessingError(f"“{key}”必须介于 1 和 1048576。")
    copy_start = column_number(settings.get("copy_start_col", "A"))
    copy_end_value = str(settings.get("copy_end_col", "")).strip()
    copy_range = None
    if copy_end_value:
        copy_end = column_number(copy_end_value)
        if copy_start > copy_end:
            raise ProcessingError("数据复制列范围的起始列不能在末列之后。")
        copy_range = (copy_start, copy_end)
    for idx, job in enumerate(jobs, start=1):
        for key in ("template_sheet", "source_sheet", "header_rows", "start_row", "formula_range", "result_col", "target_col"):
            if key not in job or not isinstance(job[key], str):
                if key in ("header_rows", "start_row") and isinstance(job.get(key), int):
                    continue
                raise ProcessingError(f"第 {idx} 个工作表配置缺少有效的“{key}”。")
        if not job["template_sheet"].strip():
            raise ProcessingError(f"第 {idx} 个目标工作表名称不能为空。")
        try:
            header_rows = int(job["header_rows"])
            start_row = int(job["start_row"])
        except (TypeError, ValueError) as exc:
            raise ProcessingError(f"第 {idx} 项“保留标题行数”和“目标/公式起始行”必须是正整数。") from exc
        if header_rows < 0 or header_rows >= 1_048_576:
            raise ProcessingError(f"第 {idx} 项保留标题行数必须介于 0 和 1048575。")
        if start_row <= header_rows or start_row > 1_048_576:
            raise ProcessingError(f"第 {idx} 项目标/公式起始行必须大于保留标题行数，且不超过 1048576。")
        formula_range = parse_column_range(job["formula_range"])
        result_col = column_number(job["result_col"])
        target_col = column_number(job["target_col"])
        if not (formula_range[0] <= result_col <= formula_range[1]):
            raise ProcessingError(f"第 {idx} 项结果列 {job['result_col']} 不在公式范围 {job['formula_range']} 内。")
        if formula_range[0] <= target_col <= formula_range[1]:
            raise ProcessingError(
                f"第 {idx} 项结果粘贴列 {job['target_col']} 与公式区域 {job['formula_range']} 重叠；"
                "为避免覆盖公式，请把结果粘贴到公式区域之外。"
            )
        if copy_range is not None:
            if ranges_overlap(copy_range, formula_range):
                raise ProcessingError(
                    f"第 {idx} 项数据复制范围 {get_column_letter(copy_range[0])}:{get_column_letter(copy_range[1])} "
                    f"与公式区域 {job['formula_range']} 重叠；为避免覆盖公式，请缩小复制范围。"
                )
            if copy_range[0] <= target_col <= copy_range[1]:
                raise ProcessingError(
                    f"第 {idx} 项“结果粘贴列”{job['target_col']} 落在数据复制范围内，最终结果会覆盖复制数据；请调整配置。"
                )
    names = [job["template_sheet"].strip() for job in jobs]
    if len(set(names)) != 3:
        raise ProcessingError("三个目标工作表名称必须互不相同。")


def _resource_path(relative: str) -> Path:
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return root / relative


def load_settings() -> dict[str, Any]:
    settings = deepcopy(DEFAULT_SETTINGS)
    if CONFIG_PATH.exists():
        try:
            saved = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            if not isinstance(saved, dict):
                raise ValueError("配置根节点不是对象")
            settings.update({k: v for k, v in saved.items() if k != "jobs"})
            if isinstance(saved.get("jobs"), list) and len(saved["jobs"]) == 3:
                settings["jobs"] = [dict(DEFAULT_SETTINGS["jobs"][i], **saved["jobs"][i]) for i in range(3)]
        except Exception as exc:
            raise ProcessingError(f"读取用户配置失败：{CONFIG_PATH}\n{exc}") from exc
    return settings


def save_settings(settings: dict[str, Any]) -> None:
    validate_static_settings(settings)
    CONFIG_PATH.write_text(json.dumps(settings, ensure_ascii=False, indent=2), encoding="utf-8")


def _check_xlsx_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise ProcessingError(f"{label}不存在或不是文件：\n{path}")
    if path.suffix.lower() != ".xlsx":
        raise ProcessingError(f"{label}必须是 .xlsx 文件（不接受含宏工作簿或其他格式）：\n{path}")
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = archive.namelist()
            if any(n.startswith("xl/externalLinks/") for n in names):
                raise ProcessingError(f"{label}包含外部工作簿链接；为避免访问外部内容，已拒绝处理：\n{path}")
            if "xl/vbaProject.bin" in names:
                raise ProcessingError(f"{label}包含 VBA 宏；程序只处理无宏 .xlsx 文件：\n{path}")
    except ProcessingError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise ProcessingError(f"{label}不是可读取的有效 XLSX 工作簿：\n{path}\n{exc}") from exc


def _find_external_formula(formula: str) -> bool:
    if re.search(r"\[[^\]]+\.(?:xlsx?|xlsm|xlsb)\]", formula, flags=re.IGNORECASE):
        return True
    if re.search(r"\b(?:WEBSERVICE|FILTERXML|RTD|CALL|REGISTER\.ID|EXEC)\s*\(", formula, flags=re.IGNORECASE):
        return True
    return False


def _normalize_path(value: str, label: str) -> Path:
    raw = str(value or "").strip()
    if not raw:
        raise ProcessingError(f"请先选择{label}。")
    return Path(raw).expanduser().resolve()


def _get_formula_templates(workbook, job: dict[str, str], index: int, formula_row: int) -> tuple[int, int, dict[int, str]]:
    sheet_name = job["template_sheet"].strip()
    if sheet_name not in workbook.sheetnames:
        raise ProcessingError(f"模板中找不到第 {index} 个目标工作表“{sheet_name}”。")
    if workbook.sheetnames[index - 1] != sheet_name:
        raise ProcessingError(
            f"模板第 {index} 个子表应为“{sheet_name}”，实际为“{workbook.sheetnames[index - 1]}”；"
            "请检查可配置的工作表名称和模板前三个子表顺序。"
        )
    ws = workbook[sheet_name]
    row = int(formula_row)
    col_start, col_end = parse_column_range(job["formula_range"])
    formulas: dict[int, str] = {}
    for col in range(col_start, col_end + 1):
        value = ws.cell(row=row, column=col).value
        if not (isinstance(value, str) and value.startswith("=")):
            letter = get_column_letter(col)
            raise ProcessingError(
                f"模板“{sheet_name}”的公式模板行 {row}、单元格 {letter}{row} 不是公式。"
                "请修正模板公式行，或在配置中填写正确的“公式模板行”。"
            )
        if _find_external_formula(value):
            raise ProcessingError(f"模板公式 {sheet_name}!{get_column_letter(col)}{row} 包含外部引用或外部访问函数，已拒绝重算。")
        formulas[col] = value
    return col_start, col_end, formulas


def _read_source_records(path: Path, sheet_name: str, start_row: int, col_start: int, col_end: int,
                         log: Callable[[str], None]) -> list[list[Any]]:
    _check_xlsx_file(path, f"源工作簿“{path.name}”")
    values_book = formula_book = None
    try:
        values_book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formula_book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        if sheet_name not in values_book.sheetnames or sheet_name not in formula_book.sheetnames:
            raise ProcessingError(f"源工作簿“{path.name}”中找不到配置的工作表“{sheet_name}”；不会猜测工作表名称。")
        val_ws = values_book[sheet_name]
        raw_ws = formula_book[sheet_name]
        max_row = max(val_ws.max_row or 0, raw_ws.max_row or 0)
        if max_row < start_row:
            raise ProcessingError(f"源工作簿“{path.name}”工作表“{sheet_name}”在第 {start_row} 行及之后没有可复制数据。")
        records: list[list[Any]] = []
        blank_formula_cache_count = 0
        raw_rows = raw_ws.iter_rows(min_row=start_row, max_row=max_row, min_col=col_start, max_col=col_end)
        value_rows = val_ws.iter_rows(min_row=start_row, max_row=max_row, min_col=col_start, max_col=col_end)
        for raw_cells, value_cells in zip(raw_rows, value_rows):
            raw = [cell.value for cell in raw_cells]
            vals = [cell.value for cell in value_cells]
            if not any(v is not None and (not isinstance(v, str) or v.strip() != "") for v in raw):
                continue
            for source_value, cached_value in zip(raw, vals):
                if isinstance(source_value, str) and source_value.startswith("=") and cached_value is None:
                    blank_formula_cache_count += 1
            records.append(vals)
        if not records:
            raise ProcessingError(f"源工作簿“{path.name}”工作表“{sheet_name}”从第 {start_row} 行起在复制范围内没有数据。")
        if blank_formula_cache_count:
            log(f"提示：{path.name} 有 {blank_formula_cache_count} 个源公式单元格没有已保存的计算缓存；这些单元格将按空值复制，不会执行源公式。")
        return records
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"读取源工作簿“{path.name}”失败：{exc}") from exc
    finally:
        for wb in (values_book, formula_book):
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass


def _validate_inputs(settings: dict[str, Any]) -> tuple[Path, list[Path], Path]:
    validate_static_settings(settings)
    template_path = _normalize_path(settings.get("template_path", ""), "模板工作簿")
    sources_raw = settings.get("sources", [])
    if not isinstance(sources_raw, list) or len(sources_raw) != 3:
        raise ProcessingError("必须分别选择 1、2、3 三个源工作簿。")
    source_paths = [_normalize_path(p, f"第 {i} 个源工作簿") for i, p in enumerate(sources_raw, 1)]
    for i, job in enumerate(settings["jobs"], 1):
        if not job["source_sheet"].strip():
            raise ProcessingError(f"第 {i} 个源工作表名称尚未配置；程序不会猜测源工作表。")
    for label, path in [("模板工作簿", template_path)] + [(f"第 {i} 个源工作簿", p) for i, p in enumerate(source_paths, 1)]:
        _check_xlsx_file(path, label)
    all_inputs = [template_path] + source_paths
    if len(set(all_inputs)) != len(all_inputs):
        raise ProcessingError("模板和三个源工作簿必须是四个不同文件。")
    output_dir = _normalize_path(settings.get("output_dir", ""), "输出目录")
    if not output_dir.is_dir():
        raise ProcessingError(f"输出目录不存在：\n{output_dir}")
    filename = str(settings.get("output_filename", "")).strip()
    if not filename:
        raise ProcessingError("输出文件名不能为空。")
    if Path(filename).name != filename or "/" in filename or "\\" in filename:
        raise ProcessingError("输出文件名只能填写文件名，不能包含目录路径。")
    if Path(filename).suffix and Path(filename).suffix.lower() != ".xlsx":
        raise ProcessingError("输出文件只支持 .xlsx 扩展名；请勿使用 .xlsm 或其他格式。")
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    final_path = (output_dir / filename).resolve()
    if final_path.exists():
        raise ProcessingError(f"输出文件已存在，为防止覆盖任何文件已停止：\n{final_path}\n请更改输出文件名或选择其他输出目录。")
    if final_path in all_inputs:
        raise ProcessingError("输出路径不能与模板或任何源工作簿相同。")
    column_number(settings["copy_start_col"])
    column_number(settings["copy_end_col"])
    return template_path, source_paths, final_path


def _create_staging_workbook(template_path: Path, settings: dict[str, Any], source_records: list[list[list[Any]]],
                             staging_path: Path, log: Callable[[str], None]) -> list[tuple[str, str, str, int, int]]:
    wb = None
    try:
        wb = load_workbook(template_path, keep_links=False, data_only=False)
        if len(wb.sheetnames) < 3:
            raise ProcessingError("模板工作簿少于三个子表，无法按流程处理。")
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith("=") and _find_external_formula(cell.value):
                    raise ProcessingError(f"模板“{ws.title}”单元格 {cell.coordinate} 含外部引用或外部访问函数；为避免外部访问已停止。")
        copy_start = column_number(settings["copy_start_col"])
        formula_jobs: list[tuple[str, str, str, int, int]] = []
        prepared = []
        # Validate every target worksheet and seed formula before changing any worksheet in the staging copy.
        for i, job in enumerate(settings["jobs"], 1):
            start_row = int(job["start_row"])
            c1, c2, formulas = _get_formula_templates(wb, job, i, start_row)
            ws = wb[job["template_sheet"].strip()]
            prepared.append((i, job, ws, start_row, int(job["header_rows"]), c1, c2, formulas))

        for i, (job, records) in enumerate(zip(settings["jobs"], source_records), 1):
            _, _, ws, start_row, header_rows, c1, c2, formulas = prepared[i - 1]
            end_row = start_row + len(records) - 1
            if end_row > 1_048_576:
                raise ProcessingError(f"第 {i} 个工作表目标末行超出 Excel 最大行数。")

            # Clear old cell contents below retained headers, but keep the formula seed cells at start_row.
            formula_columns = set(formulas)
            cleared = 0
            for cell in list(ws._cells.values()):
                if isinstance(cell, MergedCell) or cell.row <= header_rows:
                    continue
                if cell.row == start_row and cell.column in formula_columns:
                    continue
                had_content = cell.value is not None or cell.hyperlink is not None or cell.comment is not None
                if had_content:
                    cell.value = None
                    cell.hyperlink = None
                    cell.comment = None
                    cleared += 1
            log(f"已清除“{job['template_sheet']}”第 {header_rows + 1} 行起的旧内容（清除 {cleared} 个单元格），保留格式和第 {start_row} 行公式种子。")

            # Copy only stored source values into cells; template styles remain untouched.
            for offset, values in enumerate(records):
                target_row = start_row + offset
                for j, value in enumerate(values):
                    ws.cell(row=target_row, column=copy_start + j).value = value
            for col, formula in formulas.items():
                origin = f"{get_column_letter(col)}{start_row}"
                for target_row in range(start_row, end_row + 1):
                    target = f"{get_column_letter(col)}{target_row}"
                    try:
                        ws[target] = Translator(formula, origin=origin).translate_formula(target)
                    except Exception as exc:
                        raise ProcessingError(f"无法将模板公式 {job['template_sheet']}!{origin} 填充到 {target}：{exc}") from exc
            formula_jobs.append((job["template_sheet"].strip(), job["result_col"].strip().upper(),
                                 job["target_col"].strip().upper(), start_row, end_row))
            log(f"已准备工作表“{job['template_sheet']}”：复制 {len(records)} 行，填充第 {start_row} 至 {end_row} 行公式。")
        wb.save(staging_path)
        return formula_jobs
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"创建中间工作簿失败：{exc}") from exc
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def _connect_wps(log: Callable[[str], None]):
    if sys.platform != "win32":
        raise ProcessingError("WPS 自动重算只能在安装了 Windows 版 WPS 表格的 Windows 环境运行；当前环境不是 Windows，未生成输出文件。")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ProcessingError("缺少 Windows COM 依赖 pywin32。请在 Windows 运行 install_and_build.bat 安装依赖后重试。") from exc
    try:
        pythoncom.CoInitialize()
    except Exception as exc:
        raise ProcessingError(f"无法初始化 Windows COM，因而不能连接 WPS 表格：{type(exc).__name__}: {exc}") from exc
    progids = ("ket.Application", "ET.Application", "et.Application")
    failures = []
    for progid in progids:
        try:
            log(f"正在连接 WPS 表格 COM：{progid}")
            app = win32com.client.DispatchEx(progid)
            return app, pythoncom, progid
        except Exception as exc:
            failures.append(f"{progid}: {exc}")
    pythoncom.CoUninitialize()
    detail = "\n".join(failures)
    raise ProcessingError(
        "无法连接 Windows 版 WPS 表格 COM/ProgID。已尝试 ket.Application、ET.Application、et.Application。\n"
        "请确认 Windows 11 已安装并能正常启动 WPS 表格；若已安装，请修复/重新注册 WPS 表格自动化组件后重试。"
        "本程序不会改用 Microsoft Excel，也不会在连接失败后留下最终输出文件。\n\n连接错误详情：\n" + detail
    )


def _calculate_in_wps(staging_path: Path, formula_jobs: list[tuple[str, str, str, int, int]],
                      log: Callable[[str], None]) -> None:
    app = workbook = pythoncom = None
    com_initialized = False
    try:
        app, pythoncom, progid = _connect_wps(log)
        com_initialized = True
        for prop, value in (("Visible", False), ("DisplayAlerts", False), ("AskToUpdateLinks", False), ("EnableEvents", False), ("AutomationSecurity", 3)):
            try:
                setattr(app, prop, value)
            except Exception:
                pass
        try:
            app.Calculation = -4105  # xlCalculationAutomatic
        except Exception:
            pass
        try:
            workbook = app.Workbooks.Open(str(staging_path), UpdateLinks=0, ReadOnly=False, AddToMru=False)
        except Exception as exc:
            raise ProcessingError(
                f"已连接 WPS（{progid}），但无法在禁用更新外部链接的模式下打开中间工作簿：{exc}\n"
                "请检查 WPS 版本、文件权限及工作簿兼容性。"
            ) from exc
        log("已在 WPS 中打开中间工作簿，开始完整重算公式。")
        calculate_errors = []
        for method_name in ("CalculateFullRebuild", "CalculateFull", "Calculate"):
            try:
                getattr(app, method_name)()
                calculate_errors = []
                log(f"WPS 已执行 {method_name}。")
                break
            except Exception as exc:
                calculate_errors.append(f"{method_name}: {exc}")
        if calculate_errors:
            raise ProcessingError("WPS 无法触发公式重算：\n" + "\n".join(calculate_errors))
        deadline = time.time() + 180
        while time.time() < deadline:
            try:
                state = int(app.CalculationState)
                if state == 0:
                    break
            except Exception:
                break
            time.sleep(0.25)
        else:
            raise ProcessingError("等待 WPS 完成公式重算超过 180 秒，已中止以避免输出不完整文件。")
        for sheet_name, result_col, target_col, start_row, end_row in formula_jobs:
            ws = workbook.Worksheets(sheet_name)
            src = ws.Range(f"{result_col}{start_row}:{result_col}{end_row}").Value2
            count = end_row - start_row + 1
            if count == 1:
                values = [src[0][0] if isinstance(src, (tuple, list)) else src]
            else:
                values = [row[0] for row in src]
            destination = ws.Range(f"{target_col}{start_row}:{target_col}{end_row}")
            destination.Value2 = tuple((value,) for value in values)
            log(f"已将 {sheet_name}!{result_col}{start_row}:{result_col}{end_row} 的当前计算值写入 {target_col} 列。")
        workbook.Save()
        log("WPS 已保存重算后的工作簿。")
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"WPS 自动化处理中失败：{type(exc).__name__}: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        if com_initialized and pythoncom is not None:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass


def process_workbooks(settings: dict[str, Any], log: Callable[[str], None] = print) -> Path:
    """Validate, stage, recalculate in WPS, and atomically publish a new XLSX."""
    template_path, source_paths, final_path = _validate_inputs(settings)
    source_start = int(settings["source_start_row"])
    copy_start = column_number(settings["copy_start_col"])
    copy_end = column_number(settings["copy_end_col"])
    all_records = []
    for i, (path, job) in enumerate(zip(source_paths, settings["jobs"]), 1):
        log(f"读取源工作簿 {i}/3：{path.name}（工作表“{job['source_sheet']}”）")
        all_records.append(_read_source_records(path, job["source_sheet"].strip(), source_start, copy_start, copy_end, log))
    fd, temp_name = tempfile.mkstemp(prefix=f".{final_path.stem}_", suffix=".staging.xlsx", dir=str(final_path.parent))
    os.close(fd)
    staging_path = Path(temp_name)
    try:
        log("校验模板公式及工作表后，创建同目录临时工作簿。")
        formula_jobs = _create_staging_workbook(template_path, settings, all_records, staging_path, log)
        _calculate_in_wps(staging_path, formula_jobs, log)
        if not staging_path.is_file() or staging_path.stat().st_size == 0:
            raise ProcessingError("WPS 处理结束但临时工作簿不存在或为空；没有发布输出文件。")
        # os.rename does not replace an existing destination on Windows; keep the no-overwrite guarantee.
        if final_path.exists():
            raise ProcessingError(f"输出文件在处理过程中被创建；为避免覆盖已停止：\n{final_path}")
        os.rename(staging_path, final_path)
        log(f"处理成功，已另存为：{final_path}")
        return final_path
    except Exception as exc:
        try:
            staging_path.unlink(missing_ok=True)
        except Exception as cleanup_exc:
            raise ProcessingError(
                f"处理失败：{exc}\n同时无法清理临时工作簿：{staging_path}\n清理错误：{cleanup_exc}"
            ) from exc
        raise


class App:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
        self.tk = tk
        self.filedialog = filedialog
        self.messagebox = messagebox
        self.ttk = ttk
        self.root = root
        self.root.title(APP_NAME)
        self.root.geometry("1240x880")
        self.root.minsize(1000, 760)
        self.events: queue.Queue[tuple[str, Any]] = queue.Queue()
        self.settings = load_settings()
        self.vars: dict[str, Any] = {}
        self.job_vars: list[dict[str, Any]] = []
        self._build()
        self.root.after(150, self._drain_events)

    def _label_entry(self, parent, row, label, key, width=62, browse=None):
        self.ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(8, 8), pady=4)
        var = self.tk.StringVar(value=str(self.settings.get(key, "")))
        self.vars[key] = var
        self.ttk.Entry(parent, textvariable=var, width=width).grid(row=row, column=1, sticky="ew", padx=4, pady=4)
        if browse:
            self.ttk.Button(parent, text="浏览…", command=browse).grid(row=row, column=2, padx=5, pady=4)
        return var

    def _build(self):
        root, ttk = self.root, self.ttk
        outer = ttk.Frame(root, padding=10)
        outer.pack(fill="both", expand=True)
        outer.columnconfigure(0, weight=1)
        files = ttk.LabelFrame(outer, text="文件", padding=8)
        files.grid(row=0, column=0, sticky="ew")
        files.columnconfigure(1, weight=1)
        self._label_entry(files, 0, "模板工作簿（.xlsx）", "template_path", browse=self._browse_template)
        self.source_vars = []
        for i in range(3):
            row = i + 1
            ttk.Label(files, text=f"源工作簿 {i+1}（.xlsx）").grid(row=row, column=0, sticky="w", padx=(8, 8), pady=4)
            var = self.tk.StringVar(value=self.settings.get("sources", ["", "", ""])[i])
            self.source_vars.append(var)
            ttk.Entry(files, textvariable=var).grid(row=row, column=1, sticky="ew", padx=4, pady=4)
            ttk.Button(files, text="选择文件…", command=lambda n=i: self._browse_source(n)).grid(row=row, column=2, padx=5, pady=4)
        ttk.Button(files, text="从文件夹自动匹配 1.xlsx、2.xlsx、3.xlsx…", command=self._browse_source_folder).grid(row=4, column=1, sticky="w", padx=4, pady=4)
        self._label_entry(files, 5, "输出目录", "output_dir", browse=self._browse_output_dir)
        self._label_entry(files, 6, "输出文件名", "output_filename")

        shared = ttk.LabelFrame(outer, text="通用参数", padding=8)
        shared.grid(row=1, column=0, sticky="ew", pady=(10, 0))
        shared.columnconfigure(1, weight=1)
        fields = [
            ("source_start_row", "源数据起始行（默认排除第 1 行）"),
            ("copy_start_col", "数据复制起始列"),
            ("copy_end_col", "数据复制末列（必填；不确定时不要猜）"),
        ]
        for row, (key, label) in enumerate(fields):
            ttk.Label(shared, text=label).grid(row=row, column=0, sticky="w", padx=(8, 8), pady=4)
            var = self.tk.StringVar(value=str(self.settings.get(key, "")))
            self.vars[key] = var
            ttk.Entry(shared, textvariable=var, width=18).grid(row=row, column=1, sticky="w", padx=4, pady=4)
        note = "逐表起始行同时作为粘贴行和公式种子行。副本中清除标题之后旧内容，保留种子公式/格式；仅处理无宏 XLSX，失败不发布输出且不覆盖模板或已有文件。"
        ttk.Label(shared, text=note, wraplength=1050, foreground="#4b5563").grid(row=3, column=0, columnspan=2, sticky="w", padx=8, pady=(4, 2))

        table = ttk.LabelFrame(outer, text="三个目标子表 / 公式结果映射", padding=8)
        table.grid(row=2, column=0, sticky="ew", pady=(10, 0))
        headers = ["目标工作表（模板前 3 个）", "源工作表（必填）", "保留标题行数", "数据/公式起始行", "公式列范围", "结果列", "写入列"]
        for col, title in enumerate(headers):
            ttk.Label(table, text=title).grid(row=0, column=col, sticky="w", padx=5, pady=(0, 5))
        widths = [25, 18, 10, 12, 13, 9, 9]
        for i, default_job in enumerate(DEFAULT_SETTINGS["jobs"]):
            configured = self.settings["jobs"][i]
            row_vars = {}
            for col, key in enumerate(("template_sheet", "source_sheet", "header_rows", "start_row", "formula_range", "result_col", "target_col")):
                var = self.tk.StringVar(value=str(configured.get(key, default_job[key])))
                row_vars[key] = var
                ttk.Entry(table, textvariable=var, width=widths[col]).grid(row=i+1, column=col, sticky="ew", padx=4, pady=4)
            self.job_vars.append(row_vars)
        for col in range(7):
            table.columnconfigure(col, weight=1 if col < 2 else 0)

        controls = ttk.Frame(outer)
        controls.grid(row=3, column=0, sticky="ew", pady=10)
        self.run_button = ttk.Button(controls, text="开始处理", command=self._run)
        self.run_button.pack(side="left")
        ttk.Button(controls, text="保存当前配置", command=self._save_config).pack(side="left", padx=8)
        ttk.Button(controls, text="恢复默认配置", command=self._reset_config).pack(side="left")
        self.progress = ttk.Progressbar(controls, mode="indeterminate", length=180)
        self.progress.pack(side="right", padx=5)

        log_frame = ttk.LabelFrame(outer, text="处理日志", padding=6)
        log_frame.grid(row=4, column=0, sticky="nsew")
        outer.rowconfigure(4, weight=1)
        self.log_text = self.tk.Text(log_frame, height=12, wrap="word", state="disabled")
        scrollbar = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._log("请填写源工作表名称和数据复制末列。程序无法安全判断这两项时会停止，不会猜测或覆盖公式区域。")

    def _browse_template(self):
        path = self.filedialog.askopenfilename(filetypes=[("Excel 工作簿（无宏）", "*.xlsx")])
        if path:
            self.vars["template_path"].set(path)

    def _browse_source(self, index):
        path = self.filedialog.askopenfilename(filetypes=[("Excel 工作簿（无宏）", "*.xlsx")])
        if path:
            self.source_vars[index].set(path)

    def _browse_source_folder(self):
        folder = self.filedialog.askdirectory()
        if folder:
            for i in range(3):
                self.source_vars[i].set(str(Path(folder) / f"{i+1}.xlsx"))

    def _browse_output_dir(self):
        path = self.filedialog.askdirectory()
        if path:
            self.vars["output_dir"].set(path)

    def _collect(self):
        settings = {key: var.get().strip() for key, var in self.vars.items()}
        for key in ("source_start_row",):
            try:
                settings[key] = int(settings[key])
            except ValueError as exc:
                raise ProcessingError(f"“{key}”必须填写正整数。") from exc
        settings["sources"] = [v.get().strip() for v in self.source_vars]
        settings["jobs"] = []
        for row_vars in self.job_vars:
            job = {key: var.get().strip() for key, var in row_vars.items()}
            for key in ("header_rows", "start_row"):
                try:
                    job[key] = int(job[key])
                except ValueError as exc:
                    raise ProcessingError(f"“{key}”必须填写正整数。") from exc
            settings["jobs"].append(job)
        return settings

    def _save_config(self):
        try:
            settings = self._collect()
            validate_static_settings(settings)
            save_settings(settings)
            self.messagebox.showinfo(APP_NAME, f"配置已保存到：\n{CONFIG_PATH}")
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))

    def _reset_config(self):
        if not self.messagebox.askyesno(APP_NAME, "恢复程序默认参数？这不会删除已保存的配置文件。"):
            return
        defaults = deepcopy(DEFAULT_SETTINGS)
        for key in ("template_path", "output_dir", "output_filename", "source_start_row", "copy_start_col", "copy_end_col"):
            self.vars[key].set(str(defaults[key]))
        for i, var in enumerate(self.source_vars):
            var.set(defaults["sources"][i])
        for i, row in enumerate(self.job_vars):
            for key, var in row.items():
                var.set(str(defaults["jobs"][i][key]))

    def _log(self, message: str):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", message + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _run(self):
        try:
            settings = self._collect()
            validate_static_settings(settings)
            save_settings(settings)
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))
            return
        self.run_button.configure(state="disabled")
        self.progress.start(10)
        self._log("开始处理……")

        def work():
            try:
                path = process_workbooks(settings, lambda text: self.events.put(("log", text)))
                self.events.put(("success", str(path)))
            except Exception as exc:
                self.events.put(("error", f"{exc}\n\n详细诊断：\n{traceback.format_exc()}"))
        threading.Thread(target=work, daemon=True).start()

    def _drain_events(self):
        try:
            while True:
                event, payload = self.events.get_nowait()
                if event == "log":
                    self._log(str(payload))
                elif event in ("success", "error"):
                    self.progress.stop()
                    self.run_button.configure(state="normal")
                    if event == "success":
                        self._log(f"完成：{payload}")
                        self.messagebox.showinfo(APP_NAME, f"处理成功，已保存新文件：\n{payload}")
                    else:
                        self._log("处理失败；临时工作簿已清理，未发布半成品输出。")
                        self.messagebox.showerror(APP_NAME, str(payload))
        except queue.Empty:
            pass
        self.root.after(150, self._drain_events)


def self_check() -> int:
    """Linux-safe static/config tests; does not claim WPS or Windows validation."""
    try:
        validate_static_settings(DEFAULT_SETTINGS)
        expected = {"EB:EN": (132, 144), "EB:EU": (132, 151), "EB:FV": (132, 178)}
        for text, answer in expected.items():
            if parse_column_range(text) != answer:
                raise AssertionError(f"列范围换算错误：{text}")
        for job in DEFAULT_SETTINGS["jobs"]:
            lo, hi = parse_column_range(job["formula_range"])
            result = column_number(job["result_col"])
            if not lo <= result <= hi:
                raise AssertionError(f"结果列不在公式范围：{job}")
            if job["target_col"] != "DG":
                raise AssertionError("默认结果目标列应为 DG")
        if not ranges_overlap((132, 144), (140, 150)) or ranges_overlap((1, 10), (11, 20)):
            raise AssertionError("列范围重叠检测错误")
        print("静态检查通过：默认配置结构有效；EB:EN、EB:EU、EB:FV 及 EJ/EQ/FQ→DG 映射正确。")
        print("未连接 WPS COM；未读取真实业务文件；未验证 Windows GUI 或 EXE。")
        return 0
    except Exception as exc:
        print(f"静态检查失败：{exc}", file=sys.stderr)
        return 1


def main() -> int:
    if "--self-check" in sys.argv:
        return self_check()
    if sys.platform != "win32":
        print("此程序的公式重算与最终保存要求 Windows 版 WPS 表格 COM。请在 Windows 11 + WPS 表格环境运行。")
        print("当前执行环境不是 Windows；没有生成或验证 EXE。可在当前环境运行 python main.py --self-check 做静态检查。")
        return 2
    try:
        import tkinter as tk
        root = tk.Tk()
        App(root)
        root.mainloop()
        return 0
    except Exception as exc:
        print(f"启动失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
