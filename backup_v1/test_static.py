from __future__ import annotations

import tempfile
import unittest
import json
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

import main
from make_fake_data import create_fake_data


class StaticConfigurationTests(unittest.TestCase):
    def test_defaults_and_formula_mappings(self):
        main.validate_static_settings(main.DEFAULT_SETTINGS)
        expected = [("EJ", "DG", 2, 3), ("EQ", "DG", 2, 3), ("FQ", "DG", 3, 4)]
        actual = [(j["result_col"], j["target_col"], j["header_rows"], j["start_row"]) for j in main.DEFAULT_SETTINGS["jobs"]]
        self.assertEqual(expected, actual)
        self.assertEqual([(132, 144), (132, 151), (132, 178)], [main.parse_column_range(j["formula_range"]) for j in main.DEFAULT_SETTINGS["jobs"]])

    def test_excel_column_range_and_overlap_checks(self):
        self.assertEqual(111, main.column_number("DG"))
        self.assertTrue(main.ranges_overlap((132, 144), (140, 150)))
        self.assertFalse(main.ranges_overlap((1, 10), (11, 20)))
        with self.assertRaises(main.ProcessingError):
            main.parse_column_range("EN:EB")
        with self.assertRaises(main.ProcessingError):
            main.column_number("XFE")

    def test_sample_json_matches_defaults_and_unsafe_overlaps_are_rejected(self):
        sample = json.loads((Path(__file__).parent / "default_config.json").read_text(encoding="utf-8"))
        main.validate_static_settings(sample)
        self.assertEqual([j["start_row"] for j in main.DEFAULT_SETTINGS["jobs"]], [j["start_row"] for j in sample["jobs"]])
        overlapping_copy = deepcopy(sample)
        overlapping_copy["copy_end_col"] = "EN"
        with self.assertRaises(main.ProcessingError):
            main.validate_static_settings(overlapping_copy)
        overlapping_result = deepcopy(sample)
        overlapping_result["jobs"][0]["target_col"] = "EB"
        with self.assertRaises(main.ProcessingError):
            main.validate_static_settings(overlapping_result)

    def test_fake_staging_clears_contents_but_keeps_headers_styles_and_seed_formulas(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fake = create_fake_data(root / "fake_data")
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["copy_end_col"] = "C"
            settings["template_path"] = str(fake / "template.xlsx")
            settings["sources"] = [str(fake / f"{i}.xlsx") for i in (1, 2, 3)]
            settings["jobs"] = [dict(j, source_sheet="示例源表") for j in settings["jobs"]]
            records = []
            for source_path, job in zip(settings["sources"], settings["jobs"]):
                records.append(main._read_source_records(Path(source_path), job["source_sheet"], 2, 1, 3, lambda _msg: None))
            staging = root / "prepared.xlsx"
            main._create_staging_workbook(Path(settings["template_path"]), settings, records, staging, lambda _msg: None)
            wb = load_workbook(staging, data_only=False)
            try:
                for index, job in enumerate(settings["jobs"], start=1):
                    ws = wb[job["template_sheet"]]
                    start = job["start_row"]
                    self.assertIn("演示标题", ws["A1"].value)
                    self.assertIsNone(ws["A20"].value)
                    self.assertIsNone(ws["DG25"].value)
                    self.assertTrue(ws[f"EB{start}"].value.startswith("=IF("))
                    self.assertEqual(f"FAKE-{index}-01", ws[f"A{start}"].value)
                    self.assertTrue(ws[f"EB{start+1}"].value.startswith("=IF("))
                    self.assertIsNotNone(ws["A20"].fill.fgColor.rgb)
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
