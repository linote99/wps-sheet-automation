from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils.cell import column_index_from_string

JOBS = [
    ("1当天派单", 2, 3, "EB:EN", "EJ"),
    ("2当天归档", 2, 3, "EB:EU", "EQ"),
    ("3装移全业务在途", 3, 4, "EB:FV", "FQ"),
]
SOURCE_SHEET = "示例源表"


def col_range(value: str) -> tuple[int, int]:
    start, end = value.split(":")
    return column_index_from_string(start), column_index_from_string(end)


def _create_source(path: Path, prefix: str, category: str) -> None:
    source = Workbook()
    ws = source.active
    ws.title = SOURCE_SHEET
    ws.append(["测试编号", "虚构姓名", "虚构类别"])
    for row in range(1, 5):
        ws.append([f"{prefix}-{row:02d}", f"测试用户{row}", category])
    source.save(path)
    source.close()


def create_fake_data(base_dir: Path) -> Path:
    base_dir.mkdir(parents=True, exist_ok=True)
    (base_dir / "output").mkdir(exist_ok=True)

    # Four purely fabricated rows per workbook. Extra workbooks help exercise ordered appends.
    for idx in range(1, 4):
        _create_source(base_dir / f"{idx}.xlsx", f"FAKE-{idx}", f"分类{idx}")
    _create_source(base_dir / "1_extra.xlsx", "FAKE-1X", "分类1-补充")
    _create_source(base_dir / "2_extra.xlsx", "FAKE-2X", "分类2-补充")

    template = Workbook()
    template.remove(template.active)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    old_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="9AA6B2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for sheet_name, header_rows, start_row, formula_range, _result_col in JOBS:
        ws = template.create_sheet(sheet_name)
        for row in range(1, header_rows + 1):
            ws.cell(row=row, column=1, value=f"演示标题 {row}（{sheet_name}）")
            ws.cell(row=row, column=1).font = Font(bold=True)
            ws.cell(row=row, column=1).fill = header_fill
        first, last = col_range(formula_range)
        for col in range(first, last + 1):
            cell = ws.cell(row=start_row, column=col)
            cell.value = f'=IF($A{start_row}="", "", $A{start_row}&"-演示")'
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
            cell.border = border
        # Old values are deliberately present below headings; processing must clear contents, not styles.
        ws["A20"] = "待清除的旧内容"
        ws["A20"].fill = old_fill
        ws["DG25"] = "旧结果"
        ws["DG25"].fill = old_fill
        ws.freeze_panes = f"A{header_rows + 1}"
    template.save(base_dir / "template.xlsx")
    template.close()
    return base_dir


if __name__ == "__main__":
    destination = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else Path(__file__).resolve().parent / "fake_data"
    folder = create_fake_data(destination.resolve())
    print(f"假数据工作簿已创建：{folder}")
    print("源表名：示例源表；复制列建议 A:C；每个源工作簿含 4 条虚构记录，并含补充源文件。")
