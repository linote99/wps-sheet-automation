from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

import main
from test_static import FakeComWorkbook, FakeWpsApplication


OPS = ("CustomFormula", "TextToColumns", "RegionFilter", "Deduplicate", "FormatConvert")


def workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    book = Workbook()
    for index, (name, rows) in enumerate(sheets.items()):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)
    book.save(path)
    book.close()


def source_rule(sheet="Data", area="A1:A999", **extra):
    rule = {"scope": "source", "sheet": sheet, "range": area}
    rule.update(extra)
    return rule


class SelectedSourceCopyTests(unittest.TestCase):
    """These tests intentionally exercise the public source-only contract."""

    @classmethod
    def setUpClass(cls):
        if not hasattr(main, "_prepare_selected_source_copy"):
            raise unittest.SkipTest("source-selection implementation not present yet")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_single_copy_is_full_workbook_and_does_not_change_source(self):
        source = self.root / "input.xlsx"
        workbook(source, {"Data": [["id", "formula"], [1, "=1+1"]], "Notes": [["keep"]]})
        before = source.read_bytes()
        group = {"source_paths": [str(source)], "output_filename": "out.xlsx",
                 "rules": {"CustomFormula": [source_rule()]}}
        copied = Path(main._prepare_selected_source_copy(group, lambda _msg: None))
        self.addCleanup(copied.unlink, missing_ok=True)
        self.assertTrue(copied.is_file())
        self.assertEqual(source.read_bytes(), before)
        result = load_workbook(copied, data_only=False)
        self.addCleanup(result.close)
        self.assertEqual(result.sheetnames, ["Data", "Notes"])
        self.assertEqual(result["Data"]["B2"].value, "=1+1")
        self.assertEqual(result["Notes"]["A1"].value, "keep")

    def test_multi_copy_merges_only_referenced_sheet_and_keeps_every_header(self):
        first = self.root / "first.xlsx"
        second = self.root / "second.xlsx"
        workbook(first, {"Data": [["id", "value"], [1, "a"]], "Unrelated": [["x"]]})
        workbook(second, {"Other": [["z"]], "Data": [["id", "value"], [2, "b"]]})
        group = {"source_paths": [str(first), str(second)], "output_filename": "merged.xlsx",
                 "rules": {"Deduplicate": [source_rule()]}}
        copied = Path(main._prepare_selected_source_copy(group, lambda _msg: None))
        self.addCleanup(copied.unlink, missing_ok=True)
        result = load_workbook(copied, data_only=False)
        self.addCleanup(result.close)
        self.assertEqual(result.sheetnames, ["Data"])
        self.assertEqual(list(result["Data"].values), [("id", "value"), (1, "a"), ("id", "value"), (2, "b")])

    def test_missing_referenced_sheet_leaves_no_temporary_copy(self):
        first = self.root / "first.xlsx"
        second = self.root / "second.xlsx"
        workbook(first, {"Data": [[1]]})
        workbook(second, {"Other": [[2]]})
        group = {"source_paths": [str(first), str(second)], "output_filename": "merged.xlsx",
                 "rules": {"FormatConvert": [source_rule(sheet="Data")]}}
        with self.assertRaises(Exception):
            main._prepare_selected_source_copy(group, lambda _msg: None)
        self.assertEqual(list(self.root.glob("*.xlsx")), [first, second])


class SelectedSourceWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not hasattr(main, "process_selected_source_files"):
            raise unittest.SkipTest("source-only workflow implementation not present yet")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "source.xlsx"
        workbook(self.source, {"Data": [["id"], [1], [1]]})

    def tearDown(self):
        self.tmp.cleanup()

    def _settings(self, output="result.xlsx", paths=None):
        paths = [str(self.source)] if paths is None else paths
        return {"source_processing_only": True, "workflow": list(OPS),
                "CustomFormula": {"enabled": False, "rules": []},
                "TextToColumns": {"enabled": False, "rules": []},
                "RegionFilter": {"enabled": False, "rules": []},
                "Deduplicate": {"enabled": True, "rules": [source_rule(source_paths=paths,
                                                                                source_output_filename=output)]},
                "FormatConvert": {"enabled": False, "rules": []}}

    def test_preflight_rejects_missing_or_non_xlsx_before_connecting_wps(self):
        settings = self._settings(paths=[str(self.source), str(self.root / "missing.xlsx")])
        with mock.patch.object(main, "_connect_wps", side_effect=AssertionError("WPS must not start")):
            with self.assertRaises(Exception):
                main.process_selected_source_files(settings, log=lambda _msg: None)
        bad = self.root / "bad.csv"
        bad.write_text("not xlsx", encoding="utf-8")
        with mock.patch.object(main, "_connect_wps", side_effect=AssertionError("WPS must not start")):
            with self.assertRaises(Exception):
                main.process_selected_source_files(self._settings(paths=[str(bad)]), log=lambda _msg: None)

    def test_callback_receives_each_published_output_and_output_never_overwrites_input(self):
        callback_paths = []
        # The COM layer is replaced with the same fake objects used by test_static.py.
        events = []
        fake_app = FakeWpsApplication(events, {})
        with mock.patch.object(main, "_connect_wps", return_value=(fake_app, mock.Mock(), "fake")):
            outputs = main.process_selected_source_files(self._settings("source.xlsx"),
                                                         log=lambda _msg: None,
                                                         on_source_outputs=callback_paths.append)
        self.assertEqual(len(outputs), 1)
        self.assertEqual(callback_paths, [outputs])
        self.assertNotEqual(Path(outputs[0]).resolve(), self.source.resolve())
        self.assertTrue(Path(outputs[0]).is_file())

    def test_process_workbooks_source_branch_returns_first_output_and_forwards_callback(self):
        expected = self.root / "processed.xlsx"
        callback = []
        on_outputs = callback.append
        with mock.patch.object(main, "process_selected_source_files", return_value=[expected]) as process:
            settings = {"source_processing_only": True, "marker": "selected"}
            result = main.process_workbooks(settings, log=lambda _msg: None,
                                            on_source_outputs=on_outputs)
        self.assertEqual(result, expected)
        process.assert_called_once()
        self.assertIs(process.call_args.args[2], on_outputs)


class SelectedSourceConfigurationTests(unittest.TestCase):
    def test_source_processing_only_survives_save_and_load(self):
        if not hasattr(main, "save_settings") or not hasattr(main, "load_settings"):
            self.skipTest("configuration helpers unavailable")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "settings.json"
            settings = {"source_processing_only": True, "workflow": list(OPS),
                        **{key: {"enabled": False, "rules": []} for key in OPS}}
            settings["Deduplicate"] = {"enabled": True, "rules": [source_rule(source_paths=["selected.xlsx"], source_output_filename="自定义结果")] }
            main.save_settings(settings, path)
            loaded = main.load_settings(path)
            self.assertTrue(loaded["source_processing_only"])
            self.assertEqual(loaded["workflow"], main.normalize_settings(settings)["workflow"])
            self.assertEqual(json.loads(path.read_text(encoding="utf-8"))["source_processing_only"], True)

    def test_source_only_requires_selected_source_rules_and_rejects_conflicts(self):
        if not hasattr(main, "validate_static_settings"):
            self.skipTest("configuration validation unavailable")
        settings = {"source_processing_only": True, "workflow": list(OPS),
                    **{key: {"enabled": False, "rules": []} for key in OPS}}
        # Missing source_paths is not a valid source-only request.
        settings["Deduplicate"] = {"enabled": True, "rules": [source_rule()]}
        with self.assertRaises(Exception):
            main.validate_static_settings(settings)
        settings["Deduplicate"] = {"enabled": True, "rules": [
            source_rule(source_paths=["selected.xlsx"])]}
        main.validate_static_settings(settings)
        # Both mutually exclusive standalone modes must be rejected.
        settings["merge_only"] = True
        with self.assertRaises(Exception):
            main.validate_static_settings(settings)


class RegionProcessingScopeTests(unittest.TestCase):
    def test_source_rule_adapter_converts_all_five_modules_to_target_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            workbook(source, {"Data": [["x", "a|2"], ["x", "b|3"]]})
            common = {"source_paths": [str(source)], "source_output_filename": "processed.xlsx"}
            rules = {
                "CustomFormula": source_rule(area="C1:C2", formula_mode="custom", formula="=1", **common),
                "TextToColumns": source_rule(area="B1:B2", delimiter="|", output_start_col="D", column_formats=["text", "number"], **common),
                "RegionFilter": source_rule(area="A1:B2", filter_column="A", operator="equals", criteria="x", has_header=False, **common),
                "Deduplicate": source_rule(area="A1:B2", has_header=False, key_mode="row", key_columns=[], **common),
                "FormatConvert": source_rule(area="A1:A2", mode="text", **common),
            }
            settings = {"source_processing_only": True, **{key: {"enabled": True, "rules": [rules[key]]} for key in OPS}}
            observed = []
            with mock.patch.object(main, "_connect_wps", return_value=(FakeWpsApplication([], {}), mock.Mock(), "fake")), \
                 mock.patch.object(main, "_apply_region_processing", side_effect=lambda wb, cfg, jobs, log, **kw: observed.append((cfg, jobs))):
                main.process_selected_source_files(settings, log=lambda msg: None)
            converted, jobs = observed[0]
            self.assertEqual(jobs, [])
            self.assertEqual(converted["jobs"], [])
            for key in OPS:
                self.assertTrue(converted[key]["enabled"])
                self.assertEqual(converted[key]["rules"][0]["scope"], "target")
                self.assertNotIn("source_paths", converted[key]["rules"][0])
                self.assertNotIn("source_output_filename", converted[key]["rules"][0])


class SelectedSourceExecutionTests(unittest.TestCase):
    def test_multi_source_deduplication_is_global_and_saved_beside_first_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "other").mkdir()
            first, second = root / "first.xlsx", root / "other" / "second.xlsx"
            workbook(first, {"Data": [["id"], [1], [1]]})
            workbook(second, {"Data": [["id"], [2], [1]]})
            before = [first.read_bytes(), second.read_bytes()]
            settings = {"source_processing_only": True, "Deduplicate": {"enabled": True, "rules": [
                source_rule(area="A1", source_paths=[str(first), str(second)], source_output_filename="合并去重", has_header=False, key_mode="row", key_columns=[])]}}
            with mock.patch.object(main, "_connect_wps", return_value=(FakeWpsApplication([], {}), mock.Mock(), "fake")):
                output = main.process_selected_source_files(settings, log=lambda msg: None)[0]
            self.assertEqual(output, root / "合并去重.xlsx")
            result = load_workbook(output)
            values = [row[0] for row in result["Data"].values if row[0] is not None]
            self.assertEqual(values, ["id", 1, 2])
            result.close()
            self.assertEqual(before, [first.read_bytes(), second.read_bytes()])

    def test_single_source_format_conversion_and_existing_output_are_safe(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, existing = root / "source.xlsx", root / "自定义.xlsx"
            workbook(source, {"Data": [["num"], ["001"], ["2.5"]], "Notes": [["keep"]]})
            existing.write_bytes(b"do-not-overwrite")
            settings = {"source_processing_only": True, "FormatConvert": {"enabled": True, "rules": [
                source_rule(area="A2:A3", mode="number", source_paths=[str(source)], source_output_filename="自定义.xlsx")]}}
            with mock.patch.object(main, "_connect_wps", return_value=(FakeWpsApplication([], {}), mock.Mock(), "fake")):
                output = main.process_selected_source_files(settings, log=lambda msg: None)[0]
            self.assertEqual(output.name, "自定义_2.xlsx")
            self.assertEqual(existing.read_bytes(), b"do-not-overwrite")
            result = load_workbook(output)
            self.assertEqual(result["Data"]["A2"].value, 1)
            self.assertEqual(result["Data"]["A3"].value, 2.5)
            self.assertEqual(result["Notes"]["A1"].value, "keep")
            result.close()

    def test_com_failure_does_not_publish_or_modify_original(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            workbook(source, {"Data": [[1]]})
            before = source.read_bytes()
            settings = {"source_processing_only": True, "FormatConvert": {"enabled": True, "rules": [
                source_rule(area="A1", mode="text", source_paths=[str(source)], source_output_filename="failed")]}}
            with mock.patch.object(main, "_connect_wps", side_effect=main.ProcessingError("COM failed")):
                with self.assertRaisesRegex(main.ProcessingError, "COM failed"):
                    main.process_selected_source_files(settings, log=lambda msg: None)
            self.assertFalse((root / "failed.xlsx").exists())
            self.assertEqual(source.read_bytes(), before)

    def test_blank_sheet_targets_first_input_sheet_even_when_other_rule_comes_first(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first, second = root / "first.xlsx", root / "second.xlsx"
            workbook(first, {"Data": [[1]], "Other": [[2]]})
            workbook(second, {"Other": [[3]], "Data": [[4]]})
            group = {"source_paths": [str(first), str(second)], "rules": {
                "FormatConvert": [source_rule(sheet="Other"), source_rule(sheet="")]}}
            copied = main._prepare_selected_source_copy(group, lambda msg: None)
            self.addCleanup(copied.unlink, missing_ok=True)
            result = load_workbook(copied)
            self.assertEqual(result.sheetnames, ["Data", "Other"])
            self.assertEqual(list(result["Data"].values), [(1,), (4,)])
            result.close()

    def test_custom_long_filename_is_not_silently_truncated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "temp.xlsx"
            source.write_bytes(b"completed")
            custom = "a" * 100 + ".xlsx"
            output = main._publish_processed_source_copy(source, root / "original.xlsx", root,
                                                         filename_stem=custom, exact_filename=True)
            self.assertEqual(output.name, custom)


class TemplateSelectedSourceOrderingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.template = self.root / "template.xlsx"
        workbook(self.template, {"Data": [["id"], [1]]})
        self.source = self.root / "source.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _settings(self, *, source_path=None, output_name="processed.xlsx",
                  selected=True):
        source_path = self.source if source_path is None else Path(source_path)
        rule = source_rule(source_paths=[str(source_path)],
                           source_output_filename=output_name)
        expected_output = source_path.parent / output_name
        settings = {
            "template_path": str(self.template),
            "output_dir": str(self.root),
            "output_filename": "result.xlsx",
            "jobs": [{"template_sheet": "Data", "sources": [
                {"path": str(expected_output), "source_sheet": "Data"}
            ]}],
            "CustomFormula": {"enabled": False, "rules": []},
            "TextToColumns": {"enabled": False, "rules": []},
            "RegionFilter": {"enabled": False, "rules": []},
            "Deduplicate": {"enabled": selected, "rules": [rule] if selected else []},
            "FormatConvert": {"enabled": False, "rules": []},
        }
        return settings

    def test_missing_template_source_is_generated_before_template_processing(self):
        settings = self._settings()
        generated = self.root / "processed.xlsx"
        observed = {}

        def generate(processed_settings, _log, _callback):
            generated.touch()
            return [generated]

        def template_process(processed_settings, *_args):
            observed["settings"] = processed_settings
            self.assertTrue(generated.is_file())
            return self.root / "result.xlsx"

        with mock.patch.object(main, "process_selected_source_files", side_effect=generate) as source_process, \
             mock.patch.object(main, "_process_template_workbooks", side_effect=template_process) as template_process_mock:
            result = main.process_workbooks(settings, log=lambda _msg: None)

        self.assertEqual(result, self.root / "result.xlsx")
        source_process.assert_called_once()
        template_process_mock.assert_called_once()
        self.assertEqual(observed["settings"]["jobs"][0]["sources"][0]["path"], str(generated))
        self.assertEqual(observed["settings"]["Deduplicate"]["rules"], [])
        self.assertFalse(observed["settings"]["Deduplicate"]["enabled"])

    def test_existing_expected_output_maps_template_source_to_避让_output(self):
        settings = self._settings()
        expected = self.root / "processed.xlsx"
        actual = self.root / "processed_2.xlsx"
        expected.touch()
        observed = {}

        def generate(_settings, _log, _callback):
            actual.touch()
            return [actual]

        def template_process(processed_settings, *_args):
            observed["source"] = processed_settings["jobs"][0]["sources"][0]["path"]
            return self.root / "result.xlsx"

        with mock.patch.object(main, "process_selected_source_files", side_effect=generate), \
             mock.patch.object(main, "_process_template_workbooks", side_effect=template_process):
            main.process_workbooks(settings, log=lambda _msg: None)

        self.assertEqual(observed["source"], str(actual))

    def test_source_processing_failure_does_not_enter_template_validation(self):
        settings = self._settings()
        failure = RuntimeError("source processing failed")
        with mock.patch.object(main, "process_selected_source_files", side_effect=failure), \
             mock.patch.object(main, "_process_template_workbooks") as template_process:
            with self.assertRaisesRegex(RuntimeError, "source processing failed"):
                main.process_workbooks(settings, log=lambda _msg: None)
        template_process.assert_not_called()

    def test_template_without_selected_rules_does_not_call_selected_processing(self):
        settings = self._settings(selected=False)
        with mock.patch.object(main, "process_selected_source_files") as source_process, \
             mock.patch.object(main, "_process_template_workbooks", return_value=self.root / "result.xlsx") as template_process:
            result = main.process_workbooks(settings, log=lambda _msg: None)

        self.assertEqual(result, self.root / "result.xlsx")
        source_process.assert_not_called()
        template_process.assert_called_once()


if __name__ == "__main__":
    unittest.main()
