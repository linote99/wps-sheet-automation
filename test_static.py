from __future__ import annotations

import json
import re
import tempfile
import unittest
import zipfile
from copy import copy, deepcopy
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table

import main
from make_fake_data import create_fake_data


class FakeComRange:
    def __init__(self):
        self.NumberFormat = "General"
        self.Value2 = None


class FakeMappedRange:
    def __init__(self, worksheet, address, events):
        self.worksheet = worksheet
        self.address = address
        self.events = events
        self.min_col, self.min_row, self.max_col, self.max_row = range_boundaries(address)

    @property
    def Formula(self):
        values = tuple(
            tuple(self.worksheet.cell(row=row, column=col).value
                  for col in range(self.min_col, self.max_col + 1))
            for row in range(self.min_row, self.max_row + 1)
        )
        return values[0][0] if len(values) == 1 and len(values[0]) == 1 else values

    @Formula.setter
    def Formula(self, value):
        self.events.append(("formula", self.address, value))
        self.worksheet.cell(row=self.min_row, column=self.min_col).value = value

    @property
    def FormulaR1C1(self):
        return self.worksheet.cell(row=self.min_row, column=self.min_col).value

    @FormulaR1C1.setter
    def FormulaR1C1(self, value):
        self.events.append(("formula-r1c1", self.address, value))
        for row in range(self.min_row, self.max_row + 1):
            for col in range(self.min_col, self.max_col + 1):
                self.worksheet.cell(row=row, column=col).value = value

    @property
    def MergeCells(self):
        return False

    @property
    def NumberFormat(self):
        return self.worksheet.cell(row=self.min_row, column=self.min_col).number_format

    @NumberFormat.setter
    def NumberFormat(self, value):
        for row in range(self.min_row, self.max_row + 1):
            for col in range(self.min_col, self.max_col + 1):
                self.worksheet.cell(row=row, column=col).number_format = value

    @property
    def Value2(self):
        self.events.append(("read", self.address))
        values = tuple(
            tuple(self.worksheet.cell(row=row, column=col).value
                  for col in range(self.min_col, self.max_col + 1))
            for row in range(self.min_row, self.max_row + 1)
        )
        return values[0][0] if len(values) == 1 and len(values[0]) == 1 else values

    @Value2.setter
    def Value2(self, values):
        self.events.append(("write", self.address))
        if not isinstance(values, (tuple, list)):
            values = ((values,),)
        elif values and not isinstance(values[0], (tuple, list)):
            values = tuple((value,) for value in values)
        for row_offset, row_values in enumerate(values):
            for col_offset, value in enumerate(row_values):
                self.worksheet.cell(row=self.min_row + row_offset, column=self.min_col + col_offset).value = value


class FakeComWorksheet:
    def __init__(self, worksheet, events):
        self.worksheet = worksheet
        self.events = events

    @property
    def Name(self):
        return self.worksheet.title

    @property
    def UsedRange(self):
        rows = [cell.row for cell in self.worksheet._cells.values() if cell.value is not None]
        cols = [cell.column for cell in self.worksheet._cells.values() if cell.value is not None]
        first = min(rows, default=1)
        last = max(rows, default=1)
        first_col = min(cols, default=1)
        last_col = max(cols, default=1)
        return type("FakeUsedRange", (), {
            "Row": first,
            "Rows": type("FakeRows", (), {"Count": last - first + 1})(),
            "Column": first_col,
            "Columns": type("FakeColumns", (), {"Count": last_col - first_col + 1})(),
        })()

    def Range(self, address):
        return FakeMappedRange(self.worksheet, address, self.events)


class FakeComWorksheetCollection:
    def __init__(self, workbook, events):
        self.workbook = workbook
        self.events = events

    @property
    def Count(self):
        return len(self.workbook.sheetnames)

    def Item(self, name_or_index):
        if isinstance(name_or_index, int):
            title = self.workbook.sheetnames[name_or_index - 1]
        else:
            requested = str(name_or_index).casefold()
            title = next((name for name in self.workbook.sheetnames if name.casefold() == requested), None)
            if title is None:
                raise KeyError(name_or_index)
        return FakeComWorksheet(self.workbook[title], self.events)

    def __call__(self, name_or_index):
        return self.Item(name_or_index)


class FakeComWorkbook:
    def __init__(self, path, events):
        self.path = Path(path)
        self.events = events
        self.workbook = load_workbook(self.path, data_only=False, keep_links=False)
        self.Worksheets = FakeComWorksheetCollection(self.workbook, events)

    def Save(self):
        self.events.append("save")
        self.workbook.save(self.path)

    def Close(self, SaveChanges=False):
        self.events.append("close")
        self.workbook.close()


class FakeWpsApplication:
    def __init__(self, events, calculated_values):
        self.events = events
        self.calculated_values = calculated_values
        self.Workbooks = self
        self.workbook = None
        self.CalculationState = 0

    def Open(self, path, **_kwargs):
        self.events.append("open")
        self.workbook = FakeComWorkbook(path, self.events)
        return self.workbook

    def CalculateFullRebuild(self):
        self.events.append("calculate")
        if self.workbook is not None and "1当天派单" in self.workbook.workbook.sheetnames:
            worksheet = self.workbook.workbook["1当天派单"]
            for row, value in self.calculated_values.items():
                worksheet[f"EJ{row}"] = value

    def Quit(self):
        self.events.append("quit")


class FakePythonCom:
    def __init__(self, events):
        self.events = events

    def CoUninitialize(self):
        self.events.append("co-uninitialize")

    def PumpWaitingMessages(self):
        self.events.append("pump-messages")


class FakeOutputFilenameVar:
    def __init__(self, value=""):
        self.value = value

    def get(self):
        return self.value

    def set(self, value):
        self.value = value


def make_test_workbook(path: Path, sheets: list[tuple[str, dict[int, dict[int, object]]]]) -> None:
    workbook = Workbook()
    for index, (title, rows) in enumerate(sheets):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        worksheet.title = title
        for row_number, columns in rows.items():
            for column_number, value in columns.items():
                worksheet.cell(row=row_number, column=column_number, value=value)
    workbook.save(path)
    workbook.close()


def rewrite_xlsx_member(path: Path, member_name: str, transform) -> None:
    """Modify one XML member of a synthetic XLSX fixture, never a supplied workbook."""
    temporary = path.with_name(f"{path.stem}.rewritten{path.suffix}")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as destination:
        destination.comment = source.comment
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == member_name:
                data = transform(data)
            destination.writestr(copy(info), data)
    temporary.replace(path)


def corrupt_autofilter_ref(path: Path, member_name: str, old_ref: str, new_ref: str) -> None:
    import re

    old = old_ref.encode("ascii")
    new = new_ref.encode("ascii")

    def replace(data: bytes) -> bytes:
        pattern = rb'(<autoFilter\b[^>]*\bref=")[^\"]*(")'
        matched, count = re.subn(pattern, lambda m: m.group(1) + new + m.group(2), data, count=1)
        if count != 1 or old not in data:
            raise AssertionError(f"未能在 {member_name} 中定位筛选区域 {old_ref!r}")
        return matched

    rewrite_xlsx_member(path, member_name, replace)


def corrupt_sortstate_ref(path: Path, member_name: str, old_ref: str, new_ref: str) -> None:
    old = f'ref="{old_ref}"'.encode("ascii")
    new = f'ref="{new_ref}"'.encode("ascii")

    def replace(data: bytes) -> bytes:
        pattern = rb"(<sortState\b[^>]*\b)ref=" + re.escape(old[4:])
        matched, count = re.subn(pattern, lambda match: match.group(1) + new, data, count=1)
        if count != 1 or old not in data:
            raise AssertionError(f"未能在 {member_name} 中定位排序区域 {old_ref!r}")
        return matched

    rewrite_xlsx_member(path, member_name, replace)


def make_sortstate_workbook(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "整行排序"
    first.append(["标题", "值"])
    first.append(["FAKE-1", 1])
    first.append(["FAKE-2", 2])
    first.auto_filter.ref = "A1:B3"
    first.auto_filter.add_sort_condition("A2:A3")
    second = workbook.create_sheet("A1排序")
    second.append(["标题", "值"])
    second.append(["FAKE-A", 1])
    second.auto_filter.ref = "A1:B2"
    second.auto_filter.add_sort_condition("A2:A2")
    workbook.save(path)
    workbook.close()
    corrupt_sortstate_ref(path, "xl/worksheets/sheet1.xml", "A1:B3", "3:393")


class WindowsLaunchPackagingTests(unittest.TestCase):
    project_dir = Path(__file__).resolve().parent

    def test_local_workflow_module_avoids_pyinstaller_contrib_name_collision(self):
        module_path = self.project_dir / "wps_automation_workflow.py"
        legacy_module_path = self.project_dir / "workflow.py"
        main_source = (self.project_dir / "main.py").read_text(encoding="utf-8")
        installer_path = self.project_dir / "install_and_build.bat"
        if not installer_path.is_file():
            installer_path = self.project_dir / "install_and_build_main_only.bat"
        installer = installer_path.read_text(encoding="utf-8")

        self.assertTrue(module_path.is_file())
        self.assertFalse(legacy_module_path.exists())
        self.assertIn("from wps_automation_workflow import", main_source)
        self.assertNotIn("from workflow import", main_source)
        compile_line = next(line for line in installer.splitlines() if "-m py_compile" in line)
        self.assertIn("wps_automation_workflow.py", compile_line)
        self.assertNotIn(" workflow.py", compile_line)

    def test_requirements_declare_windows_pywin32(self):
        requirements = (self.project_dir / "requirements.txt").read_text(encoding="utf-8")
        self.assertRegex(
            requirements,
            r"(?im)^\s*pywin32\s*>=\s*306\s*;\s*platform_system\s*==\s*['\"]Windows['\"]",
        )

    def test_installer_pins_runtime_commands_to_venv_and_checks_com_first(self):
        installer_path = self.project_dir / "install_and_build.bat"
        if not installer_path.is_file():
            installer_path = self.project_dir / "install_and_build_main_only.bat"
        installer = installer_path.read_text(encoding="utf-8").lower()
        self.assertIn('set "venv_py=%~dp0.venv\\scripts\\python.exe"', installer)
        self.assertTrue(any(f'set "log_file=%~dp0{name}"' in installer for name in ("install_and_build.log", "build_main.log")))
        self.assertIn('"%venv_py%" -c "import sys, importlib.metadata as metadata, pythoncom, win32com.client;', installer)
        self.assertIn('"%venv_py%" -m pip install -r "%~dp0requirements.txt"', installer)
        self.assertIn('"%venv_py%" -m pyinstaller', installer)
        self.assertIn("--hidden-import=pythoncom", installer)
        self.assertIn("--hidden-import=pywintypes", installer)
        self.assertIn("--hidden-import=win32timezone", installer)
        self.assertIn("--hidden-import=win32com.client", installer)
        self.assertIn("--collect-all win32com", installer)
        self.assertIn("goto failed", installer)
        self.assertIn("endlocal & exit /b 1", installer)
        self.assertIn("endlocal & exit /b 0", installer)
        self.assertIn("if errorlevel 1 goto failed", installer)
        self.assertNotIn("call \".venv\\scripts\\activate.bat\"", installer)

        python_lines = [
            line.strip() for line in installer.splitlines()
            if any(marker in line for marker in (
                "-m pip", "-m py_compile", " main.py --self-check", "-m unittest", "-m pyinstaller",
                "-c \"import sys, importlib.metadata",
            ))
        ]
        self.assertGreaterEqual(len(python_lines), 6)
        self.assertTrue(all(line.startswith('"%venv_py%"') for line in python_lines), python_lines)
        self.assertLess(installer.index("import sys, importlib.metadata"), installer.index("-m py_compile"))
        self.assertLess(installer.index("import sys, importlib.metadata"), installer.index("-m pyinstaller"))

    def test_run_app_resolves_its_directory_and_launches_with_venv_python(self):
        launcher = (self.project_dir / "run_app.bat").read_text(encoding="utf-8").lower()
        self.assertIn('cd /d "%~dp0"', launcher)
        self.assertIn('set "venv_py=%~dp0.venv\\scripts\\python.exe"', launcher)
        self.assertIn('"%venv_py%" -c "import sys, pythoncom, win32com.client;', launcher)
        self.assertIn('"%venv_py%" "%~dp0main.py" %*', launcher)
        self.assertTrue("install_and_build.log" in launcher or "build_main.log" in launcher)
        self.assertNotRegex(launcher, r"(?m)^\s*python(?:\.exe)?\s")

    def test_startup_authorization_dialog_is_visible_with_hidden_root(self):
        licensing_source = (self.project_dir / "licensing.py").read_text(encoding="utf-8")
        main_source = (self.project_dir / "main.py").read_text(encoding="utf-8")
        self.assertIn('if root.state() != "withdrawn":', licensing_source)
        self.assertNotIn("dialog.transient(root)\n    dialog.grab_set()", licensing_source)
        self.assertLess(licensing_source.index("dialog.wait_visibility()"), licensing_source.index("dialog.grab_set()"))
        self.assertIn('root.withdraw()', main_source)
        self.assertIn('_write_startup_log("授权通过，开始构建主界面")', main_source)

    def test_windowed_startup_failures_are_logged_and_shown(self):
        main_source = (self.project_dir / "main.py").read_text(encoding="utf-8")
        self.assertIn('"logs" / "startup.log"', main_source)
        self.assertIn("messagebox.showerror(APP_NAME, message, parent=root)", main_source)
        self.assertIn("traceback.format_exc()", main_source)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "logs" / "startup.log"
            with mock.patch.object(main, "_startup_log_path", return_value=path):
                self.assertEqual(path, main._write_startup_log("构建主界面失败", "诊断堆栈"))
            self.assertIn("构建主界面失败", path.read_text(encoding="utf-8"))
            self.assertIn("诊断堆栈", path.read_text(encoding="utf-8"))

    def test_missing_pywin32_guidance_identifies_wrong_interpreter_and_launcher(self):
        with mock.patch.object(main.sys, "executable", r"C:\Program Files\Python312\python.exe"):
            message = main._pywin32_import_error_message(ModuleNotFoundError("No module named 'pythoncom'"))
        self.assertIn(r"C:\Program Files\Python312\python.exe", message)
        self.assertIn("run_app.bat", message)
        self.assertIn("不要直接双击 main.py", message)
        self.assertTrue("install_and_build.log" in message or "build_main.log" in message)


class WindowSizingTests(unittest.TestCase):
    def test_default_size_and_display_limits(self):
        self.assertEqual((1500, 960), main.normalize_window_size(None))
        self.assertEqual((1024, 680, 1318, 688), main.window_size_limits(1366, 768))
        self.assertEqual((1318, 688), main.normalize_window_size(
            {"width": 1500, "height": 960}, 1366, 768
        ))
        self.assertEqual((1024, 680), main.normalize_window_size(
            {"width": 800, "height": 600}, 1366, 768
        ))

    def test_invalid_saved_dimensions_fall_back_to_defaults(self):
        invalid_values = (
            None,
            {},
            {"width": True, "height": 900},
            {"width": "wide", "height": 900},
            {"width": 100, "height": 900},
            {"width": 1200, "height": 20_000},
        )
        for value in invalid_values:
            with self.subTest(value=value):
                self.assertEqual((1500, 960), main.normalize_window_size(value))
        self.assertEqual((1318, 688), main.normalize_window_size(
            {"width": "not-a-number", "height": 960}, 1366, 768
        ))

    def test_tiny_display_is_bounded_without_display_access(self):
        min_width, min_height, max_width, max_height = main.window_size_limits(300, 200)
        width, height = main.normalize_window_size({"width": 900, "height": 700}, 300, 200)
        self.assertLessEqual(min_width, max_width)
        self.assertLessEqual(min_height, max_height)
        self.assertLessEqual(width, 300)
        self.assertLessEqual(height, 200)

    def test_window_state_file_contains_only_width_and_height(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested" / "window_size.json"
            self.assertTrue(main.save_window_size(1280, 820, path, 1920, 1080))
            self.assertEqual((1280, 820), main.load_window_size(path, 1920, 1080))
            self.assertEqual({"width": 1280, "height": 820}, json.loads(path.read_text(encoding="utf-8")))
            path.write_text("{broken", encoding="utf-8")
            self.assertEqual((1500, 960), main.load_window_size(path))


class StaticConfigurationTests(unittest.TestCase):
    def test_default_and_legacy_configuration(self):
        main.validate_static_settings(main.DEFAULT_SETTINGS)
        expected = [(2, 3), (2, 3), (3, 4)]
        actual = [(j["header_rows"], j["start_row"]) for j in main.DEFAULT_SETTINGS["jobs"]]
        self.assertEqual(expected, actual)
        self.assertTrue(all(not main.normalize_settings({**main.DEFAULT_SETTINGS, "jobs": [job]})["jobs"][0]["sources"]
                            for job in main.DEFAULT_SETTINGS["jobs"]))
        self.assertEqual([(132, 144), (132, 151), (132, 178)], [main.parse_column_range(j["formula_range"]) for j in main.DEFAULT_SETTINGS["jobs"]])
        legacy = {
            "sources": ["1.xlsx", "2.xlsx", "3.xlsx"], "source_start_row": 4,
            "jobs": [
                {"template_sheet": "一", "source_sheet": "源", "header_rows": 2, "start_row": 3, "formula_range": "EB:EN", "result_col": "EJ", "target_col": "DG"},
                {"template_sheet": "二", "source_sheet": "源", "header_rows": 2, "start_row": 3, "formula_range": "EB:EU", "result_col": "EQ", "target_col": "DG"},
                {"template_sheet": "三", "source_sheet": "源", "header_rows": 3, "start_row": 4, "formula_range": "EB:FV", "result_col": "FQ", "target_col": "DG"},
            ],
        }
        migrated = main.normalize_settings(legacy)
        self.assertEqual("1.xlsx", migrated["jobs"][0]["sources"][0]["path"])
        self.assertEqual(4, migrated["jobs"][0]["sources"][0]["copy_start_row"])
        self.assertEqual({"result_col": "EQ", "target_col": "DG", "write_as": "value"},
                         migrated["jobs"][1]["mappings"][0])
        self.assertIn("result_mapping", migrated["workflow"])
        main.validate_static_settings(migrated)
        migrated_v5 = main.normalize_settings({"config_version": 5, "jobs": deepcopy(main.DEFAULT_SETTINGS["jobs"])})
        self.assertEqual(main.CONFIG_VERSION, migrated_v5["config_version"])
        self.assertEqual(main.DEFAULT_WORKFLOW, migrated_v5["workflow"])
        self.assertEqual(5, migrated_v5["template_prefix_length"])

        legacy_v16 = deepcopy(main.DEFAULT_SETTINGS)
        legacy_v16["config_version"] = 16
        legacy_v16["RegionFilter"] = {"enabled": True, "rules": [{
            "scope": "source", "sheet": "", "range": "", "has_header": True,
            "filter_column": "A", "operator": "not_blank", "criteria": "",
        }]}
        migrated_v16 = main.normalize_settings(legacy_v16)
        self.assertEqual(main.CONFIG_VERSION, migrated_v16["config_version"])
        self.assertEqual("sheet", main._parse_region_filter_area(
            migrated_v16["RegionFilter"]["rules"][0]["range"]
        )[1])
        migration_message = " ".join(main.configuration_migration_notes(legacy_v16))
        self.assertIn("第一张表", migration_message)
        self.assertIn("不区分大小写", migration_message)
        self.assertIn("筛选范围可留空", migration_message)

        legacy_v12 = deepcopy(main.DEFAULT_SETTINGS)
        legacy_v12["config_version"] = 12
        legacy_v12["workflow"].remove("result_mapping")
        migrated_v12 = main.normalize_settings(legacy_v12)
        self.assertLess(migrated_v12["workflow"].index("formula_recalc"), migrated_v12["workflow"].index("result_mapping"))
        self.assertLess(migrated_v12["workflow"].index("result_mapping"), migrated_v12["workflow"].index("post_mapping"))

    def test_v6_v7_formula_seed_fields_are_ignored_and_migration_is_explained(self):
        for version in (6, 7):
            with self.subTest(version=version):
                legacy = deepcopy(main.DEFAULT_SETTINGS)
                legacy["config_version"] = version
                legacy["jobs"] = [deepcopy(legacy["jobs"][0])]
                legacy["jobs"][0].update({"start_row": 4, "formula_seed_row": 4})
                migrated = main.normalize_settings(legacy)
                self.assertEqual(main.CONFIG_VERSION, migrated["config_version"])
                self.assertEqual(4, migrated["jobs"][0]["start_row"])
                self.assertNotIn("formula_seed_row", migrated["jobs"][0])
                self.assertEqual(3, main._formula_template_row(migrated["jobs"][0]))
                notes = " ".join(main.configuration_migration_notes(legacy))
                self.assertIn("start_row - 1", notes)
                self.assertIn("formula_seed_row=4", notes)
                self.assertIn("不会用于选取公式", notes)

    def test_ranges_and_formula_validation(self):
        self.assertEqual(111, main.column_number("DG"))
        self.assertTrue(main.ranges_overlap((132, 144), (140, 150)))
        self.assertFalse(main.ranges_overlap((1, 10), (11, 20)))
        with self.assertRaises(main.ProcessingError):
            main.parse_column_range("EN:EB")
        with self.assertRaises(main.ProcessingError):
            main.column_number("XFE")
        overlapping_copy_and_formula = deepcopy(main.DEFAULT_SETTINGS)
        overlapping_copy_and_formula["jobs"][0]["sources"][0]["path"] = "C:/data/source.xlsx"
        overlapping_copy_and_formula["copy_start_col"] = "EB"
        overlapping_copy_and_formula["copy_end_col"] = "EB"
        with self.assertRaisesRegex(main.ProcessingError, "与公式区域 EB:EN 重叠"):
            main.validate_static_settings(overlapping_copy_and_formula)
        invalid = deepcopy(main.DEFAULT_SETTINGS)
        invalid["jobs"][0]["sources"][0]["path"] = "C:/data/source.xlsx"
        invalid["jobs"][0]["sources"][0]["copy_end_row"] = 1
        with self.assertRaises(main.ProcessingError):
            main.validate_static_settings(invalid)
        invalid_start = deepcopy(main.DEFAULT_SETTINGS)
        invalid_start["jobs"][0].update({"header_rows": 0, "start_row": 1})
        with self.assertRaisesRegex(main.ProcessingError, "至少为第 2 行"):
            main.validate_static_settings(invalid_start)

    def test_workflow_module_add_order_and_dependency_errors(self):
        workflow = list(main.DEFAULT_SETTINGS["workflow"])
        workflow.remove("TextToColumns")
        added = main.add_workflow_module(workflow, "TextToColumns")
        self.assertIn("TextToColumns", added)
        self.assertLess(added.index("TextToColumns"), added.index("RegionFilter"))
        self.assertLess(added.index("RegionFilter"), added.index("Deduplicate"))
        with self.assertRaisesRegex(ValueError, "已在工作流中"):
            main.add_workflow_module(added, "TextToColumns")
        removed = main.remove_workflow_module(added, added.index("TextToColumns"))
        self.assertNotIn("TextToColumns", removed)
        removable_settings = deepcopy(main.DEFAULT_SETTINGS)
        removable_settings["workflow"] = removed
        main.validate_static_settings(removable_settings)
        missing_core = main.remove_workflow_module(list(main.DEFAULT_SETTINGS["workflow"]),
                                                   main.DEFAULT_SETTINGS["workflow"].index("formula_fill"))
        with self.assertRaisesRegex(main.ProcessingError, "不能移除必需模块.*公式填充"):
            main.validate_static_settings({**deepcopy(main.DEFAULT_SETTINGS), "workflow": missing_core})

        reordered = list(main.DEFAULT_SETTINGS["workflow"])
        reordered = main.move_workflow_module(reordered, reordered.index("FormatConvert"), -2)
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["workflow"] = reordered
        main.validate_static_settings(settings)
        self.assertLess(reordered.index("FormatConvert"), reordered.index("RegionFilter"))

        invalid = deepcopy(main.DEFAULT_SETTINGS)
        invalid["workflow"] = main.move_workflow_module(
            list(invalid["workflow"]), invalid["workflow"].index("formula_fill"), 1
        )
        with self.assertRaisesRegex(main.ProcessingError, "公式填充.*首次完整重算"):
            main.validate_static_settings(invalid)

        invalid = deepcopy(main.DEFAULT_SETTINGS)
        invalid["TextToColumns"] = {"enabled": True, "rules": [{
            "sheet": "T", "range": "A1:A3", "delimiter": "|", "output_start_col": "H",
            "column_formats": ["text", "text"]}]}
        invalid["Deduplicate"] = {"enabled": True, "rules": [{
            "sheet": "T", "range": "H1:I3", "has_header": False, "key_mode": "row", "key_columns": []}]}
        invalid["workflow"] = main.move_workflow_module(
            list(invalid["workflow"]), invalid["workflow"].index("Deduplicate"), -2
        )
        with self.assertRaisesRegex(main.ProcessingError, "输入/输出区域与.*重叠"):
            main.validate_static_settings(invalid)

    def test_shared_area_selector_semantics_and_excel_limits(self):
        self.assertEqual(((1, 1, 16_384, 1_048_576), "sheet"), main._parse_selection_details("A1"))
        self.assertEqual(((1, 5, 16_384, 5), "row"), main._parse_selection_details("A5"))
        self.assertEqual(((4, 1, 4, 1_048_576), "column"), main._parse_selection_details("D1"))
        self.assertEqual(((2, 3, 2, 3), "cell"), main._parse_selection_details("B3,B3"))
        self.assertEqual(((2, 3, 4, 8), "rectangle"), main._parse_selection_details("D8,B3"))
        self.assertEqual(((2, 3, 4, 8), "rectangle"), main._parse_selection_details("D8:B3"))
        self.assertEqual(((16_384, 1_048_576, 16_384, 1_048_576), "cell"),
                         main._parse_selection_details("XFD1048576"))
        for invalid in ("", "A0", "XFE1", "A1048577", "A1,B3,C5", "not-a-cell"):
            with self.subTest(invalid=invalid), self.assertRaises(main.ProcessingError):
                main._parse_selection_details(invalid)
        with self.assertRaisesRegex(main.ProcessingError, "独立 XLSX 区域输出上限"):
            main._check_region_xlsx_area_size(main._parse_selected_area("A1:DV1000"), "测试区域")
        with self.assertRaisesRegex(main.ProcessingError, "处理上限"):
            main._check_processing_area_size(main._parse_selected_area("A1"), "测试去重")

    def test_mapping_area_accepts_column_only_as_full_column(self):
        for value, expected, canonical in (
            ("A", (1, 1, 1, 1_048_576), "A:A"),
            ("$DG", (111, 1, 111, 1_048_576), "DG:DG"),
            ("A:A", (1, 1, 1, 1_048_576), "A:A"),
            ("A:C", (1, 1, 3, 1_048_576), "A:C"),
        ):
            with self.subTest(value=value):
                bounds = main._parse_area(value, "映射源")
                self.assertEqual(expected, bounds)
                self.assertEqual(canonical, main._area_address(bounds))

        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["pre_mappings"] = [{
            "source_sheet": "源", "source_range": "A",
            "target_sheet": "目标", "target_range": "B",
        }]
        settings["post_mappings"] = [{
            "source_sheet": "目标", "source_range": "C:C",
            "target_sheet": "结果", "target_range": "D",
        }]
        main.validate_static_settings(settings)

        settings["pre_mappings"][0]["target_range"] = "B1:B5"
        with self.assertRaisesRegex(main.ProcessingError, "行列数必须相同"):
            main.validate_static_settings(settings)

        for invalid in ("XFE", "A0", "A1:B", "A:C1"):
            with self.subTest(invalid=invalid), self.assertRaises(main.ProcessingError):
                main._parse_area(invalid, "映射区域")

    def test_post_mapping_full_columns_start_at_each_jobs_data_row(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        records = [[["one"]], [["two"]], [["three"], ["four"]]]
        columns = ("EJ", "EQ", "FQ")
        settings["post_mappings"] = [
            {"source_sheet": job["template_sheet"], "source_range": f"{column}:{column}",
             "target_sheet": job["template_sheet"], "target_range": "DG:DG"}
            for job, column in zip(settings["jobs"], columns)
        ]

        main._validate_post_mapping_conflicts(settings, records)
        expected_rows = ((3, 3), (3, 3), (4, 5))
        job_data_ranges = main._configured_job_data_ranges(settings["jobs"], source_records=records)
        for index, (job, column, rows) in enumerate(zip(settings["jobs"], columns, expected_rows), 1):
            resolved = main._resolve_post_mapping_areas(settings["post_mappings"][index - 1], job_data_ranges, index)
            self.assertIsNotNone(resolved)
            source_bounds, target_bounds = resolved
            self.assertEqual((column_index_from_string(column), rows[0], column_index_from_string(column), rows[1]),
                             source_bounds)
            self.assertEqual((column_index_from_string("DG"), rows[0], column_index_from_string("DG"), rows[1]),
                             target_bounds)
            self.assertEqual(job["start_row"], rows[0])

        settings["post_mappings"][0]["source_range"] = "EJ2"
        settings["post_mappings"][0]["target_range"] = "DG2"
        with self.assertRaisesRegex(main.ProcessingError, "DG2.*会覆盖保留公式区域"):
            main._validate_post_mapping_conflicts(settings, records)

    def test_post_mapping_is_optional_but_configured_rules_require_workflow_stage(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["workflow"].remove("post_mapping")
        main.validate_static_settings(settings)

        settings["post_mappings"] = [{
            "source_sheet": "1当天派单", "source_range": "EJ3:EJ3",
            "target_sheet": "1当天派单", "target_range": "DG3:DG3",
        }]
        with self.assertRaisesRegex(main.ProcessingError, "工作流中移除了“后映射”模块"):
            main.validate_static_settings(settings)

    def test_result_mapping_is_optional_but_configured_rules_require_workflow_stage(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["workflow"].remove("result_mapping")
        main.validate_static_settings(settings)

        settings["jobs"][0]["mappings"] = [{"result_col": "EJ", "target_col": "DG", "write_as": "text"}]
        with self.assertRaisesRegex(main.ProcessingError, "工作流中移除了“结果映射”模块"):
            main.validate_static_settings(settings)

        settings["workflow"] = main.add_workflow_module(settings["workflow"], "result_mapping")
        main.validate_static_settings(settings)

        settings["jobs"][0]["mappings"].append({"result_col": "EJ", "target_col": "DG", "write_as": "value"})
        with self.assertRaisesRegex(main.ProcessingError, "都写入目标列 DG"):
            main.validate_static_settings(settings)

    def test_result_mapping_targets_cannot_be_overwritten_by_pre_or_post_mapping(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["jobs"] = [deepcopy(settings["jobs"][0])]
        settings["jobs"][0]["mappings"] = [{"result_col": "EJ", "target_col": "DG", "write_as": "value"}]
        records = [[ ["one"] ]]
        settings["pre_mappings"] = [{"source_sheet": "S", "source_range": "A1",
                                      "target_sheet": "1当天派单", "target_range": "DG3"}]
        with self.assertRaisesRegex(main.ProcessingError, "与结果映射区域 DG3 冲突"):
            main._validate_pre_mapping_conflicts(settings, records)

        settings["pre_mappings"] = []
        settings["post_mappings"] = [{"source_sheet": "1当天派单", "source_range": "A3",
                                       "target_sheet": "1当天派单", "target_range": "DG3"}]
        with self.assertRaisesRegex(main.ProcessingError, "与结果映射目标 DG3 重复"):
            main._validate_post_mapping_conflicts(settings, records)

    def test_region_formula_cannot_overwrite_result_mapping_target(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "mapping-protection.xlsx"
            make_test_workbook(path, [("Data", {3: {1: "x"}, 4: {1: "y"}})])
            fake_workbook = FakeComWorkbook(path, [])
            mappings = [{"result_col": "A", "target_col": "D", "write_as": "value"}]
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [{"template_sheet": "Data", "start_row": 3, "formula_range": "A:A",
                                  "mappings": mappings}]
            settings["CustomFormula"] = {"enabled": True, "rules": [{
                "scope": "target", "sheet": "Data", "range": "D3:D4", "formula_mode": "custom",
                "formula": "=1", "seed_cell": ""}]}
            try:
                with self.assertRaisesRegex(main.ProcessingError, "与结果映射目标区域重叠"):
                    main._apply_region_processing(fake_workbook, settings,
                        [("Data", 3, 4, mappings)], lambda _msg: None)
            finally:
                fake_workbook.workbook.close()

    def test_sample_json_validates(self):
        sample = json.loads((Path(__file__).parent / "default_config.json").read_text(encoding="utf-8"))
        main.validate_static_settings(sample)
        self.assertEqual([3, 3, 4], [job["start_row"] for job in sample["jobs"]])
        self.assertEqual(["EB:EN", "EB:EU", "EB:FV"], [job["formula_range"] for job in sample["jobs"]])
        self.assertEqual(main.DEFAULT_WORKFLOW, sample["workflow"])
        self.assertEqual(main.CONFIG_VERSION, sample["config_version"])

    def test_merged_source_filename_is_optional_but_validated(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["jobs"] = [deepcopy(settings["jobs"][0])]
        settings["jobs"][0]["merged_source_filename"] = ""
        main.validate_static_settings(settings)
        settings["jobs"][0]["merged_source_filename"] = "custom-source.xlsx"
        main.validate_static_settings(settings)
        for filename, error in (("../escape.xlsx", "不能包含目录路径"),
                                ("source.xlsm", "只能使用 .xlsx"),
                                ("CON.xlsx", "保留设备名")):
            settings["jobs"][0]["merged_source_filename"] = filename
            with self.subTest(filename=filename), self.assertRaisesRegex(main.ProcessingError, error):
                main.validate_static_settings(settings)

    def test_missing_source_path_is_located_and_reported_before_processing(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["jobs"][1]["sources"][0]["path"] = "C:/data/second.xlsx"
        self.assertIsNone(main._first_missing_source_location(settings))
        settings["jobs"][0]["sources"][0]["path"] = "  "
        settings["jobs"][0]["sources"][0].pop("_optional_placeholder", None)
        self.assertEqual((0, 0), main._first_missing_source_location(settings))
        settings["jobs"][0]["sources"][0]["path"] = "C:/data/first.xlsx"
        settings["jobs"][2]["sources"][0].pop("_optional_placeholder", None)
        self.assertEqual((2, 0), main._first_missing_source_location(settings))
        settings["jobs"][2]["sources"][0]["path"] = "C:/data/third.xlsx"
        self.assertIsNone(main._first_missing_source_location(settings))

        with tempfile.TemporaryDirectory() as td:
            fake = create_fake_data(Path(td) / "fake")
            template_only = deepcopy(main.DEFAULT_SETTINGS)
            template_only["template_path"] = str(fake / "template.xlsx")
            template_only["output_dir"] = str(fake / "output")
            template_only["copy_start_col"] = ""
            validated_template, _final = main._validate_inputs(template_only)
            self.assertEqual(Path(template_only["template_path"]).resolve(), validated_template)

    def test_output_name_matching_template_in_same_directory_is_rejected_without_overwrite(self):
        with tempfile.TemporaryDirectory() as td:
            fake = create_fake_data(Path(td) / "fake")
            template = fake / "template.xlsx"
            original = template.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"template_path": str(template), "output_dir": str(fake),
                             "output_filename": template.name, "output_filename_custom": False,
                             "copy_end_col": "C"})
            for index, job in enumerate(settings["jobs"], 1):
                job["sources"][0]["path"] = str(fake / f"{index}.xlsx")
            with self.assertRaisesRegex(main.ProcessingError, "不能覆盖模板"):
                main._validate_inputs(settings)
            self.assertEqual(original, template.read_bytes())

    def test_save_load_reusable_configuration_round_trip(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "reusable.json"
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"][0]["sources"] = [
                {"path": "C:/data/one.xlsx", "source_sheet": "甲", "copy_start_row": 5, "copy_end_row": 8, "keep_first_row": True},
                {"path": "C:/data/two.xlsx", "source_sheet": "乙", "copy_start_row": 2, "copy_end_row": None, "keep_first_row": False},
            ]
            settings["pre_mappings"] = [{"source_sheet": "模板甲", "source_range": "A1:B2", "target_sheet": "模板乙", "target_range": "C3:D4"}]
            settings["post_mappings"] = [{"source_sheet": "模板乙", "source_range": "A1:B2", "target_sheet": "模板丙", "target_range": "C3:D4"}]
            settings["region_xlsx_outputs"] = [{"sheet": "模板丙", "range": "A1:B2"}]
            settings["text_outputs"] = [{"sheet": "模板丙", "range": "A1"}]
            settings["retain_merged_sources"] = True
            settings["retain_edited_sources"] = True
            settings["jobs"][0]["merged_source_filename"] = "源数据副本.xlsx"
            settings["TextToColumns"] = {"enabled": True, "rules": [{"scope": "source", "sheet": "模板甲", "range": "A2:A8", "delimiter": "|", "output_start_col": "D", "column_formats": ["text", "number"]}]}
            settings["RegionFilter"] = {"enabled": True, "rules": [{"scope": "target", "sheet": "模板乙", "range": "A1:C8", "has_header": True, "filter_column": "A", "operator": "equals", "criteria": "OK"}]}
            settings["Deduplicate"] = {"enabled": True, "rules": [{"scope": "target", "sheet": "模板乙", "range": "A1:C8", "has_header": True, "key_mode": "columns", "key_columns": ["A", "C"]}]}
            settings["FormatConvert"] = {"enabled": True, "rules": [{"scope": "target", "sheet": "模板丙", "range": "D2:D8", "mode": "date"}]}
            settings["config_name"] = "月度汇总配置"
            settings["template_prefix_length"] = 7
            main.save_settings(settings, path)
            loaded = main.load_settings(path)
            self.assertEqual(settings["jobs"][0], loaded["jobs"][0])
            self.assertEqual(main.CONFIG_VERSION, loaded["config_version"])
            for key in ("pre_mappings", "post_mappings", "region_xlsx_outputs", "text_outputs", "TextToColumns", "RegionFilter", "Deduplicate", "FormatConvert", "workflow", "config_name", "template_prefix_length", "retain_edited_sources"):
                self.assertEqual(settings[key], loaded[key])
            self.assertTrue(loaded["retain_merged_sources"])
        self.assertEqual([], loaded["jobs"][0]["mappings"])

    def test_named_profile_save_load_and_template_prefix_match_cardinality(self):
        with tempfile.TemporaryDirectory() as td:
            profile_dir = Path(td) / "profiles"
            daily = deepcopy(main.DEFAULT_SETTINGS)
            daily.update({"config_name": "日报流程", "template_path": "C:/templates/AB日报模板.xlsx", "template_prefix_length": 4})
            weekly = deepcopy(main.DEFAULT_SETTINGS)
            weekly.update({"config_name": "周报流程", "template_path": "C:/templates/AB周报模板.xlsx", "template_prefix_length": 2})
            main.save_profile(daily, "日报流程", profile_dir)
            main.save_profile(weekly, "周报流程", profile_dir)
            self.assertEqual(["周报流程", "日报流程"], main.list_profiles(profile_dir))
            loaded = main.load_profile("日报流程", profile_dir)
            self.assertEqual("日报流程", loaded["config_name"])
            self.assertEqual(4, loaded["template_prefix_length"])
            self.assertEqual("C:/templates/AB日报模板.xlsx", loaded["template_path"])

            selected = "D:/input/AB日报_20261005.xlsx"
            self.assertEqual([], main.matching_profiles("D:/input/ZZ库存.xlsx", 2, profile_dir))
            self.assertEqual(["日报流程"], main.matching_profiles(selected, 4, profile_dir))
            self.assertEqual(["周报流程", "日报流程"], main.matching_profiles(selected, 2, profile_dir))
            with self.assertRaisesRegex(ValueError, "介于 1 和 120"):
                main.matching_profiles(selected, 0, profile_dir)

    def test_delete_named_profile_removes_only_the_selected_profile(self):
        with tempfile.TemporaryDirectory() as td:
            profile_dir = Path(td) / "profiles"
            for name in ("月报配置", "周报配置"):
                main.save_profile({**deepcopy(main.DEFAULT_SETTINGS), "config_name": name}, name, profile_dir)
            deleted_path = main.delete_profile("月报配置", profile_dir)
            self.assertFalse(deleted_path.exists())
            self.assertEqual(["周报配置"], main.list_profiles(profile_dir))
            self.assertEqual("周报配置", main.load_profile("周报配置", profile_dir)["config_name"])
            with self.assertRaisesRegex(ValueError, "找不到已保存配置"):
                main.delete_profile("月报配置", profile_dir)

    def test_template_selection_updates_automatic_name_but_keeps_explicit_custom_name(self):
        app = main.App.__new__(main.App)
        app.vars = {
            "template_path": FakeOutputFilenameVar(""),
            "output_filename": FakeOutputFilenameVar("处理结果.xlsx"),
        }
        app.output_filename_custom = False
        app._suppress_output_filename_trace = False

        app._set_template_selected("C:/templates/装移机在途通报V92026.10.5.xlsx")
        self.assertEqual("装移机在途通报V92026.10.5.xlsx", app.vars["output_filename"].get())
        app._set_template_selected("D:/templates/another.xlsx")
        self.assertEqual("another.xlsx", app.vars["output_filename"].get())

        app.vars["output_filename"].set("手动命名.xlsx")
        app._on_output_filename_edited()
        self.assertTrue(app.output_filename_custom)
        app._set_template_selected("E:/templates/third.xlsx")
        self.assertEqual("手动命名.xlsx", app.vars["output_filename"].get())
        self.assertEqual("E:/templates/third.xlsx", app.vars["template_path"].get())

    def test_profile_load_uses_selected_template_for_auto_name_and_keeps_profile_custom_name(self):
        with tempfile.TemporaryDirectory() as td:
            profile_dir = Path(td) / "profiles"
            auto = deepcopy(main.DEFAULT_SETTINGS)
            auto.update({"config_name": "自动命名", "template_path": "C:/old/old.xlsx",
                         "output_filename": "old.xlsx", "output_filename_custom": False})
            custom = deepcopy(main.DEFAULT_SETTINGS)
            custom.update({"config_name": "自定义命名", "template_path": "C:/old/old.xlsx",
                           "output_filename": "保留名称.xlsx", "output_filename_custom": True})
            main.save_profile(auto, "自动命名", profile_dir)
            main.save_profile(custom, "自定义命名", profile_dir)

            captured = {}
            app = main.App.__new__(main.App)
            app._apply_settings_to_gui = lambda settings, name: captured.update(settings=settings, name=name)
            with mock.patch.object(main, "PROFILE_DIR", profile_dir):
                app._load_profile_named("自动命名", selected_template_path="D:/selected/新模板.xlsx")
                self.assertEqual("D:/selected/新模板.xlsx", captured["settings"]["template_path"])
                self.assertEqual("新模板.xlsx", captured["settings"]["output_filename"])
                self.assertFalse(captured["settings"]["output_filename_custom"])

                app._load_profile_named("自定义命名", selected_template_path="D:/selected/另一模板.xlsx")
                self.assertEqual("D:/selected/另一模板.xlsx", captured["settings"]["template_path"])
                self.assertEqual("保留名称.xlsx", captured["settings"]["output_filename"])
                self.assertTrue(captured["settings"]["output_filename_custom"])

    def test_legacy_output_filename_migration_and_explicit_custom_marker(self):
        legacy_auto = deepcopy(main.DEFAULT_SETTINGS)
        legacy_auto.pop("output_filename_custom", None)
        legacy_auto.update({"template_path": r"C:\\templates\\legacy.xlsx", "output_filename": "处理结果.xlsx"})
        migrated_auto = main.normalize_settings(legacy_auto)
        self.assertEqual("legacy.xlsx", migrated_auto["output_filename"])
        self.assertFalse(migrated_auto["output_filename_custom"])

        legacy_custom = deepcopy(legacy_auto)
        legacy_custom["output_filename"] = "用户手动名称.xlsx"
        migrated_custom = main.normalize_settings(legacy_custom)
        self.assertEqual("用户手动名称.xlsx", migrated_custom["output_filename"])
        self.assertTrue(migrated_custom["output_filename_custom"])

    def test_pure_region_processing_and_text_area_format(self):
        split, columns = main.split_text_column(["001|2024-02-01", "002|2024-02-02"], "|", ["number", "date"])
        self.assertEqual(2, columns)
        self.assertEqual([[1, __import__("datetime").datetime(2024, 2, 1)],
                          [2, __import__("datetime").datetime(2024, 2, 2)]], split)
        deduplicated, removed = main.deduplicate_matrix(
            [["ID", "Name"], ["1", "first"], ["1", "duplicate"], ["2", "next"]],
            has_header=True, key_mode="columns", key_indexes=[0],
        )
        self.assertEqual(1, removed)
        self.assertEqual([["ID", "Name"], ["1", "first"], ["2", "next"], [None, None]], deduplicated)
        self.assertEqual(1234.5, main.convert_cell_value("1,234.5", "number"))
        self.assertEqual("42", main.convert_cell_value(42, "text"))
        self.assertEqual(__import__("datetime").datetime(2024, 3, 5), main.convert_cell_value("20240305", "date"))
        self.assertEqual("x\t1\ny\t", main._format_text_matrix([["x", 1], ["y", None]]))
        filtered, removed, keep_indexes = main.filter_matrix(
            [["ID", "score"], ["a", 12], ["b", 7], ["c", 15]], True, 1, "greater_equal", "12")
        self.assertEqual([["ID", "score"], ["a", 12], ["c", 15], [None, None]], filtered)
        self.assertEqual(1, removed)
        self.assertEqual({0, 1, 3}, keep_indexes)
        self.assertTrue(main._filter_value_matches("Alpha", "contains", "ph"))
        self.assertTrue(main._filter_value_matches(None, "blank", ""))
        self.assertFalse(main._filter_value_matches("Alpha", "equals", "alphaX"))

    def test_invalid_processing_step_rules_are_rejected(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["TextToColumns"] = {"enabled": True, "rules": [{"sheet": "T", "range": "A1:B2", "delimiter": ",", "output_start_col": "C", "column_formats": ["text"]}]}
        with self.assertRaisesRegex(main.ProcessingError, "源范围必须是单列"):
            main.validate_static_settings(settings)
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["Deduplicate"] = {"enabled": True, "rules": [{"sheet": "T", "range": "A1:C4", "has_header": False, "key_mode": "columns", "key_columns": ["D"]}]}
        with self.assertRaisesRegex(main.ProcessingError, "不在所选区域内"):
            main.validate_static_settings(settings)
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["RegionFilter"] = {"enabled": True, "rules": [{"scope": "target", "sheet": "T", "range": "B2:D9", "filter_column": "A", "operator": "equals", "criteria": "x"}]}
        with self.assertRaisesRegex(main.ProcessingError, "条件列必须位于筛选区域内"):
            main.validate_static_settings(settings)

    def test_pre_mapping_reads_original_template_before_cleanup_and_uses_independent_sheets(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            template_path = fake / "template.xlsx"
            workbook = load_workbook(template_path)
            workbook["1当天派单"]["A20"] = "PRE-SOURCE-ONE"
            workbook["2当天归档"]["B20"] = "PRE-SOURCE-TWO"
            workbook.save(template_path)
            workbook.close()
            original_bytes = template_path.read_bytes()

            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["template_path"] = str(template_path)
            settings["pre_mappings"] = [
                {"source_sheet": "1当天派单", "source_range": "A20", "target_sheet": "2当天归档", "target_range": "Z1"},
                {"source_sheet": "2当天归档", "source_range": "B20", "target_sheet": "3装移全业务在途", "target_range": "Y1"},
            ]
            staging = root / "pre-mapped.xlsx"
            records = [[ ["D1", "x", 1] ], [["D2", "y", 2]], [["D3", "z", 3]]]
            main._create_staging_workbook(template_path, settings, records, staging, lambda _msg: None)
            result = load_workbook(staging, data_only=False)
            try:
                self.assertIsNone(result["1当天派单"]["A20"].value)
                self.assertEqual("PRE-SOURCE-ONE", result["2当天归档"]["Z1"].value)
                self.assertEqual("PRE-SOURCE-TWO", result["3装移全业务在途"]["Y1"].value)
            finally:
                result.close()
            self.assertEqual(original_bytes, template_path.read_bytes())

    def test_pre_mapping_conflict_with_main_write_area_is_rejected_before_staging(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["copy_end_col"] = "C"
        settings["pre_mappings"] = [{"source_sheet": "S", "source_range": "A1", "target_sheet": "1当天派单", "target_range": "B3"}]
        records = [[["ROW", "x", 1]], [["ROW2", "y", 2]], [["ROW3", "z", 3]]]
        with self.assertRaisesRegex(main.ProcessingError, "前映射目标.*与主数据复制区域.*冲突"):
            main._validate_pre_mapping_conflicts(settings, records)

    def test_pre_mapping_cannot_overwrite_formula_template_row(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["copy_end_col"] = "C"
        settings["jobs"] = [deepcopy(settings["jobs"][0])]
        settings["jobs"][0]["start_row"] = 3
        settings["pre_mappings"] = [{
            "source_sheet": "源", "source_range": "A1", "target_sheet": "1当天派单", "target_range": "EB2"
        }]
        with self.assertRaisesRegex(main.ProcessingError, "前映射目标.*与公式模板区域.*冲突"):
            main._validate_pre_mapping_conflicts(settings, [[ ["DATA"] ]])

    def test_multiple_region_outputs_are_standalone_xlsx_and_never_overwrite(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "工单表"
            worksheet["A1"] = "header"
            worksheet["B2"] = 12
            worksheet["D3"] = "right"
            worksheet["E4"] = 45
            workbook.save(source_path)
            workbook.close()

            existing = output_dir / "工单表1.xlsx"
            output_dir.mkdir()
            existing.write_bytes(b"USER-FILE")
            paths = main._export_table_areas(
                source_path,
                [{"sheet": "工单表", "range": "A1:B2"}, {"sheet": "工单表", "range": "D3:E4"}],
                output_dir, lambda _msg: None,
            )
            self.assertEqual(
                ["工单表1_2.xlsx", "工单表2.xlsx"],
                [path.name for path in paths],
            )
            self.assertEqual(b"USER-FILE", existing.read_bytes())
            first = load_workbook(paths[0], data_only=False)
            second = load_workbook(paths[1], data_only=False)
            try:
                self.assertEqual(["工单表"], first.sheetnames)
                self.assertEqual("header", first.active["A1"].value)
                self.assertEqual("12", first.active["B2"].value)
                self.assertEqual("s", first.active["B2"].data_type)
                self.assertEqual("right", second.active["A1"].value)
                self.assertEqual("45", second.active["B2"].value)
            finally:
                first.close()
                second.close()

    def test_region_export_sanitizes_invalid_autofilter_on_copy_without_touching_saved_source(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "2026.10.6_fake.staging.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "已保存结果"
            worksheet.append(["名称", "数量"])
            worksheet.append(["FAKE-ROW", 42])
            worksheet.auto_filter.ref = "A1:B2"
            workbook.save(source_path)
            workbook.close()
            corrupt_autofilter_ref(source_path, "xl/worksheets/sheet1.xml", "A1:B2", "A1:BAD")
            saved_bytes = source_path.read_bytes()
            saved_mtime_ns = source_path.stat().st_mtime_ns
            logs = []
            load_paths = []
            real_load = main.load_workbook

            def tracked_load(path, *args, **kwargs):
                load_paths.append(Path(path))
                return real_load(path, *args, **kwargs)

            with mock.patch.object(main, "load_workbook", side_effect=tracked_load):
                paths = main._export_table_areas(
                    source_path, [{"sheet": "已保存结果", "range": "A1:B2"}], output_dir, logs.append
                )
                previews = main._read_text_output_areas(
                    source_path, [{"sheet": "已保存结果", "range": "A1:B2"}], logs.append
                )

            self.assertEqual(1, len(paths))
            self.assertEqual("已保存结果.xlsx", paths[0].name)
            self.assertEqual(2, len(load_paths))
            self.assertTrue(all(path.name.startswith("wps_sheet_autofilter_") for path in load_paths))
            self.assertNotIn(source_path, load_paths)
            self.assertEqual(saved_bytes, source_path.read_bytes())
            self.assertEqual(saved_mtime_ns, source_path.stat().st_mtime_ns)
            self.assertTrue(any("自动筛选区域 'A1:BAD' 无效" in message for message in logs))
            self.assertEqual("名称\t数量\nFAKE-ROW\t42", previews[0]["text"])
            exported = load_workbook(paths[0], data_only=False)
            try:
                self.assertEqual("FAKE-ROW", exported["已保存结果"]["A2"].value)
                self.assertEqual("42", exported["已保存结果"]["B2"].value)
                self.assertEqual("s", exported["已保存结果"]["B2"].data_type)
            finally:
                exported.close()

    def test_single_region_export_uses_template_sheet_name_and_reserves_main_workbook_name(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "saved.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "工资表"
            worksheet["A1"] = "员工"
            worksheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
            worksheet["A1"].font = Font(color="FF0000FF", bold=True)
            worksheet["A2"] = 12.3
            worksheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FF00FF00")
            worksheet["A2"].font = Font(color="FFFF0000")
            worksheet["A2"].number_format = "0.00"
            workbook.save(source_path)
            workbook.close()
            reserved_main_path = output_dir / "工资表.xlsx"
            paths = main._export_table_areas(
                source_path, [{"sheet": "工资表", "range": "A1:A2"}], output_dir,
                lambda _msg: None, reserved_paths={reserved_main_path},
            )
            self.assertEqual("工资表_2.xlsx", paths[0].name)
            self.assertFalse(reserved_main_path.exists())
            self.assertTrue(paths[0].is_file())
            exported = load_workbook(paths[0], data_only=False)
            try:
                ws = exported["工资表"]
                self.assertEqual("FFFF0000", ws["A1"].fill.fgColor.rgb)
                self.assertEqual("FF0000FF", ws["A1"].font.color.rgb)
                self.assertEqual("FF00FF00", ws["A2"].fill.fgColor.rgb)
                self.assertEqual("FFFF0000", ws["A2"].font.color.rgb)
                self.assertEqual("0.00", ws["A2"].number_format)
                self.assertEqual("12.30", ws["A2"].value)
            finally:
                exported.close()

    def test_region_xlsx_is_exported_without_image_generation(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"copy_end_col": "C", "template_path": str(fake / "template.xlsx"),
                             "output_dir": str(fake / "output"), "output_filename": "main-only.xlsx",
                             "region_xlsx_outputs": [{"sheet": "1当天派单", "range": "A1:B2"}]})
            for idx, job in enumerate(settings["jobs"], 1):
                job["sources"] = [{"path": str(fake / f"{idx}.xlsx"), "source_sheet": "示例源表",
                                   "copy_start_row": 2, "copy_end_row": 3, "keep_first_row": True}]
            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _settings):
                workbook = load_workbook(path, data_only=False)
                workbook.save(path); workbook.close()
            region_outputs = []
            with mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate), \
                 mock.patch.object(main, "_read_wps_display_text", return_value={}):
                result = main.process_workbooks(settings, lambda _msg: None,
                                                on_region_outputs=region_outputs.extend)
            self.assertTrue(result.is_file())
            self.assertTrue((Path(settings["output_dir"]) / "1当天派单.xlsx").is_file())
            self.assertEqual([Path(settings["output_dir"]) / "1当天派单.xlsx"], region_outputs)
            self.assertEqual([".xlsx"], [path.suffix.lower() for path in region_outputs])

    def test_region_export_configuration_is_preserved_and_old_auto_open_removed(self):
        old = deepcopy(main.DEFAULT_SETTINGS)
        old["config_version"] = 8
        old["open_region_exports"] = True
        old["export_region_pngs"] = True
        old.pop("region_xlsx_outputs")
        old["image_outputs"] = [{"sheet": "旧表", "range": "A1:B2"}]
        old["workflow"].remove("region_xlsx_output")
        old["workflow"].append("image_output")
        normalized = main.normalize_settings(old)
        self.assertEqual(main.CONFIG_VERSION, normalized["config_version"])
        self.assertNotIn("open_region_exports", normalized)
        self.assertNotIn("export_region_pngs", normalized)
        self.assertNotIn("image_outputs", normalized)
        self.assertEqual(old["image_outputs"], normalized["region_xlsx_outputs"])
        self.assertIn("region_xlsx_output", normalized["workflow"])
        self.assertTrue(any("自动打开区域 XLSX" in note
                            for note in main.configuration_migration_notes(old)))

    def test_region_xlsx_preserves_styles_merges_dimensions_and_hidden_axes(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "工单表"
            worksheet.merge_cells("A1:C1")
            worksheet["A1"] = "标题"
            worksheet["A1"].font = Font(bold=True, color="FFFFFF")
            worksheet["A1"].fill = PatternFill("solid", fgColor="224466")
            worksheet["A1"].alignment = Alignment(horizontal="center")
            worksheet["A2"] = "001"
            worksheet["B2"] = 1234
            worksheet["B2"].number_format = "0000"
            worksheet["B2"].font = Font(italic=True, color="112233")
            worksheet["C2"] = "HIDDEN-COL"
            worksheet["A3"] = "HIDDEN-ROW"
            worksheet["B3"] = 99
            worksheet["A4"] = "尾行"
            worksheet["B4"] = 42
            worksheet["C4"] = "HIDDEN-COL"
            worksheet.row_dimensions[2].height = 26
            worksheet.row_dimensions[3].hidden = True
            worksheet.row_dimensions[4].height = 31
            worksheet.column_dimensions["A"].width = 20
            worksheet.column_dimensions["B"].width = 17
            worksheet.column_dimensions["C"].hidden = True
            workbook.save(source_path)
            workbook.close()

            # This test validates the number-format fallback; never launch a real WPS COM server.
            with mock.patch.object(main, "_read_wps_display_text", return_value={}):
                paths = main._export_table_areas(
                    source_path, [{"sheet": "工单表", "range": "A1:C4"}], output_dir, lambda _msg: None
                )
            exported = load_workbook(paths[0], data_only=False)
            original = load_workbook(source_path, data_only=False)
            try:
                result = exported["工单表"]
                self.assertEqual({"A1:C1"}, {str(area) for area in result.merged_cells.ranges})
                self.assertEqual((4, 3), (result.max_row, result.max_column))
                self.assertEqual(("标题", "001", "HIDDEN-ROW", "尾行"),
                                 (result["A1"].value, result["A2"].value, result["A3"].value, result["A4"].value))
                self.assertEqual(("1234", "99", "42"), (result["B2"].value, result["B3"].value, result["B4"].value))
                self.assertEqual("HIDDEN-COL", result["C2"].value)
                self.assertEqual("s", result["B2"].data_type)
                self.assertEqual("0000", result["B2"].number_format)
                self.assertTrue(result["B2"].font.italic)
                self.assertEqual("112233", result["B2"].font.color.rgb[-6:])
                self.assertTrue(result["A1"].font.bold)
                self.assertEqual("224466", result["A1"].fill.fgColor.rgb[-6:])
                self.assertEqual("center", result["A1"].alignment.horizontal)
                self.assertEqual(26, result.row_dimensions[2].height)
                self.assertEqual(31, result.row_dimensions[4].height)
                self.assertEqual(20, result.column_dimensions["A"].width)
                self.assertEqual(17, result.column_dimensions["B"].width)
                self.assertTrue(result.row_dimensions[3].hidden)
                self.assertTrue(result.column_dimensions["C"].hidden)
                self.assertEqual("112233", result["B2"].font.color.rgb[-6:])
                self.assertTrue(original["工单表"].row_dimensions[3].hidden)
                self.assertTrue(original["工单表"].column_dimensions["C"].hidden)
            finally:
                exported.close()
                original.close()

    def test_region_xlsx_text_values_keep_number_format_display_precision(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "formats.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "显示格式"
            worksheet["A1"] = 1234
            worksheet["A1"].number_format = "#,##0"
            worksheet["A1"].font = Font(color="123456")
            worksheet["A2"] = 0.125
            worksheet["A2"].number_format = "0.00%"
            worksheet["A2"].font = Font(color="AA2244")
            worksheet["A3"] = 12.34567
            worksheet["A3"].number_format = "0.000"
            worksheet["A4"] = 42
            worksheet["A4"].number_format = "000000"
            worksheet["A5"] = __import__("datetime").datetime(2025, 4, 3)
            worksheet["A5"].number_format = "yyyy/mm/dd"
            worksheet["A6"] = True
            worksheet["A7"] = 0
            worksheet["A7"].number_format = "#,##0.00;(#,##0.00);-"
            worksheet.column_dimensions["A"].width = 1
            workbook.save(source_path)
            workbook.close()

            paths = main._export_table_areas(
                source_path, [{"sheet": "显示格式", "range": "A1:A7"}], output_dir, lambda _msg: None
            )
            exported = load_workbook(paths[0], data_only=False)
            try:
                result = exported.active
                self.assertEqual(
                    ["1,234", "12.50%", "12.346", "000042", "2025/04/03", "TRUE", "-"],
                    [result.cell(row=row, column=1).value for row in range(1, 8)],
                )
                self.assertTrue(all(result.cell(row=row, column=1).data_type == "s" for row in range(1, 8)))
                self.assertTrue(all("#" not in result.cell(row=row, column=1).value for row in range(1, 8)))
                self.assertEqual("123456", result["A1"].font.color.rgb[-6:])
                self.assertEqual("AA2244", result["A2"].font.color.rgb[-6:])
                self.assertEqual("0.00%", result["A2"].number_format)
                self.assertEqual("000000", result["A4"].number_format)
                self.assertEqual(1, result.column_dimensions["A"].width)
            finally:
                exported.close()

    def test_wps_display_text_uses_text_and_ignores_hash_overflow(self):
        class FakeSheet:
            def Cells(self, row, col):
                text = "1.234,50" if (row, col) == (1, 1) else "####"
                return type("FakeCell", (), {"Text": text})()

        fake_app = mock.Mock()
        fake_workbook = mock.Mock()
        fake_workbook.Worksheets.return_value = FakeSheet()
        fake_app.Workbooks.Open.return_value = fake_workbook
        fake_pythoncom = mock.Mock()
        safe_copy_path = Path("/tmp/sanitized-read-copy.xlsx")
        with mock.patch.object(main.sys, "platform", "win32"), \
                mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")):
            display = main._read_wps_display_text(
                safe_copy_path, [("S", 1, 1), ("S", 1, 2)], lambda _message: None
            )
        self.assertEqual({("S", 1, 1): "1.234,50"}, display)
        fake_app.Workbooks.Open.assert_called_once_with(
            str(safe_copy_path), UpdateLinks=0, ReadOnly=True, AddToMru=False
        )
        fake_workbook.Close.assert_called_once_with(SaveChanges=False)
        fake_app.Quit.assert_called_once_with()
        fake_pythoncom.CoUninitialize.assert_called_once_with()

    def test_region_xlsx_merge_closure_includes_the_complete_merged_title(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "统计"
            worksheet.merge_cells("AH1:CL1")
            worksheet["AH1"] = "标题"
            worksheet["BO44"] = 125
            workbook.save(source_path)
            workbook.close()

            paths = main._export_table_areas(
                source_path, [{"sheet": "统计", "range": "AH1:BO44"}], output_dir, lambda _msg: None
            )
            self.assertEqual("统计.xlsx", paths[0].name)
            exported = load_workbook(paths[0], data_only=False)
            try:
                result = exported["统计"]
                right_edge = column_index_from_string("CL") - column_index_from_string("AH") + 1
                merged = f"A1:{get_column_letter(right_edge)}1"
                self.assertIn(merged, {str(area) for area in result.merged_cells.ranges})
                value_col = column_index_from_string("BO") - column_index_from_string("AH") + 1
                self.assertEqual("125", result.cell(44, value_col).value)
                self.assertEqual("s", result.cell(44, value_col).data_type)
            finally:
                exported.close()

    def test_region_xlsx_clips_whole_sheet_shortcut_and_cleans_partial_outputs_on_error(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "整表"
            worksheet["A1"] = "head"
            worksheet["B2"] = 5
            worksheet.row_dimensions[3].hidden = True
            workbook.save(source_path)
            workbook.close()

            paths = main._export_table_areas(
                source_path, [{"sheet": "整表", "range": "A1"}], output_dir, lambda _msg: None
            )
            exported = load_workbook(paths[0], data_only=False)
            try:
                self.assertEqual((2, 2), (exported.active.max_row, exported.active.max_column))
                self.assertEqual("5", exported.active["B2"].value)
            finally:
                exported.close()

            for path in paths:
                path.unlink()
            with self.assertRaisesRegex(main.ProcessingError, "工作表“缺失”在已保存工作簿中不存在"):
                main._export_table_areas(
                    source_path,
                    [{"sheet": "整表", "range": "A1:B2"}, {"sheet": "缺失", "range": "A1"}],
                    output_dir, lambda _msg: None,
                )
            self.assertEqual([], list(output_dir.iterdir()))

    def test_legacy_export_helper_preserves_hidden_row_when_called_directly(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            output_dir = root / "exports"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "空白"
            worksheet["A1"] = "hidden"
            worksheet.row_dimensions[1].hidden = True
            workbook.save(source_path)
            workbook.close()
            paths = main._export_table_areas(
                source_path, [{"sheet": "空白", "range": "A1"}], output_dir, lambda _msg: None
            )
            exported = load_workbook(paths[0], data_only=False)
            try:
                self.assertEqual("hidden", exported.active["A1"].value)
                self.assertTrue(exported.active.row_dimensions[1].hidden)
            finally:
                exported.close()
            for path in paths:
                path.unlink()

    def test_text_area_reader_emits_tab_columns_newline_rows_and_blanks(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "text-area.xlsx"
            make_test_workbook(path, [("文本", {1: {1: "name", 2: "count"}, 2: {1: "row-1", 2: 5}, 3: {1: "row-2"}})])
            previews = main._read_text_output_areas(path, [{"sheet": "文本", "range": "A1:B3"}])
            self.assertEqual("name\tcount\nrow-1\t5\nrow-2\t", previews[0]["text"])

    def test_whole_sheet_text_is_resolved_to_used_range_and_chunked(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "used-range.xlsx"
            sheet_data = {1: {1: "header", 2: "value"}}
            sheet_data.update({row: {1: f"r{row}", 2: row} for row in range(2, 506)})
            make_test_workbook(path, [("整表", sheet_data)])
            preview = main._read_text_output_areas(path, [{"sheet": "整表", "range": "A1"}])[0]
            self.assertEqual("A1:B505", preview["resolved_range"])
            self.assertEqual(2, len(preview["chunks"]))
            self.assertEqual((1, 500), (preview["chunks"][0]["start_row"], preview["chunks"][0]["end_row"]))
            self.assertEqual((501, 505), (preview["chunks"][1]["start_row"], preview["chunks"][1]["end_row"]))
            self.assertTrue(preview["chunks"][1]["text"].startswith("r501\t501"))

    def test_large_text_output_and_processing_regions_are_rejected(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["text_outputs"] = [{"sheet": "工作表", "range": "A1:DV1000"}]
        with self.assertRaisesRegex(main.ProcessingError, "文本预览/复制上限"):
            main.validate_static_settings(settings)
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["RegionFilter"] = {"enabled": True, "rules": [{"sheet": "工作表", "range": "A1:XFD1000", "filter_column": "A", "operator": "blank"}]}
        with self.assertRaisesRegex(main.ProcessingError, "处理上限"):
            main.validate_static_settings(settings)
        settings["RegionFilter"]["rules"][0]["range"] = "A1"
        main.validate_static_settings(settings)

    def test_text_split_refuses_to_overwrite_nonempty_destination(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "split-conflict.xlsx"
            make_test_workbook(path, [("数据", {1: {1: "x|y", 2: "KEEP-ME"}})])
            events = []
            fake_workbook = FakeComWorkbook(path, events)
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = []
            settings["TextToColumns"] = {"enabled": True, "rules": [{
                "sheet": "数据", "range": "A1:A1", "delimiter": "|", "output_start_col": "B", "column_formats": ["text", "text"]}]}
            with self.assertRaisesRegex(main.ProcessingError, "不会静默覆盖"):
                main._apply_region_processing(fake_workbook, settings, [], lambda _msg: None)
            self.assertEqual("KEEP-ME", fake_workbook.workbook["数据"]["B1"].value)
            fake_workbook.workbook.close()

    def test_preview_copy_writes_exact_tabs_and_newlines_to_clipboard(self):
        text = "a\t1\nb\t2"
        clipboard = {"text": None}
        root = mock.Mock()
        root.clipboard_clear.side_effect = lambda: clipboard.update(text="")
        root.clipboard_append.side_effect = lambda value: clipboard.update(text=value)
        stub_app = type("ClipboardApp", (), {})()
        stub_app.root = root
        stub_app._log = lambda _message: None
        stub_app.messagebox = mock.Mock()
        main.App._copy_preview_to_clipboard(stub_app, text)
        self.assertEqual(text, clipboard["text"])
        root.clipboard_clear.assert_called_once()
        root.clipboard_append.assert_called_once_with(text)

    def test_formula_results_are_not_mapped_over_copied_columns(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            source_path = root / "source_dg.xlsx"
            make_test_workbook(source_path, [("数据", {
                2: {111: "COPIED-1"},
                3: {111: "COPIED-2"},
            })])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["copy_start_col"] = "DG"
            settings["copy_end_col"] = "DG"
            settings["template_path"] = str(fake / "template.xlsx")
            settings["jobs"][0]["sources"] = [
                {"path": str(source_path), "source_sheet": "数据", "copy_start_row": 2, "copy_end_row": 3}
            ]
            main.validate_static_settings(settings)
            records = main._read_source_records(source_path, "数据", 2, 111, 111, lambda _msg: None, end_row=3)
            staging = root / "mapping-overwrite.xlsx"
            formula_jobs = main._create_staging_workbook(
                Path(settings["template_path"]), settings, [records], staging, lambda _msg: None
            )

            staged = load_workbook(staging, data_only=False)
            try:
                self.assertEqual("COPIED-1", staged["1当天派单"]["DG3"].value)
                self.assertEqual("COPIED-2", staged["1当天派单"]["DG4"].value)
            finally:
                staged.close()

            events = []
            fake_app = FakeWpsApplication(events, {3: "CALCULATED-1", 4: "CALCULATED-2"})
            fake_pythoncom = FakePythonCom(events)
            with mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")):
                main._calculate_in_wps(staging, formula_jobs, lambda _msg: None)

            final = load_workbook(staging, data_only=False)
            try:
                self.assertEqual("COPIED-1", final["1当天派单"]["DG3"].value)
                self.assertEqual("COPIED-2", final["1当天派单"]["DG4"].value)
            finally:
                final.close()
            self.assertIn("calculate", events)
            self.assertFalse(any(isinstance(event, tuple) and event[0] in ("read", "write") and "EJ" in event[1]
                                 for event in events))

    def test_post_mapping_column_alias_only_writes_data_rows_and_preserves_template_formula(self):
        with tempfile.TemporaryDirectory() as td:
            staging = Path(td) / "post-column.xlsx"
            make_test_workbook(staging, [("1当天派单", {
                2: {111: "=A2*2", 140: "SEED"},
                3: {111: "OLD-3", 140: "OLD-SOURCE-3"},
                4: {111: "OLD-4", 140: "OLD-SOURCE-4"},
            })])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["jobs"][0].update({"template_sheet": "1当天派单", "start_row": 3, "formula_range": "EB:EN"})
            rule = {"source_sheet": "1当天派单", "source_range": "EJ:EJ",
                    "target_sheet": "1当天派单", "target_range": "DG:DG"}
            settings["post_mappings"] = [rule]
            formula_jobs = [("1当天派单", 3, 4, [])]
            events = []
            fake_app = FakeWpsApplication(events, {3: "CALC-3", 4: "CALC-4"})
            fake_pythoncom = FakePythonCom(events)
            with mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")):
                main._calculate_in_wps(staging, formula_jobs, lambda _message: None, [rule], settings)

            output = load_workbook(staging, data_only=False)
            try:
                self.assertEqual("=A2*2", output["1当天派单"]["DG2"].value)
                self.assertEqual("CALC-3", output["1当天派单"]["DG3"].value)
                self.assertEqual("CALC-4", output["1当天派单"]["DG4"].value)
            finally:
                output.close()
            self.assertIn(("read", "EJ3:EJ4"), events)
            self.assertIn(("write", "DG3:DG4"), events)
            self.assertNotIn(("write", "DG1:DG1048576"), events)
            self.assertNotIn(("write", "DG2"), events)

    def test_wps_order_is_recalc_post_mapping_three_independent_functions_recalc_save(self):
        with tempfile.TemporaryDirectory() as td:
            staging = Path(td) / "ordered.xlsx"
            make_test_workbook(staging, [
                ("1当天派单", {3: {1: "11|x"}, 4: {1: "11|x"}, 5: {1: "12|y"}}),
                ("2当天归档", {1: {1: "POST-A", 2: "POST-B"}}),
                ("3装移全业务在途", {1: {4: None, 5: None}}),
            ])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["jobs"][0].update({"template_sheet": "1当天派单", "start_row": 3, "formula_range": "EB:EN"})
            settings["TextToColumns"] = {"enabled": True, "rules": [{
                "sheet": "1当天派单", "range": "A3:A5", "delimiter": "|", "output_start_col": "H", "column_formats": ["number", "text"]}]}
            settings["Deduplicate"] = {"enabled": True, "rules": [{
                "sheet": "1当天派单", "range": "H3:I5", "has_header": False, "key_mode": "row", "key_columns": []}]}
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "sheet": "1当天派单", "range": "H3:H5", "mode": "text"}]}
            post_mappings = [{"source_sheet": "2当天归档", "source_range": "A1:B1",
                              "target_sheet": "3装移全业务在途", "target_range": "D1:E1"}]
            formula_jobs = [("1当天派单", 3, 5, [{"result_col": "EJ", "target_col": "DG", "write_as": "text"}])]
            events = []
            fake_app = FakeWpsApplication(events, {3: "CALC-3", 4: "CALC-4", 5: "CALC-5"})
            fake_pythoncom = FakePythonCom(events)
            with mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")):
                main._calculate_in_wps(staging, formula_jobs, lambda _msg: None, post_mappings, settings)

            output = load_workbook(staging, data_only=False)
            try:
                self.assertEqual(("POST-A", "POST-B"), (output["3装移全业务在途"]["D1"].value, output["3装移全业务在途"]["E1"].value))
                self.assertEqual(("CALC-3", "CALC-4", "CALC-5"), tuple(output["1当天派单"][f"DG{row}"].value for row in range(3, 6)))
                self.assertEqual("@", output["1当天派单"]["DG3"].number_format)
                self.assertEqual("11", output["1当天派单"]["H3"].value)
                self.assertEqual("12", output["1当天派单"]["H4"].value)
                self.assertIsNone(output["1当天派单"]["H5"].value)
                self.assertEqual("@", output["1当天派单"]["H3"].number_format)
            finally:
                output.close()

            first_recalc = events.index("calculate")
            result_read = events.index(("read", "EJ3:EJ5"))
            result_write = events.index(("write", "DG3:DG5"))
            post_write = events.index(("write", "D1:E1"))
            split_write = events.index(("write", "H3:I5"))
            convert_write = events.index(("write", "H3:H5"))
            second_recalc = events.index("calculate", first_recalc + 1)
            save = events.index("save")
            self.assertLess(first_recalc, result_read)
            self.assertLess(result_read, result_write)
            self.assertLess(result_write, post_write)
            self.assertLess(first_recalc, post_write)
            self.assertLess(post_write, split_write)
            self.assertLess(split_write, convert_write)
            self.assertLess(convert_write, second_recalc)
            self.assertLess(second_recalc, save)

    def test_independent_region_modules_execute_in_custom_workflow_order(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "custom-region-order.xlsx"
            make_test_workbook(path, [("数据", {
                1: {1: "x|y", 5: "dup", 9: "42"},
                2: {5: "dup"},
                3: {5: "keep"},
            })])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["TextToColumns"] = {"enabled": True, "rules": [{
                "sheet": "数据", "range": "A1:A1", "delimiter": "|", "output_start_col": "C",
                "column_formats": ["text", "text"]}]}
            settings["Deduplicate"] = {"enabled": True, "rules": [{
                "sheet": "数据", "range": "E1:E3", "has_header": False, "key_mode": "row", "key_columns": []}]}
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "sheet": "数据", "range": "I1:I1", "mode": "number"}]}
            order = list(settings["workflow"])
            region_start = order.index("TextToColumns")
            for module_id in ("TextToColumns", "Deduplicate", "FormatConvert"):
                order.remove(module_id)
            order[region_start:region_start] = ["FormatConvert", "Deduplicate", "TextToColumns"]
            settings["workflow"] = order
            main.validate_static_settings(settings)

            events = []
            fake_workbook = FakeComWorkbook(path, events)
            logs = []
            self.assertTrue(main._apply_region_processing(fake_workbook, settings, [], logs.append))
            output = fake_workbook.workbook["数据"]
            self.assertEqual(("x", "y"), (output["C1"].value, output["D1"].value))
            self.assertEqual(("dup", "keep", None), (output["E1"].value, output["E2"].value, output["E3"].value))
            self.assertEqual(42, output["I1"].value)
            action_logs = [line for line in logs if line.startswith("已执行")]
            self.assertIn("已执行格式转换", action_logs[0])
            self.assertIn("已执行去重", action_logs[1])
            self.assertIn("已执行文本分列", action_logs[2])
            fake_workbook.workbook.close()

    def test_text_output_runs_after_saved_workbook(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"copy_end_col": "C", "template_path": str(fake / "template.xlsx"),
                             "output_dir": str(fake / "output"), "output_filename": "terminal-outputs.xlsx",
                             "text_outputs": [{"sheet": "2当天归档", "range": "A1:B2"}]})
            for idx, job in enumerate(settings["jobs"], 1):
                job["sources"] = [{"path": str(fake / f"{idx}.xlsx"), "source_sheet": "示例源表",
                                   "copy_start_row": 2, "copy_end_row": 3}]
            events = []
            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _settings):
                events.append("workbook-save")
                workbook = load_workbook(path, data_only=False); workbook.save(path); workbook.close()
            def fake_text(path, _outputs, _log=None):
                self.assertTrue(path.is_file())
                events.append("text-read")
                return [{"sheet": "2当天归档", "range": "A1:B2", "text": "preview"}]
            previews = []
            with mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate), \
                 mock.patch.object(main, "_read_text_output_areas", side_effect=fake_text):
                result = main.process_workbooks(settings, lambda _msg: None, previews.extend)
            self.assertTrue(result.is_file())
            self.assertEqual(["workbook-save", "text-read"], events)
            self.assertEqual("preview", previews[0]["text"])
            self.assertEqual([result], list(Path(settings["output_dir"]).iterdir()))

    def test_filter_can_run_on_target_workbook_range(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "target-filter.xlsx"
            make_test_workbook(path, [("数据", {1: {1: "ID", 2: "状态"}, 2: {1: 1, 2: "ok"},
                                      3: {1: 2, 2: "skip"}, 4: {1: 3, 2: "ok"}})])
            fake_workbook = FakeComWorkbook(path, [])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = []
            settings["RegionFilter"] = {"enabled": True, "rules": [{"scope": "target", "sheet": "数据",
                "range": "A1:B4", "has_header": True, "filter_column": "B", "operator": "equals", "criteria": "ok"}]}
            self.assertTrue(main._apply_region_processing(fake_workbook, settings, [], lambda _msg: None))
            sheet = fake_workbook.workbook["数据"]
            self.assertEqual(["ID", "状态", 1, "ok", 3, "ok"],
                             [sheet.cell(row=row, column=col).value for row, col in ((1,1),(1,2),(2,1),(2,2),(3,1),(3,2))])
            self.assertIsNone(sheet["A4"].value)
            fake_workbook.workbook.close()

    def test_text_output_failure_keeps_existing_no_final_publish_policy(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"copy_end_col": "C", "template_path": str(fake / "template.xlsx"),
                             "output_dir": str(fake / "output"), "output_filename": "text-failure.xlsx",
                             "text_outputs": [{"sheet": "1当天派单", "range": "A1:B2"}]})
            for idx, job in enumerate(settings["jobs"], 1):
                job["sources"] = [{"path": str(fake / f"{idx}.xlsx"), "source_sheet": "示例源表",
                                   "copy_start_row": 2, "copy_end_row": 3}]

            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _region_processing):
                workbook = load_workbook(path, data_only=False)
                workbook.save(path)
                workbook.close()

            with mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate), \
                 mock.patch.object(main, "_read_text_output_areas",
                                   side_effect=main.ProcessingError("模拟文本输出失败")):
                with self.assertRaisesRegex(main.ProcessingError, "模拟文本输出失败"):
                    main.process_workbooks(settings, lambda _msg: None)
            self.assertEqual([], list(Path(settings["output_dir"]).iterdir()))

    def test_workbook_publish_never_overwrites_existing_destination(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            staging = root / "staging.xlsx"
            final_path = root / "result.xlsx"
            staging.write_bytes(b"NEW-WORKBOOK")
            final_path.write_bytes(b"EXISTING-USER-FILE")

            with self.assertRaises(main.ProcessingError) as caught:
                main._publish_staging_workbook(staging, final_path)

            self.assertIn("为避免覆盖", str(caught.exception))
            self.assertIn(str(final_path), str(caught.exception))
            self.assertEqual(b"EXISTING-USER-FILE", final_path.read_bytes())
            self.assertEqual(b"NEW-WORKBOOK", staging.read_bytes())

    def test_template_only_staging_preserves_existing_cells_without_sources(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            template_path = root / "template.xlsx"
            make_test_workbook(template_path, [("模板", {1: {1: "标题"}, 5: {1: "已有模板数据"}})])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["jobs"][0].update({"template_sheet": "模板", "sources": [], "header_rows": 0,
                                         "start_row": 3, "formula_range": "EB:EN"})
            settings["copy_start_col"] = ""
            settings["copy_end_col"] = ""
            main.validate_static_settings(settings)
            staging = root / "template-only.xlsx"
            formula_jobs = main._create_staging_workbook(template_path, settings, [[]], staging, lambda _msg: None)
            self.assertEqual([], formula_jobs)
            output = load_workbook(staging, data_only=False)
            try:
                self.assertEqual("标题", output["模板"]["A1"].value)
                self.assertEqual("已有模板数据", output["模板"]["A5"].value)
            finally:
                output.close()

            settings.update({"template_path": str(template_path), "output_dir": str(root),
                             "output_filename": "template-only-final.xlsx", "copy_end_col": ""})
            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _region_processing):
                workbook = load_workbook(path, data_only=False)
                workbook.save(path)
                workbook.close()
            with mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate):
                final_path = main.process_workbooks(settings, lambda _msg: None)
            self.assertTrue(final_path.is_file())
            final = load_workbook(final_path, data_only=False)
            try:
                self.assertEqual("已有模板数据", final["模板"]["A5"].value)
            finally:
                final.close()

    def test_source_rules_run_on_wps_copy_and_optional_retained_copy_preserves_original(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            make_test_workbook(source_path, [("Data", {1: {1: "001"}, 2: {1: "2"}})])
            original_bytes = source_path.read_bytes()
            output_dir = root / "output"
            output_dir.mkdir()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["workflow"] = ["FormatConvert"]
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "A1", "mode": "number"
            }]}
            events = []
            fake_app = FakeWpsApplication(events, {})
            fake_pythoncom = FakePythonCom(events)
            retained_paths = []
            used_rules = set()
            with mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")):
                records = main._read_source_records_with_rules(
                    source_path, "Data", 1, 1, 1, lambda _msg: None, 2, settings,
                    used_rules, True, output_dir, output_dir / "result.xlsx", retained_paths.append
                )
            self.assertEqual([[1], [2]], records)
            self.assertEqual({("FormatConvert", 0)}, used_rules)
            self.assertEqual(original_bytes, source_path.read_bytes())
            self.assertEqual(1, len(retained_paths))
            retained = load_workbook(retained_paths[0], data_only=True)
            try:
                self.assertEqual(1, retained["Data"]["A1"].value)
                self.assertEqual(2, retained["Data"]["A2"].value)
            finally:
                retained.close()
            self.assertIn("save", events)

    def test_optional_merged_source_save_uses_target_default_or_custom_name(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            make_test_workbook(source_path, [("Data", {1: {1: "ID"}, 2: {1: "123"}})])
            original_bytes = source_path.read_bytes()
            template_path = root / "template.xlsx"
            template = Workbook()
            template.active.title = "Target Default"
            template.active["B1"] = "=ROW()"
            template_custom = template.create_sheet("TargetCustom")
            template_custom["B1"] = "=ROW()"
            template.save(template_path)
            template.close()
            output_dir = root / "output"
            output_dir.mkdir()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"template_path": str(template_path), "output_dir": str(output_dir),
                             "output_filename": "template-result.xlsx", "copy_start_col": "A",
                             "copy_end_col": "A", "retain_merged_sources": True})
            source = {"path": str(source_path), "source_sheet": "Data", "copy_start_row": 1,
                      "copy_end_row": None, "keep_first_row": True}
            settings["jobs"] = [
                {"template_sheet": "Target Default", "merged_source_filename": "", "header_rows": 0,
                 "start_row": 2, "formula_range": "B:B", "sources": [dict(source)], "mappings": []},
                {"template_sheet": "TargetCustom", "merged_source_filename": "custom-source.xlsx",
                 "header_rows": 0, "start_row": 2, "formula_range": "B:B",
                 "sources": [dict(source)], "mappings": []},
            ]
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "A2", "mode": "number",
            }]}
            merged_paths = []
            real_merge = main._merge_source_workbooks
            events = []

            def fake_connect(_log):
                return FakeWpsApplication(events, {}), FakePythonCom(events), "fake.WPS"

            def capture_merge(*args, **kwargs):
                merged = real_merge(*args, **kwargs)
                merged_paths.append(merged["path"])
                return merged

            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _region_processing):
                self.assertEqual(2, len(merged_paths))
                self.assertTrue(all(merged_path.is_file() for merged_path in merged_paths))
                workbook = load_workbook(path, data_only=False)
                try:
                    self.assertEqual("ID", workbook["Target Default"]["A2"].value)
                    self.assertEqual(123, workbook["Target Default"]["A3"].value)
                    self.assertEqual("ID", workbook["TargetCustom"]["A2"].value)
                    self.assertEqual(123, workbook["TargetCustom"]["A3"].value)
                    workbook.save(path)
                finally:
                    workbook.close()

            with mock.patch.object(main, "_merge_source_workbooks", side_effect=capture_merge), \
                    mock.patch.object(main, "_connect_wps", side_effect=fake_connect), \
                    mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate):
                result = main.process_workbooks(settings, lambda _msg: None)

            self.assertTrue(result.is_file())
            self.assertEqual(["Target Default.xlsx", "custom-source.xlsx"],
                             sorted(path.name for path in output_dir.glob("*.xlsx") if path != result))
            self.assertTrue(all(not merged_path.exists() for merged_path in merged_paths))
            self.assertEqual(original_bytes, source_path.read_bytes())
            for filename in ("Target Default.xlsx", "custom-source.xlsx"):
                saved = load_workbook(output_dir / filename, read_only=True, data_only=True)
                try:
                    self.assertEqual([("ID",), (123,)], list(saved["合并源数据"].values))
                finally:
                    saved.close()

    def test_processes_source_copy_before_editing_template(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            make_test_workbook(source_path, [("Data", {1: {1: "001"}, 2: {1: "2"}})])
            original_bytes = source_path.read_bytes()
            template_path = root / "template.xlsx"
            template = Workbook()
            template.active.title = "Target"
            template.active["EB1"] = "=ROW()"
            template.save(template_path)
            template.close()
            output_dir = root / "output"
            output_dir.mkdir()

            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings.update({"template_path": str(template_path), "output_dir": str(output_dir),
                             "output_filename": "source-first.xlsx", "copy_start_col": "A",
                             "copy_end_col": "A"})
            settings["jobs"] = [{
                "template_sheet": "Target", "header_rows": 0, "start_row": 2,
                "formula_range": "EB:EB", "sources": [{
                    "path": str(source_path), "source_sheet": "Data",
                    "copy_start_row": 1, "copy_end_row": 2,
                }],
            }]
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "A1", "mode": "number",
            }]}

            events = []
            fake_app = FakeWpsApplication(events, {})
            fake_pythoncom = FakePythonCom(events)
            merged_paths = []
            source_read_paths = []
            real_merge = main._merge_source_workbooks
            real_read_source = main._read_source_records

            def capture_merge(*args, **kwargs):
                merged = real_merge(*args, **kwargs)
                merged_paths.append(merged["path"])
                return merged

            def fake_calculate(path, _formula_jobs, _log, _post_mappings, _settings):
                self.assertEqual(1, len(merged_paths))
                self.assertTrue(merged_paths[0].is_file())
                self.assertTrue((output_dir / "Target.xlsx").is_file())
                self.assertEqual([output_dir / "Target.xlsx"], source_read_paths)
                events.append("template-processing")
                workbook = load_workbook(path, data_only=False)
                workbook.save(path)
                workbook.close()

            def capture_source_read(path, *args, **kwargs):
                source_path_used = Path(path)
                self.assertTrue(source_path_used.is_file())
                source_read_paths.append(source_path_used)
                return real_read_source(path, *args, **kwargs)

            with mock.patch.object(main, "_merge_source_workbooks", side_effect=capture_merge), \
                    mock.patch.object(main, "_connect_wps", return_value=(fake_app, fake_pythoncom, "fake.WPS")), \
                    mock.patch.object(main, "_read_source_records", side_effect=capture_source_read), \
                    mock.patch.object(main, "_calculate_in_wps", side_effect=fake_calculate):
                final_path = main.process_workbooks(settings, lambda _msg: None)

            self.assertLess(events.index("save"), events.index("template-processing"))
            self.assertTrue(merged_paths)
            self.assertTrue(all(not merged_path.exists() for merged_path in merged_paths))
            self.assertTrue((output_dir / "Target.xlsx").is_file())
            self.assertEqual([output_dir / "Target.xlsx"], source_read_paths)
            self.assertEqual(original_bytes, source_path.read_bytes())
            result = load_workbook(final_path, data_only=False)
            try:
                self.assertEqual(1, result["Target"]["A2"].value)
                self.assertEqual(2, result["Target"]["A3"].value)
            finally:
                result.close()

    def test_processed_source_copy_deletion_waits_for_success_and_is_opt_in(self):
        for template_fails in (False, True):
            with self.subTest(template_fails=template_fails), tempfile.TemporaryDirectory() as td:
                root = Path(td)
                source_path = root / "source.xlsx"
                make_test_workbook(source_path, [("Data", {1: {1: "ID"}, 2: {1: "001"}})])
                source_bytes = source_path.read_bytes()
                template_path = root / "template.xlsx"
                template = Workbook()
                template.active.title = "Target"
                template.active["EB1"] = "=ROW()"
                template.save(template_path)
                template.close()
                template_bytes = template_path.read_bytes()
                output_dir = root / "output"
                output_dir.mkdir()

                settings = deepcopy(main.DEFAULT_SETTINGS)
                settings.update({
                    "template_path": str(template_path), "output_dir": str(output_dir),
                    "output_filename": "result.xlsx", "copy_start_col": "A", "copy_end_col": "A",
                    "delete_processed_sources_after_success": True,
                })
                settings["jobs"] = [{
                    "template_sheet": "Target", "header_rows": 0, "start_row": 2,
                    "formula_range": "EB:EB", "sources": [{
                        "path": str(source_path), "source_sheet": "Data",
                        "copy_start_row": 1, "copy_end_row": 2,
                    }],
                }]
                settings["FormatConvert"] = {"enabled": True, "rules": [{
                    "scope": "source", "sheet": "Data", "range": "A2", "mode": "number",
                }]}
                events = []
                fake_app = FakeWpsApplication(events, {})
                fake_pythoncom = FakePythonCom(events)
                processed_copy = output_dir / "Target.xlsx"

                def calculate_template(path, _formula_jobs, _log, _post_mappings, _settings):
                    self.assertTrue(processed_copy.is_file())
                    if template_fails:
                        raise main.ProcessingError("模拟模板处理失败")
                    workbook = load_workbook(path, data_only=False)
                    workbook.save(path)
                    workbook.close()

                with mock.patch.object(main, "_connect_wps",
                                       return_value=(fake_app, fake_pythoncom, "fake.WPS")), \
                        mock.patch.object(main, "_calculate_in_wps", side_effect=calculate_template):
                    if template_fails:
                        with self.assertRaisesRegex(main.ProcessingError, "模拟模板处理失败"):
                            main.process_workbooks(settings, lambda _msg: None)
                    else:
                        final_path = main.process_workbooks(settings, lambda _msg: None)
                        self.assertTrue(final_path.is_file())

                self.assertEqual(source_bytes, source_path.read_bytes())
                self.assertEqual(template_bytes, template_path.read_bytes())
                self.assertEqual(template_fails, processed_copy.is_file())

    def test_formula_templates_use_previous_row_and_fill_only_pasted_rows(self):
        for start_row in (3, 4):
            with self.subTest(start_row=start_row), tempfile.TemporaryDirectory() as td:
                root = Path(td)
                template_path = root / "template.xlsx"
                template = Workbook()
                ws = template.active
                ws.title = "模板"
                template_row = start_row - 1
                ws.cell(template_row, 1).value = f"KEEP-TEMPLATE-ROW-{template_row}"
                original_formulas = {}
                for col in (column_index_from_string("EB"), column_index_from_string("EC")):
                    letter = get_column_letter(col)
                    formula = f'=IF($A{template_row}="", "", $A{template_row}&"-{letter}")'
                    ws.cell(template_row, col).value = formula
                    original_formulas[letter] = formula
                template.save(template_path)
                template.close()

                settings = deepcopy(main.DEFAULT_SETTINGS)
                settings["jobs"] = [deepcopy(settings["jobs"][0])]
                settings["jobs"][0].update({
                    "template_sheet": "模板", "header_rows": 0, "start_row": start_row,
                    "formula_range": "EB:EC",
                })
                records = [[f"DATA-{start_row}-1"], [f"DATA-{start_row}-2"], [f"DATA-{start_row}-3"]]
                staging_path = root / "staging.xlsx"
                formula_jobs = main._create_staging_workbook(
                    template_path, settings, [records], staging_path, lambda _msg: None
                )
                self.assertEqual([("模板", start_row, start_row + len(records) - 1, [])], formula_jobs)

                result = load_workbook(staging_path, data_only=False)
                try:
                    result_ws = result["模板"]
                    self.assertEqual(f"KEEP-TEMPLATE-ROW-{template_row}", result_ws.cell(template_row, 1).value)
                    for letter, seed_formula in original_formulas.items():
                        self.assertEqual(seed_formula, result_ws[f"{letter}{template_row}"].value)
                        for target_row in range(start_row, start_row + len(records)):
                            expected = main.Translator(seed_formula, origin=f"{letter}{template_row}").translate_formula(
                                f"{letter}{target_row}"
                            )
                            self.assertEqual(expected, result_ws[f"{letter}{target_row}"].value)
                    self.assertEqual(records[0][0], result_ws.cell(start_row, 1).value)
                    self.assertEqual(records[-1][0], result_ws.cell(start_row + len(records) - 1, 1).value)
                finally:
                    result.close()

    def test_formula_template_missing_columns_are_reported_together(self):
        workbook = Workbook()
        ws = workbook.active
        ws.title = "模板"
        ws["EB2"] = "=A2"
        ws["ED2"] = "=C2"
        with self.assertRaisesRegex(main.ProcessingError, r"第 2 行.*EC、EE"):
            main._get_formula_templates(
                workbook,
                {"template_sheet": "模板", "start_row": 3, "formula_range": "EB:EE"},
                1,
                3,
            )
        workbook.close()

    def test_source_bounds_multi_source_order_and_formula_seed_fill(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            template_path = fake / "template.xlsx"
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["copy_end_col"] = "C"
            settings["template_path"] = str(template_path)
            settings["output_dir"] = str(fake / "output")
            settings["output_filename"] = "not-published.xlsx"
            settings["jobs"][0].update({"header_rows": 2, "start_row": 5})
            settings["jobs"][1].update({"header_rows": 2, "start_row": 6})
            settings["jobs"][2].update({"header_rows": 3, "start_row": 7})
            settings["jobs"][0]["sources"] = [
                {"path": str(fake / "1.xlsx"), "source_sheet": "示例源表", "copy_start_row": 2, "copy_end_row": 3},
                {"path": str(fake / "1_extra.xlsx"), "source_sheet": "示例源表", "copy_start_row": 4, "copy_end_row": 5},
            ]
            settings["jobs"][1]["sources"] = [
                {"path": str(fake / "2.xlsx"), "source_sheet": "示例源表", "copy_start_row": 3, "copy_end_row": 4},
            ]
            settings["jobs"][2]["sources"] = [
                {"path": str(fake / "3.xlsx"), "source_sheet": "示例源表", "copy_start_row": 5, "copy_end_row": None},
            ]
            # Move formula seeds in a template copy to distinct configured destination rows.
            wb = load_workbook(template_path)
            for job in settings["jobs"]:
                ws = wb[job["template_sheet"]]
                first, last = main.parse_column_range(job["formula_range"])
                row = job["start_row"] - 1
                for col in range(first, last + 1):
                    ws.cell(row=row, column=col).value = f'=IF($A{row}="", "", $A{row}&"-演示")'
            wb.save(template_path)
            wb.close()

            all_records = []
            for job in settings["jobs"]:
                records = []
                for source in job["sources"]:
                    start, end = source["copy_start_row"], source["copy_end_row"]
                    records.extend(main._read_source_records(Path(source["path"]), source["source_sheet"], start, 1, 3,
                                                             lambda _msg: None, end))
                all_records.append(records)
            ids = [record[0] for record in all_records[0]]
            self.assertEqual(["FAKE-1-01", "FAKE-1-02", "FAKE-1X-03", "FAKE-1X-04"], ids)
            self.assertEqual(["FAKE-2-02", "FAKE-2-03"], [r[0] for r in all_records[1]])
            self.assertEqual(["FAKE-3-04"], [r[0] for r in all_records[2]])

            staging = root / "prepared.xlsx"
            main._create_staging_workbook(template_path, settings, all_records, staging, lambda _msg: None)
            out = load_workbook(staging, data_only=False)
            try:
                for job, records in zip(settings["jobs"], all_records):
                    ws = out[job["template_sheet"]]
                    start = job["start_row"]
                    last = start + len(records) - 1
                    self.assertIn("演示标题", ws["A1"].value)
                    self.assertEqual(records[0][0], ws.cell(start, 1).value)
                    self.assertEqual(records[-1][0], ws.cell(last, 1).value)
                    first_formula_col, _ = main.parse_column_range(job["formula_range"])
                    template_row = start - 1
                    self.assertTrue(ws.cell(template_row, first_formula_col).value.startswith("=IF("))
                    self.assertTrue(ws.cell(start, first_formula_col).value.startswith("=IF("))
                    self.assertTrue(ws.cell(last, first_formula_col).value.startswith("=IF("))
                    self.assertIn(f"${get_column_letter(1)}{last}", ws.cell(last, first_formula_col).value)
                    self.assertIsNone(ws["A20"].value)
                    self.assertIsNone(ws["DG25"].value)
                    self.assertIsNotNone(ws["A20"].fill.fgColor.rgb)
            finally:
                out.close()

    def test_all_four_region_operations_can_process_source_file_without_modifying_it(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.xlsx"
            make_test_workbook(path, [("Data", {
                1: {1: "ID", 2: "Name", 3: "Value", 4: "Pair"},
                2: {1: 1, 2: "old", 3: 10, 4: "A|10"},
                3: {1: 2, 2: "dup", 3: 20, 4: "B|20"},
                4: {1: 2, 2: "dup", 3: 20, 4: "B|20"},
                5: {1: 3, 2: "keep", 3: 30, 4: "C|30"},
            })])
            original = path.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["workflow"] = ["TextToColumns", "RegionFilter", "Deduplicate", "FormatConvert"]
            settings["TextToColumns"] = {"enabled": True, "rules": [{"scope": "source", "sheet": "Data", "range": "D1:D5", "delimiter": "|", "output_start_col": "E", "column_formats": ["text", "number"]}]}
            settings["RegionFilter"] = {"enabled": True, "rules": [{"scope": "source", "sheet": "Data", "range": "A1:F5", "has_header": True, "filter_column": "C", "operator": "greater_equal", "criteria": "20"}]}
            settings["Deduplicate"] = {"enabled": True, "rules": [{"scope": "source", "sheet": "Data", "range": "A1:F5", "has_header": True, "key_mode": "columns", "key_columns": ["A"]}]}
            settings["FormatConvert"] = {"enabled": True, "rules": [{"scope": "source", "sheet": "Data", "range": "C2:C5", "mode": "number"}]}
            used = set()
            records = main._read_source_records(path, "Data", 1, 1, 6, lambda _msg: None,
                                               source_processing=settings, source_rules_used=used)
            self.assertEqual([
                ["ID", "Name", "Value", "Pair", "Pair", None],
                [2, "dup", 20, "B|20", "B", 20],
                [3, "keep", 30, "C|30", "C", 30],
            ], records)
            self.assertEqual({("TextToColumns", 0), ("RegionFilter", 0), ("Deduplicate", 0), ("FormatConvert", 0)}, used)
            self.assertEqual(original, path.read_bytes())

    def test_blank_source_sheet_name_selects_first_worksheet(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "two_sheets.xlsx"
            make_test_workbook(source_path, [
                ("首个工作表", {2: {1: "FIRST-SHEET"}}),
                ("第二个工作表", {2: {1: "SECOND-SHEET"}}),
            ])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"][0]["sources"][0].update({"path": str(source_path), "source_sheet": "   "})
            main.validate_static_settings(settings)
            messages = []
            records = main._read_source_records(source_path, "", 2, 1, 3, messages.append)
            self.assertEqual("FIRST-SHEET", records[0][0])
            self.assertTrue(any("默认使用第一个工作表“首个工作表”" in message for message in messages))

    def test_blank_row_boundary_checks_all_used_columns_and_next_source_appends(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            first_source = root / "first.xlsx"
            second_source = root / "second.xlsx"
            make_test_workbook(first_source, [("数据", {
                2: {1: "SRC1-1", 2: "A"},
                3: {4: "copy-range-empty-but-row-not-empty"},
                4: {1: "SRC1-2", 2: "B"},
                6: {1: "AFTER-ALL-COLUMN-BLANK-BOUNDARY"},
            })])
            make_test_workbook(second_source, [("数据", {
                2: {1: "SRC2-1"},
                3: {1: "SRC2-2"},
            })])

            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["template_path"] = str(fake / "template.xlsx")
            settings["output_dir"] = str(fake / "output")
            settings["output_filename"] = "append-boundary-result.xlsx"
            job = settings["jobs"][0]
            job.update({"header_rows": 2, "start_row": 5})
            job["sources"] = [
                {"path": str(first_source), "source_sheet": "数据", "copy_start_row": 2, "copy_end_row": None},
                {"path": str(second_source), "source_sheet": "数据", "copy_start_row": 2, "copy_end_row": None},
            ]

            template = load_workbook(settings["template_path"])
            ws = template[job["template_sheet"]]
            formula_start, formula_end = main.parse_column_range(job["formula_range"])
            for col in range(formula_start, formula_end + 1):
                ws.cell(row=4, column=col).value = '=IF($A4="", "", $A4&"-演示")'
            template.save(settings["template_path"])
            template.close()
            main.validate_static_settings(settings)

            records = []
            messages = []
            for source in job["sources"]:
                records.extend(main._read_source_records(Path(source["path"]), source["source_sheet"],
                                                         source["copy_start_row"], 1, 3, messages.append))
            self.assertEqual(["SRC1-1", "SRC1-2", "SRC2-1", "SRC2-2"], [row[0] for row in records])
            self.assertTrue(any("第 5 行遇到整行空白边界" in message for message in messages))

            staging = root / "prepared.xlsx"
            main._create_staging_workbook(Path(settings["template_path"]), settings, [records], staging, lambda _msg: None)
            output = load_workbook(staging, data_only=False)
            try:
                out_ws = output[job["template_sheet"]]
                self.assertEqual(["SRC1-1", "SRC1-2", "SRC2-1", "SRC2-2"],
                                 [out_ws.cell(row=row, column=1).value for row in range(5, 9)])
                self.assertIsNone(out_ws.cell(row=9, column=1).value)
                self.assertEqual('=IF($A4="", "", $A4&"-演示")', out_ws.cell(4, formula_start).value)
            finally:
                output.close()

    def test_empty_end_row_without_blank_boundary_uses_last_nonempty_row(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "no_boundary.xlsx"
            make_test_workbook(source_path, [("数据", {
                2: {1: "FIRST"},
                3: {4: "outside-copy-columns"},
                4: {1: "LAST"},
            })])
            messages = []
            records = main._read_source_records(source_path, "数据", 2, 1, 3, messages.append)
            self.assertEqual(["FIRST", "LAST"], [row[0] for row in records])
            self.assertTrue(any("最后非空行 4 未遇到整行空白边界" in message for message in messages))

    def test_explicit_end_row_includes_data_across_blank_rows_but_not_beyond_end(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "explicit_end.xlsx"
            make_test_workbook(source_path, [("数据", {
                2: {1: "IN-1"},
                4: {1: "IN-2"},
                6: {1: "AFTER-EXPLICIT-END"},
            })])
            records = main._read_source_records(source_path, "数据", 2, 1, 3, lambda _msg: None, end_row=4)
            self.assertEqual(["IN-1", "IN-2"], [row[0] for row in records])

    def test_all_blank_source_fails_without_creating_output(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            blank_source = root / "blank.xlsx"
            make_test_workbook(blank_source, [("空表", {})])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["copy_end_col"] = "C"
            settings["template_path"] = str(fake / "template.xlsx")
            settings["output_dir"] = str(fake / "output")
            settings["output_filename"] = "must-not-exist.xlsx"
            settings["jobs"][0]["sources"] = [
                {"path": str(blank_source), "source_sheet": "", "copy_start_row": 2, "copy_end_row": None}
            ]
            with self.assertRaisesRegex(main.ProcessingError, "没有可复制数据"):
                main.process_workbooks(settings, lambda _msg: None)
            self.assertEqual([], list((fake / "output").iterdir()))

    def test_failure_before_wps_save_leaves_no_final_or_staging_file(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["copy_end_col"] = "C"
            settings["template_path"] = str(fake / "template.xlsx")
            settings["output_dir"] = str(fake / "output")
            settings["output_filename"] = "failure.xlsx"
            for idx, job in enumerate(settings["jobs"], 1):
                job["sources"] = [{"path": str(fake / f"{idx}.xlsx"), "source_sheet": "示例源表", "copy_start_row": 2, "copy_end_row": 3}]
            with mock.patch.object(main, "_calculate_in_wps", side_effect=main.ProcessingError("模拟 COM 故障")):
                with self.assertRaisesRegex(main.ProcessingError, "模拟 COM 故障"):
                    main.process_workbooks(settings, lambda _msg: None)
            output_files = list((fake / "output").iterdir())
            self.assertEqual([], output_files)

    def test_invalid_autofilter_is_removed_only_in_template_read_copy(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            fake = create_fake_data(root / "fake")
            template_path = fake / "template.xlsx"
            template = load_workbook(template_path)
            template["1当天派单"].auto_filter.ref = "A1:C20"
            template.save(template_path)
            template.close()
            corrupt_autofilter_ref(template_path, "xl/worksheets/sheet1.xml", "A1:C20", "A1-BROKEN")
            original_bytes = template_path.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["template_path"] = str(template_path)
            staging = root / "staging.xlsx"
            messages = []

            main._create_staging_workbook(
                template_path, settings, [[["TEMPLATE-FAKE", "value", 1]]], staging, messages.append
            )
            self.assertEqual(original_bytes, template_path.read_bytes())
            self.assertTrue(any(
                '模板工作簿“template.xlsx”工作表“1当天派单”的自动筛选区域 \'A1-BROKEN\'' in message
                and "临时读取副本" in message for message in messages
            ))
            result = load_workbook(staging, data_only=False)
            try:
                self.assertEqual("TEMPLATE-FAKE", result["1当天派单"]["A3"].value)
            finally:
                result.close()

    def test_invalid_autofilter_is_removed_only_in_source_read_copy(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            make_test_workbook(source_path, [("源数据", {1: {1: "标题"}, 2: {1: "ROW-1", 2: "x"}, 3: {1: "ROW-2", 2: "y"}})])
            source = load_workbook(source_path)
            source.active.auto_filter.ref = "A1:B3"
            source.save(source_path)
            source.close()
            corrupt_autofilter_ref(source_path, "xl/worksheets/sheet1.xml", "A1:B3", "not-an-area")
            original_bytes = source_path.read_bytes()
            messages = []

            records = main._read_source_records(source_path, "源数据", 2, 1, 2, messages.append, end_row=3)
            self.assertEqual([["ROW-1", "x"], ["ROW-2", "y"]], records)
            self.assertEqual(original_bytes, source_path.read_bytes())
            self.assertTrue(any(
                '源工作簿“source.xlsx”工作表“源数据”的自动筛选区域 \'not-an-area\'' in message
                and "临时读取副本" in message for message in messages
            ))

    def test_valid_filters_table_filter_and_unfiltered_workbook_are_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            workbook_path = root / "mixed.xlsx"
            workbook = Workbook()
            invalid_sheet = workbook.active
            invalid_sheet.title = "需要临时修复"
            invalid_sheet.append(["A", "B"])
            invalid_sheet.append([1, 2])
            invalid_sheet.auto_filter.ref = "A1:B2"
            valid_sheet = workbook.create_sheet("有效筛选")
            valid_sheet.append(["A", "B"])
            valid_sheet.append([1, 2])
            valid_sheet.auto_filter.ref = "A1:B2"
            table = Table(displayName="FakeTable", ref="A1:B2")
            valid_sheet.add_table(table)
            workbook.save(workbook_path)
            workbook.close()
            corrupt_autofilter_ref(workbook_path, "xl/worksheets/sheet1.xml", "A1:B2", "A1:BAD")
            original_bytes = workbook_path.read_bytes()
            with zipfile.ZipFile(workbook_path) as archive:
                original_invalid_sheet = archive.read("xl/worksheets/sheet1.xml")
                original_valid_sheet = archive.read("xl/worksheets/sheet2.xml")
                original_table = archive.read("xl/tables/table1.xml")
            expected_invalid_sheet, removed_count = re.subn(
                rb'<autoFilter\b[^>]*\bref="A1:BAD"[^>]*/>',
                b"", original_invalid_sheet, count=1,
            )
            self.assertEqual(1, removed_count)

            messages = []
            with main._temporary_workbook_copy(workbook_path, "测试工作簿", messages.append) as clean_copy:
                with zipfile.ZipFile(clean_copy) as archive:
                    self.assertEqual(expected_invalid_sheet, archive.read("xl/worksheets/sheet1.xml"))
                    self.assertEqual(original_valid_sheet, archive.read("xl/worksheets/sheet2.xml"))
                    self.assertEqual(original_table, archive.read("xl/tables/table1.xml"))
                    self.assertRegex(archive.read("xl/worksheets/sheet2.xml"), rb'<autoFilter\b[^>]*\bref="A1:B2"[^>]*/>')
                    self.assertRegex(archive.read("xl/tables/table1.xml"), rb'<autoFilter\b[^>]*\bref="A1:B2"[^>]*/>')
                loaded = load_workbook(clean_copy, data_only=False)
                try:
                    self.assertEqual("A1:B2", loaded["有效筛选"].auto_filter.ref)
                finally:
                    loaded.close()
            self.assertEqual(original_bytes, workbook_path.read_bytes())
            self.assertEqual(1, len(messages))
            self.assertIn("需要临时修复", messages[0])

            no_filter_path = root / "no-filter.xlsx"
            make_test_workbook(no_filter_path, [("普通表", {1: {1: "HEADER"}, 2: {1: "FAKE"}})])
            no_filter_bytes = no_filter_path.read_bytes()
            messages = []
            with main._temporary_workbook_copy(no_filter_path, "测试工作簿", messages.append) as clean_copy:
                with zipfile.ZipFile(no_filter_path) as source_zip, zipfile.ZipFile(clean_copy) as copy_zip:
                    self.assertEqual(source_zip.read("xl/worksheets/sheet1.xml"), copy_zip.read("xl/worksheets/sheet1.xml"))
            self.assertEqual(no_filter_bytes, no_filter_path.read_bytes())
            self.assertEqual([], messages)

    def test_row_sortstate_ref_normalizes_and_valid_a1_sortstate_is_unchanged(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            workbook_path = root / "sortstate.xlsx"
            make_sortstate_workbook(workbook_path)
            original_bytes = workbook_path.read_bytes()
            messages = []

            with main._temporary_workbook_copy(workbook_path, "测试工作簿", messages.append) as clean_copy:
                with zipfile.ZipFile(clean_copy) as archive:
                    row_sorted_xml = archive.read("xl/worksheets/sheet1.xml")
                    a1_sorted_xml = archive.read("xl/worksheets/sheet2.xml")
                self.assertIn(b'<sortState ref="A3:XFD393"', row_sorted_xml)
                self.assertIn(b'<sortState ref="A1:B2"', a1_sorted_xml)
                loaded = load_workbook(clean_copy, data_only=False, keep_links=False)
                try:
                    self.assertEqual("A3:XFD393", loaded["整行排序"].auto_filter.sortState.ref)
                    self.assertEqual("A1:B2", loaded["A1排序"].auto_filter.sortState.ref)
                finally:
                    loaded.close()

            self.assertEqual(original_bytes, workbook_path.read_bytes())
            self.assertEqual(1, len(messages))
            self.assertIn(str(workbook_path), messages[0])
            self.assertIn("整行排序", messages[0])
            self.assertIn("'3:393'", messages[0])
            self.assertIn("'A3:XFD393'", messages[0])

    def test_sortstate_preprocessing_is_used_for_template_and_source_entries(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            make_sortstate_workbook(source_path)
            source_bytes = source_path.read_bytes()
            source_messages = []
            records = main._read_source_records(
                source_path, "整行排序", 2, 1, 2, source_messages.append, end_row=3
            )
            self.assertEqual([["FAKE-1", 1], ["FAKE-2", 2]], records)
            self.assertEqual(source_bytes, source_path.read_bytes())
            self.assertTrue(any(
                str(source_path) in message and "整行排序" in message
                and "'3:393'" in message and "'A3:XFD393'" in message
                for message in source_messages
            ))

            fake = create_fake_data(root / "fake")
            template_path = fake / "template.xlsx"
            template = load_workbook(template_path)
            first = template["1当天派单"]
            first.auto_filter.ref = "A1:C10"
            first.auto_filter.add_sort_condition("A3:A10")
            second = template["2当天归档"]
            second.auto_filter.ref = "A1:C10"
            second.auto_filter.add_sort_condition("A3:A10")
            template.save(template_path)
            template.close()
            corrupt_sortstate_ref(template_path, "xl/worksheets/sheet1.xml", "A1:C10", "3:393")
            template_bytes = template_path.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["template_path"] = str(template_path)
            staging = root / "staging.xlsx"
            template_messages = []

            main._create_staging_workbook(
                template_path, settings, [[ ["FAKE-TEMPLATE", "value", 1] ]],
                staging, template_messages.append,
            )
            self.assertEqual(template_bytes, template_path.read_bytes())
            self.assertTrue(any(
                str(template_path) in message and "1当天派单" in message
                and "'3:393'" in message and "'A3:XFD393'" in message
                for message in template_messages
            ))
            prepared = load_workbook(staging, data_only=False, keep_links=False)
            try:
                self.assertEqual("A3:XFD393", prepared["1当天派单"].auto_filter.sortState.ref)
                self.assertEqual("A1:C10", prepared["2当天归档"].auto_filter.sortState.ref)
            finally:
                prepared.close()

    def test_temporary_copy_closes_source_and_template_workbooks_before_unlink(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            make_test_workbook(source_path, [("数据", {1: {1: "标题"}, 2: {1: "FAKE-1"}})])
            fake = create_fake_data(root / "fake")
            template_path = fake / "template.xlsx"
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [deepcopy(settings["jobs"][0])]
            settings["template_path"] = str(template_path)

            def expect_processing_error(operation):
                with self.assertRaises(main.ProcessingError):
                    operation()

            def assert_close_precedes_unlink(operation, expected_book_count):
                events = []
                real_load = main.load_workbook
                real_unlink = Path.unlink

                def observed_load(path, *args, **kwargs):
                    workbook = real_load(path, *args, **kwargs)
                    if Path(path).name.startswith("wps_sheet_autofilter_"):
                        original_close = workbook.close

                        def observed_close():
                            events.append("close")
                            return original_close()

                        workbook.close = observed_close
                    return workbook

                def observed_unlink(path, *args, **kwargs):
                    if path.name.startswith("wps_sheet_autofilter_"):
                        events.append("unlink")
                    return real_unlink(path, *args, **kwargs)

                with mock.patch.object(main, "load_workbook", side_effect=observed_load), \
                     mock.patch.object(Path, "unlink", new=observed_unlink):
                    operation()
                self.assertEqual(["close"] * expected_book_count + ["unlink"], events)

            assert_close_precedes_unlink(
                lambda: main._read_source_records(source_path, "数据", 2, 1, 1, lambda _msg: None, end_row=2),
                2,
            )
            assert_close_precedes_unlink(
                lambda: expect_processing_error(
                    lambda: main._read_source_records(source_path, "不存在", 2, 1, 1, lambda _msg: None, end_row=2)
                ),
                2,
            )
            assert_close_precedes_unlink(
                lambda: main._create_staging_workbook(
                    template_path, settings, [[ ["FAKE-ROW", "x", 1] ]], root / "staging.xlsx", lambda _msg: None
                ),
                1,
            )
            assert_close_precedes_unlink(
                lambda: main._export_table_areas(
                    source_path, [{"sheet": "数据", "range": "A1:A2"}], root / "region-exports", lambda _msg: None
                ),
                1,
            )
            broken_settings = deepcopy(settings)
            broken_settings["jobs"][0]["template_sheet"] = "不存在"
            assert_close_precedes_unlink(
                lambda: expect_processing_error(
                    lambda: main._create_staging_workbook(
                        template_path, broken_settings, [[["FAKE-ROW", "x", 1]]], root / "broken-staging.xlsx", lambda _msg: None
                    )
                ),
                1,
            )

    def test_locked_temporary_copy_retries_and_warns_without_masking_read_result(self):
        with tempfile.TemporaryDirectory() as td:
            source_path = Path(td) / "source.xlsx"
            make_test_workbook(source_path, [("数据", {1: {1: "标题"}, 2: {1: "FAKE-1"}})])
            real_unlink = Path.unlink

            for requested_sheet, expected_success in (("数据", True), ("不存在", False)):
                with self.subTest(expected_success=expected_success):
                    messages = []
                    attempts = []
                    locked_paths = []

                    def locked_unlink(path, *args, **kwargs):
                        if path.name.startswith("wps_sheet_autofilter_"):
                            attempts.append(path)
                            locked_paths.append(path)
                            raise PermissionError(13, "simulated sharing violation", str(path))
                        return real_unlink(path, *args, **kwargs)

                    with mock.patch.object(Path, "unlink", new=locked_unlink), \
                         mock.patch.object(main.time, "sleep"):
                        if expected_success:
                            records = main._read_source_records(
                                source_path, requested_sheet, 2, 1, 1, messages.append, end_row=2
                            )
                            self.assertEqual([["FAKE-1"]], records)
                        else:
                            with self.assertRaises(main.ProcessingError) as caught:
                                main._read_source_records(
                                    source_path, requested_sheet, 2, 1, 1, messages.append, end_row=2
                                )
                            self.assertIn("找不到配置的工作表", str(caught.exception))
                            self.assertNotIn("无法清理", str(caught.exception))

                    self.assertEqual(len(main._TEMP_COPY_CLEANUP_RETRY_DELAYS) + 1, len(attempts))
                    self.assertTrue(locked_paths[-1].exists())
                    warning = next(message for message in messages if "临时读取副本" in message and "无法清理" in message)
                    self.assertIn(str(locked_paths[-1]), warning)
                    self.assertIn("手动删除", warning)
                    if expected_success:
                        self.assertIn("读取/处理已成功", warning)
                        self.assertNotIn("Traceback", warning)
                    else:
                        self.assertIn("读取/处理本身已失败", warning)
                    real_unlink(locked_paths[-1])

    def test_locked_region_read_copy_retries_and_warns_without_failing_export(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "saved-staging.xlsx"
            make_test_workbook(source_path, [("数据", {1: {1: "标题"}, 2: {1: "FAKE-1"}})])
            output_dir = root / "exports"
            messages = []
            attempts = []
            locked_paths = []
            real_unlink = Path.unlink

            def locked_unlink(path, *args, **kwargs):
                if path.name.startswith("wps_sheet_autofilter_"):
                    attempts.append(path)
                    locked_paths.append(path)
                    raise PermissionError(13, "simulated sharing violation", str(path))
                return real_unlink(path, *args, **kwargs)

            with mock.patch.object(Path, "unlink", new=locked_unlink), \
                 mock.patch.object(main.time, "sleep"):
                paths = main._export_table_areas(
                    source_path, [{"sheet": "数据", "range": "A1:A2"}], output_dir, messages.append
                )

            self.assertEqual(1, len(paths))
            self.assertTrue(paths[0].is_file())
            self.assertEqual(len(main._TEMP_COPY_CLEANUP_RETRY_DELAYS) + 1, len(attempts))
            self.assertTrue(locked_paths[-1].exists())
            warning = next(message for message in messages if "临时读取副本" in message and "无法清理" in message)
            self.assertIn("读取/处理已成功", warning)
            self.assertIn(str(locked_paths[-1]), warning)
            self.assertNotIn("Traceback", warning)
            real_unlink(locked_paths[-1])

    def test_temporary_copy_names_remain_unique(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.xlsx"
            make_test_workbook(path, [("数据", {1: {1: "FAKE"}})])
            with main._temporary_workbook_copy(path, "测试工作簿", lambda _msg: None) as first:
                with main._temporary_workbook_copy(path, "测试工作簿", lambda _msg: None) as second:
                    self.assertNotEqual(first.path, second.path)
                    self.assertTrue(first.path.name.startswith("wps_sheet_autofilter_"))
                    self.assertTrue(second.path.name.startswith("wps_sheet_autofilter_"))

    def test_invalid_sortstate_ref_fails_explicitly_without_deleting_sort_info(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            for invalid_ref in ("0:393", "1048577:1048578", "393:3", "3:393x"):
                with self.subTest(ref=invalid_ref):
                    workbook_path = root / f"invalid-{len(list(root.iterdir()))}.xlsx"
                    make_sortstate_workbook(workbook_path)
                    corrupt_sortstate_ref(workbook_path, "xl/worksheets/sheet1.xml", "3:393", invalid_ref)
                    original_bytes = workbook_path.read_bytes()
                    with self.assertRaises(main.ProcessingError) as caught:
                        with main._temporary_workbook_copy(workbook_path, "测试工作簿", lambda _msg: None):
                            self.fail("不安全的 sortState 引用不应被静默删除或接受")
                    self.assertIn("sortState ref", str(caught.exception))
                    self.assertIn(invalid_ref, str(caught.exception))
                    self.assertIn("未删除该排序定义", str(caught.exception))
                    self.assertEqual(original_bytes, workbook_path.read_bytes())

    def test_safe_copy_rejects_zip_traversal_member_count_and_uncompressed_size(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            path = root / "fixture.xlsx"
            make_test_workbook(path, [("数据", {1: {1: "FAKE"}})])
            original = path.read_bytes()
            traversal_path = root / "traversal.xlsx"
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(traversal_path, "w") as destination:
                for info in source.infolist():
                    destination.writestr(copy(info), source.read(info.filename))
                destination.writestr("../outside.xml", b"test")
            temp_copies_before = set(Path(tempfile.gettempdir()).glob("wps_sheet_autofilter_*.xlsx"))
            with self.assertRaisesRegex(main.ProcessingError, "路径不安全") as caught:
                with main._temporary_workbook_copy(traversal_path, "测试工作簿", lambda _msg: None):
                    self.fail("不安全 ZIP 路径不应被接受")
            self.assertIn("在 WPS 中打开该文件的副本", str(caught.exception))
            self.assertEqual(temp_copies_before, set(Path(tempfile.gettempdir()).glob("wps_sheet_autofilter_*.xlsx")))

            with mock.patch.object(main, "MAX_XLSX_MEMBERS", 1):
                with self.assertRaisesRegex(main.ProcessingError, "成员数"):
                    with main._temporary_workbook_copy(path, "测试工作簿", lambda _msg: None):
                        self.fail("超限 ZIP 成员数不应被接受")
            with mock.patch.object(main, "MAX_XLSX_UNCOMPRESSED_BYTES", 1):
                with self.assertRaisesRegex(main.ProcessingError, "解压总大小"):
                    with main._temporary_workbook_copy(path, "测试工作簿", lambda _msg: None):
                        self.fail("超限 ZIP 解压大小不应被接受")
            self.assertEqual(original, path.read_bytes())


class CustomFormulaFillTests(unittest.TestCase):
    def _run_formula_rule(self, rule, initial_cells=None):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "formula.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            for address, value in (initial_cells or {}).items():
                worksheet[address] = value
            workbook.save(path)
            workbook.close()
            events = []
            com_workbook = FakeComWorkbook(path, events)
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = []
            settings["CustomFormula"] = {"enabled": True, "rules": [rule]}
            app = FakeWpsApplication(events, {})
            app.workbook = com_workbook
            logs = []
            try:
                applied = main._apply_region_processing(com_workbook, settings, [], logs.append, application=app)
                values = {
                    address: com_workbook.workbook["Data"][address].value
                    for address in ("B2", "C2", "B3", "C3")
                }
            finally:
                com_workbook.workbook.close()
            return applied, events, values, logs

    def test_custom_formula_fills_used_range_as_relative_r1c1(self):
        rule = {"scope": "target", "sheet": "Data", "range": "A1", "formula_mode": "custom",
                "formula": "=$A2+B$1", "seed_cell": ""}
        applied, events, values, logs = self._run_formula_rule(
            rule, {"B2": 1, "C2": 2, "B3": 3, "C3": 4})
        self.assertTrue(applied)
        self.assertEqual({address: "=$A2+B$1" for address in ("B2", "C2", "B3", "C3")}, values)
        self.assertIn(("formula-r1c1", "B2:C3", "=$A2+B$1"), events)
        self.assertIn("calculate", events)
        self.assertTrue(any("按相对引用向下/向右填充" in message for message in logs))

    def test_seed_cell_formula_fills_selected_region(self):
        rule = {"scope": "target", "sheet": "Data", "range": "B2:C3", "formula_mode": "cell",
                "formula": "", "seed_cell": "A1"}
        applied, events, values, _logs = self._run_formula_rule(
            rule, {"A1": "=$A$1+1", "B2": 1, "C2": 2, "B3": 3, "C3": 4})
        self.assertTrue(applied)
        self.assertEqual({address: "=$A$1+1" for address in ("B2", "C2", "B3", "C3")}, values)
        self.assertIn(("formula-r1c1", "B2:C3", "=$A$1+1"), events)

    def test_seed_cell_without_formula_is_rejected(self):
        rule = {"scope": "target", "sheet": "Data", "range": "B2:C3", "formula_mode": "cell",
                "formula": "", "seed_cell": "A1"}
        with self.assertRaisesRegex(main.ProcessingError, "不含公式"):
            self._run_formula_rule(rule, {"A1": "普通文本", "B2": 1, "C2": 2, "B3": 3, "C3": 4})

    def test_custom_formula_can_override_template_formula_area_and_protects_it_later(self):
        rule = {"scope": "target", "sheet": "Data", "range": "B3:C4", "formula_mode": "custom",
                "formula": "=A3+1", "seed_cell": ""}
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "formula-area.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["B3"] = "old"
            worksheet["C3"] = "old"
            worksheet["B4"] = "old"
            worksheet["C4"] = "old"
            workbook.save(path)
            workbook.close()
            events = []
            com_workbook = FakeComWorkbook(path, events)
            app = FakeWpsApplication(events, {})
            app.workbook = com_workbook
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = [{"template_sheet": "Data", "start_row": 3, "formula_range": "A:C"}]
            settings["CustomFormula"] = {"enabled": True, "rules": [rule]}
            formula_jobs = [("Data", 3, 4, [])]
            try:
                self.assertTrue(main._apply_region_processing(
                    com_workbook, settings, formula_jobs, lambda _msg: None, application=app))
                settings["FormatConvert"] = {"enabled": True, "rules": [{
                    "scope": "target", "sheet": "Data", "range": "B3:C4", "mode": "text"}]}
                with self.assertRaisesRegex(main.ProcessingError, "与已配置的公式区域重叠"):
                    main._apply_region_processing(
                        com_workbook, settings, formula_jobs, lambda _msg: None, application=app)
            finally:
                com_workbook.workbook.close()

    def test_workflow_requires_formula_module_and_places_it_before_region_operations(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["CustomFormula"] = {"enabled": True, "rules": [{
            "scope": "target", "sheet": "1当天派单", "range": "B5", "formula_mode": "custom",
            "formula": "=1+1", "seed_cell": ""}]}
        main.validate_workflow(settings)
        workflow = list(settings["workflow"])
        workflow.remove("CustomFormula")
        workflow.insert(workflow.index("TextToColumns") + 1, "CustomFormula")
        settings["workflow"] = workflow
        with self.assertRaisesRegex(ValueError, "区域公式填充.*文本分列"):
            main.validate_workflow(settings)

    def test_source_formula_runs_in_wps_copy_before_read_and_original_is_preserved(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source_path = root / "source.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Data"
            worksheet["A1"] = 2
            worksheet["A2"] = 3
            workbook.save(source_path)
            workbook.close()
            original_bytes = source_path.read_bytes()
            output_dir = root / "output"
            output_dir.mkdir()
            final_path = output_dir / "result.xlsx"
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["retain_edited_sources"] = True
            settings["CustomFormula"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "B1:B2", "formula_mode": "custom",
                "formula": "=A1+1", "seed_cell": ""}]}
            events = []
            fake_app = FakeWpsApplication(events, {})
            used = set()

            def read_after_wps(_path, *_args):
                events.append("read-records")
                return [[2], [3]]

            with mock.patch.object(main, "_connect_wps", return_value=(fake_app, None, "fake.WPS")), \
                 mock.patch.object(main, "_read_source_records", side_effect=read_after_wps):
                records = main._read_source_records_with_rules(
                    source_path, "Data", 1, 1, 1, lambda _msg: None, None, settings, used,
                    True, output_dir, final_path)

            self.assertEqual([[2], [3]], records)
            self.assertIn(("CustomFormula", 0), used)
            self.assertLess(events.index("calculate"), events.index("save"))
            self.assertLess(events.index("save"), events.index("read-records"))
            self.assertEqual(original_bytes, source_path.read_bytes())
            retained = list(output_dir.glob("source_已编辑*.xlsx"))
            self.assertEqual(1, len(retained))
            saved = load_workbook(retained[0], data_only=False)
            try:
                self.assertTrue(str(saved["Data"]["B1"].value).startswith("="))
                self.assertTrue(str(saved["Data"]["B2"].value).startswith("="))
            finally:
                saved.close()

    def test_formula_rule_static_validation_accepts_both_modes_and_rejects_unsafe_input(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["CustomFormula"] = {"enabled": True, "rules": [{
            "scope": "target", "sheet": "1当天派单", "range": "B5:C6", "formula_mode": "custom",
            "formula": "=SUM(B1:B2)", "seed_cell": ""}]}
        main.validate_static_settings(settings)
        settings["CustomFormula"]["rules"][0].update({"formula_mode": "cell", "formula": "", "seed_cell": "B2"})
        main.validate_static_settings(settings)
        for mode, formula, seed, expected in (
            ("custom", "SUM(A1:A2)", "", "以 = 开头"),
            ("custom", '=WEBSERVICE("https://example.invalid")', "", "外部工作簿引用或外部访问函数"),
            ("cell", "", "A1:B2", "单个单元格"),
        ):
            settings["CustomFormula"]["rules"][0].update(
                {"formula_mode": mode, "formula": formula, "seed_cell": seed})
            with self.subTest(mode=mode, formula=formula, seed=seed):
                with self.assertRaisesRegex(main.ProcessingError, expected):
                    main.validate_static_settings(settings)


if __name__ == "__main__":
    unittest.main(verbosity=2)


class RegionFilterRangeShorthandTests(unittest.TestCase):
    def test_region_filter_range_shorthand_and_legacy_rectangle(self):
        bounds, kind = main._parse_region_filter_area("")
        self.assertEqual((1, 1, 16_384, 1_048_576), bounds)
        self.assertEqual("sheet", kind)
        bounds, kind = main._parse_region_filter_area("A")
        self.assertEqual((1, 1, 1, 1_048_576), bounds)
        self.assertEqual("column", kind)
        self.assertEqual(((1, 1, 1, 1_048_576), "column"), main._parse_region_filter_area("a"))
        bounds, kind = main._parse_region_filter_area("a1")
        self.assertEqual((1, 1, 1, 1), bounds)
        self.assertEqual("cell", kind)
        self.assertEqual(((1, 1, 1, 1), "cell"), main._parse_region_filter_area("A1"))
        bounds, kind = main._parse_region_filter_area("A1:F100")
        self.assertEqual((1, 1, 6, 100), bounds)
        self.assertEqual("rectangle", kind)
        bounds, kind = main._parse_region_filter_area("A1")
        self.assertEqual("cell", kind)
        self.assertEqual((1, 1, 1, 1), bounds)

    def test_blank_filter_sheet_and_range_use_first_sheet_and_used_range(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "default-sheet.xlsx"
            make_test_workbook(path, [
                ("Primary", {1: {1: "ID", 2: "Status"},
                             2: {1: "KEEP-1", 2: "Keep"},
                             3: {1: "DROP", 2: "Drop"}}),
                ("Secondary", {1: {1: "OTHER", 2: "Keep"}}),
            ])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "target", "sheet": "", "range": "", "has_header": True,
                "filter_column": "B", "operator": "equals", "criteria": "Keep",
            }]}
            main.validate_static_settings(settings)
            messages = []
            fake = FakeComWorkbook(path, [])
            self.assertTrue(main._apply_region_processing(fake, settings, [], messages.append))
            self.assertEqual("KEEP-1", fake.workbook["Primary"]["A2"].value)
            self.assertIsNone(fake.workbook["Primary"]["A3"].value)
            self.assertEqual("OTHER", fake.workbook["Secondary"]["A1"].value)
            self.assertTrue(any("默认使用第一个工作表“Primary”" in message for message in messages))
            fake.workbook.close()

            settings["RegionFilter"]["rules"][0]["sheet"] = "pRiMaRy"
            fake = FakeComWorkbook(path, [])
            self.assertTrue(main._apply_region_processing(fake, settings, [], lambda _msg: None))
            self.assertIsNone(fake.workbook["Primary"]["A3"].value)
            fake.workbook.close()

    def test_region_filter_single_column_validates_without_full_column_cell_limit(self):
        settings = deepcopy(main.DEFAULT_SETTINGS)
        settings["RegionFilter"] = {"enabled": True, "rules": [{
            "scope": "target", "sheet": "数据", "range": "A", "has_header": True,
            "filter_column": "A", "operator": "equals", "criteria": "保留",
        }]}
        main.validate_static_settings(settings)

    def test_upper_and_lowercase_single_cell_filter_ranges_are_equivalent(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "single-cell-filter.xlsx"
            make_test_workbook(path, [("Primary", {
                1: {1: "Drop", 2: "Keep-other-cell"},
                2: {1: "Keep-row", 2: "Keep-other-row"},
            })])
            for area in ("A1", "a1"):
                with self.subTest(area=area):
                    settings = deepcopy(main.DEFAULT_SETTINGS)
                    settings["RegionFilter"] = {"enabled": True, "rules": [{
                        "scope": "target", "sheet": "primary", "range": area,
                        "has_header": False, "filter_column": "A",
                        "operator": "equals", "criteria": "Keep",
                    }]}
                    main.validate_static_settings(settings)
                    fake = FakeComWorkbook(path, [])
                    self.assertTrue(main._apply_region_processing(fake, settings, [], lambda _msg: None))
                    sheet = fake.workbook["Primary"]
                    self.assertIsNone(sheet["A1"].value)
                    self.assertEqual("Keep-other-cell", sheet["B1"].value)
                    self.assertEqual("Keep-row", sheet["A2"].value)
                    fake.workbook.close()

    def test_sheet_row_and_column_selectors_resolve_to_used_range(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet["B2"] = "used-start"
        worksheet["D4"] = "used-end"
        fake_sheet = FakeComWorksheet(worksheet, [])
        self.assertEqual((2, 2, 4, 4), main._resolve_com_selection_bounds(
            fake_sheet, (1, 1, 16_384, 1_048_576), "sheet", "整表测试"))
        self.assertEqual((2, 3, 4, 3), main._resolve_com_selection_bounds(
            fake_sheet, (1, 3, 16_384, 3), "row", "整行测试"))
        self.assertEqual((3, 1, 3, 4), main._resolve_com_selection_bounds(
            fake_sheet, (3, 1, 3, 1_048_576), "column", "整列测试"))
        workbook.close()

    def test_whole_sheet_region_filter_runs_only_on_used_range(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "whole-sheet-filter.xlsx"
            make_test_workbook(path, [("数据", {
                2: {2: "名称", 3: "状态"},
                3: {2: "保留-1", 3: "保留"},
                4: {2: "删除", 3: "丢弃"},
            })])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "target", "sheet": "数据", "range": "", "has_header": True,
                "filter_column": "C", "operator": "equals", "criteria": "保留",
            }]}
            main.validate_static_settings(settings)
            settings["jobs"] = []
            events = []
            fake_workbook = FakeComWorkbook(path, events)
            self.assertTrue(main._apply_region_processing(fake_workbook, settings, [], lambda _msg: None))
            self.assertIn(("read", "B2:C4"), events)
            sheet = fake_workbook.workbook["数据"]
            self.assertEqual(["名称", "保留-1", None], [sheet.cell(row=row, column=2).value for row in range(2, 5)])
            self.assertEqual(["状态", "保留", None], [sheet.cell(row=row, column=3).value for row in range(2, 5)])
            fake_workbook.workbook.close()

    def test_whole_sheet_format_conversion_runs_only_on_used_range(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "whole-sheet-format.xlsx"
            make_test_workbook(path, [("数据", {2: {2: "12", 3: "13.5"}, 4: {4: "9"}})])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "target", "sheet": "数据", "range": "A1", "mode": "number",
            }]}
            main.validate_static_settings(settings)
            settings["jobs"] = []
            events = []
            fake_workbook = FakeComWorkbook(path, events)
            self.assertTrue(main._apply_region_processing(fake_workbook, settings, [], lambda _msg: None))
            self.assertIn(("read", "B2:D4"), events)
            sheet = fake_workbook.workbook["数据"]
            self.assertEqual(12, sheet["B2"].value)
            self.assertEqual(13.5, sheet["C2"].value)
            self.assertEqual(9, sheet["D4"].value)
            fake_workbook.workbook.close()

    def test_target_region_filter_single_column_uses_used_rows_only(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "target-filter-column.xlsx"
            make_test_workbook(path, [("数据", {
                1: {1: "状态", 2: "关联列"},
                2: {1: "保留", 2: "记录-1"},
                3: {1: "丢弃", 2: "记录-2"},
                4: {1: "保留", 2: "记录-3"},
            })])
            fake_workbook = FakeComWorkbook(path, [])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["jobs"] = []
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "target", "sheet": "数据", "range": "A", "has_header": True,
                "filter_column": "A", "operator": "equals", "criteria": "保留",
            }]}
            self.assertTrue(main._apply_region_processing(fake_workbook, settings, [], lambda _msg: None))
            sheet = fake_workbook.workbook["数据"]
            self.assertEqual(["状态", "保留", "保留", None], [sheet.cell(row=row, column=1).value for row in range(1, 5)])
            self.assertEqual(["关联列", "记录-1", "记录-2", "记录-3"], [sheet.cell(row=row, column=2).value for row in range(1, 5)])
            fake_workbook.workbook.close()

    def test_source_region_filter_single_column_covers_blank_key_rows_through_sheet_data(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source-filter-column.xlsx"
            make_test_workbook(path, [("Data", {
                1: {1: "状态", 2: "记录"},
                2: {1: "保留", 2: "first"},
                3: {1: "丢弃", 2: "remove"},
                4: {1: "保留", 2: "last"},
                5: {2: "empty-key-row"},
            })])
            original = path.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "A", "has_header": True,
                "filter_column": "A", "operator": "equals", "criteria": "保留",
            }]}
            records = main._read_source_records(path, "Data", 1, 1, 2, lambda _msg: None,
                                               source_processing=settings)
            self.assertEqual([["状态", "记录"], ["保留", "first"], ["保留", "last"]], records)
            self.assertEqual(original, path.read_bytes())


class SourceMergeTests(unittest.TestCase):
    def _make_source(self, path, title, rows):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = title
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    def test_merge_is_top_to_bottom_keeps_first_source_header_and_optionally_skips_later(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            first = root / "first.xlsx"
            second = root / "second.xlsx"
            self._make_source(first, "Data", [["ID", "Value"], ["A", 1]])
            self._make_source(second, "Other", [["ID", "Value"], ["B", 2]])
            originals = {path: path.read_bytes() for path in (first, second)}
            sources = [
                {"path": str(first), "source_sheet": "Data", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": False},
                {"path": str(second), "source_sheet": "Other", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": False},
            ]
            merged = main._merge_source_workbooks(
                sources, deepcopy(main.DEFAULT_SETTINGS), 1, 2, lambda _msg: None
            )
            try:
                records, retained = main._process_merged_source_workbook(
                    merged, main.DEFAULT_SETTINGS, 1, 2, lambda _msg: None,
                    False, root, root / "result.xlsx"
                )
                self.assertIsNone(retained)
                self.assertEqual([["ID", "Value"], ["A", 1], ["B", 2]], records)
                self.assertEqual(originals, {path: path.read_bytes() for path in originals})
            finally:
                merged["path"].unlink(missing_ok=True)

    def test_later_source_header_can_be_retained_and_source_rules_map_to_merged_rows(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            first, second = root / "first.xlsx", root / "second.xlsx"
            self._make_source(first, "Data", [["ID", "Value"], ["A", 1]])
            self._make_source(second, "Data", [["ID", "Value"], ["B", 2]])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "Data", "range": "A1:B20", "mode": "text",
            }]}
            sources = [
                {"path": str(first), "source_sheet": "Data", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": True},
                {"path": str(second), "source_sheet": "Data", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": True},
            ]
            merged = main._merge_source_workbooks(sources, settings, 1, 2, lambda _msg: None)
            try:
                self.assertEqual(
                    ["A1:B2", "A3:B4"],
                    [rule["range"] for rule in merged["settings"]["FormatConvert"]["rules"]],
                )
                self.assertTrue(all(rule["sheet"] == "合并源数据" and rule["scope"] == "target"
                                    for rule in merged["settings"]["FormatConvert"]["rules"]))
                records, _retained = main._process_merged_source_workbook(
                    {**merged, "has_rules": False}, settings, 1, 2, lambda _msg: None,
                    False, root, root / "result.xlsx"
                )
                self.assertEqual([["ID", "Value"], ["A", 1], ["ID", "Value"], ["B", 2]], records)
            finally:
                merged["path"].unlink(missing_ok=True)

    def test_blank_source_filter_uses_first_sheet_full_range_before_template_records_are_read(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            source = root / "source.xlsx"
            make_test_workbook(source, [
                ("Primary", {1: {1: "ID", 2: "Status"},
                             2: {1: "KEEP", 2: "Keep"},
                             3: {1: "DROP", 2: "Drop"}}),
                ("Secondary", {1: {1: "MUST-NOT-READ", 2: "Keep"}}),
            ])
            original = source.read_bytes()
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "", "range": "", "has_header": True,
                "filter_column": "B", "operator": "equals", "criteria": "Keep",
            }]}
            merged = main._merge_source_workbooks([{
                "path": str(source), "source_sheet": "", "copy_start_row": 1,
                "copy_end_row": None, "keep_first_row": True,
            }], settings, 1, 2, lambda _msg: None)
            try:
                mapped = merged["settings"]["RegionFilter"]["rules"][0]
                self.assertEqual("合并源数据", mapped["sheet"])
                self.assertEqual("A1:B3", mapped["range"])
                self.assertIn(("RegionFilter", 0), merged["used_rules"])

                events = []
                fake_app = FakeWpsApplication(events, {})
                fake_pythoncom = FakePythonCom(events)
                with mock.patch.object(main, "_connect_wps",
                                        return_value=(fake_app, fake_pythoncom, "Fake WPS")):
                    records, retained = main._process_merged_source_workbook(
                        merged, settings, 1, 2, lambda _msg: None,
                        False, root, root / "result.xlsx",
                    )
                self.assertIsNotNone(retained)
                self.assertTrue(retained.is_file())
                self.assertNotEqual(merged["path"], retained)
                self.assertEqual([["ID", "Status"], ["KEEP", "Keep"]], records)
                self.assertEqual(original, source.read_bytes())
            finally:
                merged["path"].unlink(missing_ok=True)

    def test_all_five_source_rule_types_map_first_sheet_and_match_names_case_insensitively(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "mixed-case.xlsx"
            make_test_workbook(source, [("Data", {
                1: {1: "ID", 2: "Amount", 3: "Pair"},
                2: {1: "A", 2: "10", 3: "X|1"},
                3: {1: "A", 2: "10", 3: "X|1"},
            }), ("Other", {1: {1: "NOT-FIRST-SHEET"}})])
            settings = deepcopy(main.DEFAULT_SETTINGS)
            settings["CustomFormula"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "dAtA", "range": "D2:D3",
                "formula_mode": "custom", "formula": '=A2&"-F"',
            }]}
            settings["TextToColumns"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "", "range": "C2:C3", "delimiter": "|",
                "output_start_col": "E", "column_formats": ["text", "text"],
            }]}
            settings["RegionFilter"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "", "range": "", "has_header": True,
                "filter_column": "A", "operator": "not_blank", "criteria": "",
            }]}
            settings["Deduplicate"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "DATA", "range": "A1:C3",
                "has_header": True, "key_mode": "columns", "key_columns": ["A"],
            }]}
            settings["FormatConvert"] = {"enabled": True, "rules": [{
                "scope": "source", "sheet": "", "range": "B2:B3", "mode": "text",
            }]}
            merged = main._merge_source_workbooks([{
                "path": str(source), "source_sheet": "dAtA", "copy_start_row": 1,
                "copy_end_row": None, "keep_first_row": True,
            }], settings, 1, 3, lambda _msg: None)
            try:
                expected = {"CustomFormula", "TextToColumns", "RegionFilter", "Deduplicate", "FormatConvert"}
                self.assertEqual(expected, {key for key, _index in merged["used_rules"]})
                for key in expected:
                    rules = merged["settings"][key]["rules"]
                    self.assertTrue(rules)
                    self.assertTrue(all(rule["sheet"] == "合并源数据" and rule["scope"] == "target"
                                        for rule in rules))
                self.assertEqual("A1:C3", merged["settings"]["RegionFilter"]["rules"][0]["range"])
            finally:
                merged["path"].unlink(missing_ok=True)


class StandaloneMergeTests(unittest.TestCase):
    def _make_source(self, path, sheet_name, rows):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    def _settings(self, root, sources, **updates):
        settings = {
            "config_version": main.CONFIG_VERSION,
            "merge_only": True,
            "template_path": "",
            "output_dir": str(root / "out"),
            "merge_output_filename": "custom-merge",
            "merge_sources": sources,
            "jobs": [],
            "workflow": [],
        }
        settings.update(updates)
        return settings

    def test_csv_then_xlsx_merges_without_template_or_workflow_and_preserves_sources(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output_dir = root / "out"
            output_dir.mkdir()
            csv_path = root / "first.csv"
            csv_path.write_bytes("\ufeff编号,名称\n001,CSV 记录\n003,\"=SUM(1,2)\"\n".encode("utf-8"))
            xlsx_path = root / "later.xlsx"
            self._make_source(xlsx_path, "首表", [["编号", "名称"], ["002", "XLSX 记录"]])
            original_bytes = {path: path.read_bytes() for path in (csv_path, xlsx_path)}
            settings = self._settings(root, [
                {"path": str(csv_path), "source_sheet": "", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": False},
                {"path": str(xlsx_path), "source_sheet": "", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": False},
            ], FormatConvert={"enabled": True, "rules": [{
                "scope": "target", "sheet": "does-not-exist", "range": "A1", "mode": "invalid",
            }]})
            result = main.process_workbooks(settings, lambda _message: None)
            self.assertEqual("custom-merge.xlsx", result.name)
            workbook = load_workbook(result, read_only=True, data_only=False)
            try:
                self.assertEqual([("编号", "名称"), ("001", "CSV 记录"), ("003", "=SUM(1,2)"),
                                  ("002", "XLSX 记录")],
                                 list(workbook["合并结果"].values))
                self.assertEqual("s", workbook["合并结果"]["B3"].data_type)
            finally:
                workbook.close()
            self.assertEqual(original_bytes, {path: path.read_bytes() for path in original_bytes})

    def test_gb18030_csv_and_later_header_retention(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output_dir = root / "out"
            output_dir.mkdir()
            csv_path = root / "first.csv"
            csv_path.write_bytes("列名;值\n甲;1\n".encode("gb18030"))
            xlsx_path = root / "later.xlsx"
            self._make_source(xlsx_path, "Data", [["列名", "值"], ["乙", 2]])
            result = main.process_merge_only(self._settings(root, [
                {"path": str(csv_path), "source_sheet": "", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": False},
                {"path": str(xlsx_path), "source_sheet": "Data", "copy_start_row": 1,
                 "copy_end_row": None, "keep_first_row": True},
            ], merge_output_filename="retained.xlsx"), lambda _message: None)
            workbook = load_workbook(result, read_only=True, data_only=True)
            try:
                self.assertEqual([("列名", "值"), ("甲", "1"), ("列名", "值"), ("乙", 2)],
                                 list(workbook.active.values))
            finally:
                workbook.close()

    def test_merge_only_requires_csv_only_in_first_position_and_prevents_overwrite(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            output_dir = root / "out"
            output_dir.mkdir()
            csv_path = root / "first.csv"
            csv_path.write_text("a\n1\n", encoding="utf-8")
            another_csv = root / "later.csv"
            another_csv.write_text("b\n2\n", encoding="utf-8")
            with self.assertRaisesRegex(main.ProcessingError, "后续文件必须是 .xlsx"):
                main.validate_static_settings(self._settings(root, [
                    {"path": str(csv_path)}, {"path": str(another_csv)},
                ]))
            xlsx_path = root / "source.xlsx"
            self._make_source(xlsx_path, "Data", [["x"]])
            source_before = xlsx_path.read_bytes()
            existing_result = output_dir / "custom-merge.xlsx"
            existing_result.write_bytes(b"existing")
            with self.assertRaisesRegex(main.ProcessingError, "已存在"):
                main.process_merge_only(self._settings(root, [
                    {"path": str(xlsx_path), "source_sheet": "", "copy_start_row": 1,
                     "copy_end_row": None, "keep_first_row": True},
                ]), lambda _message: None)
            self.assertEqual(b"existing", existing_result.read_bytes())
            self.assertEqual(source_before, xlsx_path.read_bytes())

    def test_v14_configuration_migrates_without_enabling_standalone_mode(self):
        legacy = deepcopy(main.DEFAULT_SETTINGS)
        legacy["config_version"] = 14
        for key in ("merge_only", "merge_sources", "merge_output_filename", "retain_merged_sources",
                    "delete_processed_sources_after_success"):
            legacy.pop(key, None)
        legacy["retain_edited_sources"] = True
        for job in legacy["jobs"]:
            job.pop("merged_source_filename", None)
        migrated = main.normalize_settings(legacy)
        self.assertEqual(main.CONFIG_VERSION, migrated["config_version"])
        self.assertFalse(migrated["merge_only"])
        self.assertEqual([], migrated["merge_sources"])
        self.assertEqual("合并结果.xlsx", migrated["merge_output_filename"])
        self.assertTrue(migrated["retain_merged_sources"])
        self.assertFalse(migrated["delete_processed_sources_after_success"])
        self.assertTrue(all(job["merged_source_filename"] == "" for job in migrated["jobs"]))
        self.assertIn("独立源文件合并", " ".join(main.configuration_migration_notes(legacy)))


class UiPreferencesAndBrandingTests(unittest.TestCase):
    def test_version_notice_preference_is_saved_and_loaded_separately(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "nested" / "ui_preferences.json"
            self.assertFalse(main.load_ui_preferences(path)["dont_show_version_notice"])
            main.save_ui_preferences(True, path)
            self.assertTrue(main.load_ui_preferences(path)["dont_show_version_notice"])
            self.assertFalse(path.with_name(path.name + ".tmp").exists())

    def test_configuration_notice_preference_is_independent(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "ui_preferences.json"
            main.save_ui_preferences(True, path)
            main.save_ui_preferences(dont_show_configuration_notice=True, path=path)
            preferences = main.load_ui_preferences(path)
            self.assertTrue(preferences["dont_show_version_notice"])
            self.assertTrue(preferences["dont_show_configuration_notice"])

            main.save_ui_preferences(False, path)
            preferences = main.load_ui_preferences(path)
            self.assertFalse(preferences["dont_show_version_notice"])
            self.assertTrue(preferences["dont_show_configuration_notice"])

    def test_malformed_ui_preferences_default_to_showing_notice(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "ui_preferences.json"
            path.write_text("{not-json", encoding="utf-8")
            self.assertFalse(main.load_ui_preferences(path)["dont_show_version_notice"])
            self.assertFalse(main.load_ui_preferences(path)["dont_show_configuration_notice"])
            path.write_text('{"dont_show_version_notice":"false"}', encoding="utf-8")
            self.assertFalse(main.load_ui_preferences(path)["dont_show_version_notice"])

    def test_version_and_configuration_notices_share_one_message(self):
        message = main.format_startup_notice(True, ["配置已升级到 v18：迁移提醒测试。"])
        self.assertIn("v26", message)
        self.assertIn("本次更新", message)
        self.assertIn("配置迁移说明", message)
        self.assertIn("迁移提醒测试", message)

        migration_only = main.format_startup_notice(False, ["仅配置提醒。"])
        self.assertNotIn("本次更新", migration_only)
        self.assertIn("仅配置提醒", migration_only)

    def test_lino_branding_assets_and_about_links_are_available(self):
        for relative_path in (
            "assets/linote-logo.png",
            "assets/linote-icon.png",
            "assets/linote-logo.ico",
        ):
            self.assertTrue(main.app_resource_path(relative_path).is_file(), relative_path)
        self.assertEqual("v26", main.APP_VERSION)
        self.assertEqual("https://www.linote.top/", main.LINO_HOME_URL)
        self.assertEqual("https://github.com/linote99", main.GITHUB_PROFILE_URL)
        self.assertEqual("https://github.com/linote99/wps-sheet-automation", main.PROJECT_REPOSITORY_URL)
        source = Path(main.__file__).read_text(encoding="utf-8")
        project_dir = Path(main.__file__).resolve().parent
        installer_path = project_dir / "install_and_build_main_only.bat"
        if not installer_path.is_file():
            installer_path = installer_path.with_name("install_and_build.bat")
        installer = installer_path.read_text(encoding="utf-8").lower()
        installer_definition = (project_dir / "WpsSheetAutomation.iss").read_text(
            encoding="utf-8").lower()
        self.assertIn('version_label = "不再提醒版本更新" if has_migration_notices else "不再提示"', source)
        self.assertIn('else "不再提醒"', source)
        self.assertIn("dont_show_configuration_notice", source)
        self.assertIn("root.after(300, app._show_startup_notices)", source)
        self.assertNotIn("root.after(200, self._show_configuration_migration_notices", source)
        self.assertIn('--icon="assets\\linote-logo.ico"', installer)
        self.assertIn('--add-data "assets\\linote-logo.png;assets"', installer)
        self.assertIn("iscc.exe", installer)
        self.assertIn("wpssheetautomation.iss", installer)
        self.assertIn("wpssheetautomation_setup_v26.exe", installer)
        self.assertIn("outputbasefilename=wpssheetautomation_setup_v26", installer_definition)
        self.assertIn('source: "dist\\wpssheetautomation\\*"', installer_definition)
        self.assertNotIn("issuer_private_key.pem", installer)
        self.assertNotIn("WpsSheetAutomation_KeyGen", installer)
