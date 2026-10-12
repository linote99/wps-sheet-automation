from __future__ import annotations

import json
import os
import posixpath
import queue
import re
import sys
import tempfile
import threading
import time
import traceback
import zipfile
from contextlib import contextmanager
from copy import copy, deepcopy
from datetime import date, datetime, time as datetime_time, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path, PurePosixPath
from typing import Any, Callable
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET
from xml.parsers import expat

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel, WINDOWS_EPOCH
from workflow import (DEFAULT_WORKFLOW, MODULE_LABELS, add_workflow_module, list_profiles,
                      load_profile, matching_profiles, move_workflow_module,
                      remove_workflow_module, save_profile, validate_workflow)

APP_NAME = "WPS 表格多源汇总处理器"
CONFIG_PATH = Path.home() / ".wps_sheet_automation.json"
PROFILE_DIR = Path.home() / ".wps_sheet_automation_profiles"
DEFAULT_WINDOW_WIDTH = 1500
DEFAULT_WINDOW_HEIGHT = 960
MIN_WINDOW_WIDTH = 1024
MIN_WINDOW_HEIGHT = 680


def _window_size_path() -> Path:
    if os.name == "nt":
        base = Path(os.environ.get("APPDATA") or Path.home() / "AppData" / "Roaming")
        return base / "WpsSheetAutomation" / "window_size.json"
    xdg_config = os.environ.get("XDG_CONFIG_HOME", "")
    base = Path(xdg_config) if xdg_config and Path(xdg_config).is_absolute() else Path.home() / ".config"
    return base / "wps_sheet_automation" / "window_size.json"


WINDOW_SIZE_PATH = _window_size_path()
CONFIG_VERSION = 8
LEGACY_FORMULA_ROW_KEYS = ("formula_seed_row", "formula_template_row", "formula_row", "seed_row")
MAX_PROCESSING_CELLS = 500_000
MAX_REGION_XLSX_CELLS = 100_000
MAX_REGION_XLSX_CHARS = 2_000_000
MAX_TEXT_CELLS = 100_000
MAX_TEXT_CHARS = 2_000_000
TEXT_COPY_CHUNK_ROWS = 500
MAX_XLSX_MEMBERS = 20_000
MAX_XLSX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_XLSX_WORKSHEET_XML_BYTES = 64 * 1024 * 1024
MAX_XLSX_METADATA_XML_BYTES = 8 * 1024 * 1024
DEFAULT_SETTINGS: dict[str, Any] = {
    "config_version": CONFIG_VERSION,
    "config_name": "",
    "template_prefix_length": 5,
    "workflow": list(DEFAULT_WORKFLOW),
    "template_path": "",
    "output_dir": "",
    "output_filename": "处理结果.xlsx",
    "output_filename_custom": None,
    "export_region_pngs": False,
    "copy_start_col": "A",
    "copy_end_col": "",
    "pre_mappings": [],
    "post_mappings": [],
    "image_outputs": [],
    "text_outputs": [],
    "TextToColumns": {"enabled": False, "rules": []},
    "Deduplicate": {"enabled": False, "rules": []},
    "FormatConvert": {"enabled": False, "rules": []},
    # Kept for migration of first-version configurations.
    "source_start_row": 2,
    "jobs": [
        {"template_sheet": "1当天派单", "header_rows": 2, "start_row": 3, "formula_range": "EB:EN",
         "sources": [{"path": "", "source_sheet": "", "copy_start_row": 2, "copy_end_row": None}],
         "mappings": [{"result_col": "EJ", "target_col": "DG", "write_as": "value"}]},
        {"template_sheet": "2当天归档", "header_rows": 2, "start_row": 3, "formula_range": "EB:EU",
         "sources": [{"path": "", "source_sheet": "", "copy_start_row": 2, "copy_end_row": None}],
         "mappings": [{"result_col": "EQ", "target_col": "DG", "write_as": "value"}]},
        {"template_sheet": "3装移全业务在途", "header_rows": 3, "start_row": 4, "formula_range": "EB:FV",
         "sources": [{"path": "", "source_sheet": "", "copy_start_row": 2, "copy_end_row": None}],
         "mappings": [{"result_col": "FQ", "target_col": "DG", "write_as": "value"}]},
    ],
}


class ProcessingError(RuntimeError):
    """An actionable error safe to show to the user."""


def window_size_limits(screen_width: Any = None, screen_height: Any = None) -> tuple[int, int, int, int]:
    """Return minimum and maximum window dimensions, safely bounded to the display."""
    def positive_dimension(value: Any, fallback: int) -> int:
        try:
            dimension = int(value)
        except (TypeError, ValueError, OverflowError):
            return fallback
        return dimension if dimension > 0 else fallback

    screen_width = positive_dimension(screen_width, DEFAULT_WINDOW_WIDTH + 64)
    screen_height = positive_dimension(screen_height, DEFAULT_WINDOW_HEIGHT + 96)
    max_width = screen_width - 48 if screen_width >= 400 else screen_width
    max_height = screen_height - 80 if screen_height >= 320 else screen_height
    max_width, max_height = max(1, max_width), max(1, max_height)
    return (min(MIN_WINDOW_WIDTH, max_width), min(MIN_WINDOW_HEIGHT, max_height), max_width, max_height)


def normalize_window_size(saved: Any, screen_width: Any = None,
                          screen_height: Any = None) -> tuple[int, int]:
    """Validate saved width/height and clamp them to the current display."""
    min_width, min_height, max_width, max_height = window_size_limits(screen_width, screen_height)

    def parse_dimension(value: Any, minimum: int) -> int | None:
        if isinstance(value, bool):
            return None
        if isinstance(value, int):
            dimension = value
        elif isinstance(value, str) and re.fullmatch(r"[0-9]{1,5}", value.strip()):
            dimension = int(value.strip())
        else:
            return None
        return dimension if minimum <= dimension <= 16_384 else None

    width = height = None
    if isinstance(saved, dict):
        width = parse_dimension(saved.get("width"), 320)
        height = parse_dimension(saved.get("height"), 240)
    if width is None or height is None:
        width, height = DEFAULT_WINDOW_WIDTH, DEFAULT_WINDOW_HEIGHT
    return (min(max(width, min_width), max_width), min(max(height, min_height), max_height))


def load_window_size(path: Path | str | None = None, screen_width: Any = None,
                     screen_height: Any = None) -> tuple[int, int]:
    """Load only the prior window dimensions; malformed or missing state uses defaults."""
    state_path = Path(path) if path is not None else WINDOW_SIZE_PATH
    try:
        saved = json.loads(state_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        saved = None
    return normalize_window_size(saved, screen_width, screen_height)


def save_window_size(width: Any, height: Any, path: Path | str | None = None,
                     screen_width: Any = None, screen_height: Any = None) -> bool:
    """Persist only validated dimensions, atomically; return False on I/O failure."""
    state_path = Path(path) if path is not None else WINDOW_SIZE_PATH
    size = normalize_window_size({"width": width, "height": height}, screen_width, screen_height)
    temporary = state_path.with_name(state_path.name + ".tmp")
    try:
        state_path.parent.mkdir(parents=True, exist_ok=True)
        temporary.write_text(json.dumps({"width": size[0], "height": size[1]}), encoding="utf-8")
        os.replace(temporary, state_path)
        return True
    except OSError:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        return False


def column_number(value: str) -> int:
    value = str(value).strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", value):
        raise ProcessingError(f"列号“{value or '(空)'}”无效；请使用 A、DG、XFD 这样的 Excel 列名。")
    try:
        number = column_index_from_string(value)
    except ValueError as exc:
        raise ProcessingError(f"列号“{value}”超出 Excel 列范围。") from exc
    if number > 16384:
        raise ProcessingError(f"列号“{value}”超出 Excel 最大列 XFD。")
    return number


def parse_column_range(value: str) -> tuple[int, int]:
    parts = str(value).strip().upper().split(":")
    if len(parts) != 2:
        raise ProcessingError(f"列范围“{value}”格式错误；请按 A:C 或 EB:EN 填写。")
    start, end = (column_number(part) for part in parts)
    if start > end:
        raise ProcessingError(f"列范围“{value}”的起始列不能在末列之后。")
    return start, end


def ranges_overlap(first: tuple[int, int], second: tuple[int, int]) -> bool:
    return first[0] <= second[1] and second[0] <= first[1]


def _positive_row(value: Any, label: str, allow_zero: bool = False) -> int:
    try:
        result = int(value)
    except (TypeError, ValueError) as exc:
        raise ProcessingError(f"{label}必须是{'非负' if allow_zero else '正'}整数。") from exc
    lower = 0 if allow_zero else 1
    if result < lower or result > 1_048_576:
        raise ProcessingError(f"{label}必须介于 {lower} 和 1048576。")
    return result


def _formula_template_row(job: dict[str, Any]) -> int:
    """Formula templates always come from the row immediately above pasted data."""
    start_row = _positive_row(job.get("start_row"), "粘贴起始行")
    if start_row <= 1:
        raise ProcessingError("公式填充要求粘贴起始行至少为第 2 行；公式模板行自动取粘贴起始行的上一行。")
    return start_row - 1


def configuration_migration_notes(raw: Any) -> list[str]:
    """Explain legacy row semantics without trusting any saved seed-row field."""
    if not isinstance(raw, dict):
        return []
    try:
        version = int(raw.get("config_version", 1))
    except (TypeError, ValueError):
        version = 1
    jobs = raw.get("jobs", [])
    if not isinstance(jobs, list):
        jobs = []
    notes: list[str] = []
    if version < 7:
        notes.append(
            f"旧配置版本 v{version} 已按 v7 规则迁移：每个目标表的 start_row 保持为粘贴起始行，"
            "公式模板行自动取 start_row - 1；旧版本中将起始行当作公式种子行的含义不再沿用。"
        )
    if "open_region_exports" in raw and "export_region_pngs" not in raw:
        notes.append(
            "旧的“自动打开区域 XLSX”选项已迁移为 PNG 导出开关：启用后会导出 PNG，"
            "并在图片成功生成后删除本轮对应的独立区域 XLSX。"
        )
    for index, job in enumerate(jobs, 1):
        if not isinstance(job, dict):
            continue
        legacy_values = [(key, job[key]) for key in LEGACY_FORMULA_ROW_KEYS if key in job]
        if not legacy_values:
            continue
        try:
            template_row = _formula_template_row(job)
        except ProcessingError:
            template_row = "无法计算（粘贴起始行必须至少为 2）"
        old_fields = "、".join(f"{key}={value!r}" for key, value in legacy_values)
        notes.append(
            f"第 {index} 个目标表的旧公式行字段（{old_fields}）不会用于选取公式；"
            f"实际公式模板行按 start_row - 1 自动确定为 {template_row}。"
        )
    return notes


def _validate_sheet_name(value: Any, label: str) -> str:
    name = str(value or "").strip()
    if not name:
        raise ProcessingError(f"{label}工作表名不能为空。")
    if len(name) > 31 or any(char in name for char in "[]:*?/\\"):
        raise ProcessingError(f"{label}工作表名“{name}”不符合 Excel 命名规则。")
    return name


def _parse_area(value: Any, label: str) -> tuple[int, int, int, int]:
    """Parse one worksheet-local A1 rectangle as (min_col, min_row, max_col, max_row)."""
    text = str(value or "").strip().upper()
    if not re.fullmatch(r"\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?", text):
        raise ProcessingError(f"{label}区域“{text or '(空)'}”无效；请填写 A1 或 A1:C10 这样的单一单元格区域。")
    try:
        bounds = range_boundaries(text)
    except ValueError as exc:
        raise ProcessingError(f"{label}区域“{text}”无效。") from exc
    if any(value is None for value in bounds):
        raise ProcessingError(f"{label}区域“{text}”必须包含起止行和列。")
    min_col, min_row, max_col, max_row = bounds
    if min_col < 1 or min_row < 1 or max_col > 16_384 or max_row > 1_048_576:
        raise ProcessingError(f"{label}区域“{text}”超出 Excel 工作表边界。")
    if min_col > max_col or min_row > max_row:
        raise ProcessingError(f"{label}区域“{text}”的起点不能在终点之后。")
    return min_col, min_row, max_col, max_row


def _parse_selection_details(value: Any, label: str = "区域") -> tuple[tuple[int, int, int, int], str]:
    """Parse the shared UI/JSON selector, normalizing reversed endpoint rectangles."""
    text = str(value or "").strip().upper()
    coordinate = r"\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}"

    def point(token: str) -> tuple[int, int]:
        token = token.strip().upper()
        if not re.fullmatch(coordinate, token):
            raise ProcessingError(f"{label}端点“{token or '(空)'}”无效；请使用如 A5、D1 或 B3,B3 的 Excel 坐标。")
        try:
            col, row = range_boundaries(token)[0:2]
            column_number(get_column_letter(col))
            _positive_row(row, f"{label}行号")
        except (ValueError, TypeError) as exc:
            raise ProcessingError(f"{label}端点“{token}”超出 Excel 范围（XFD 列、1048576 行）。") from exc
        return col, row

    if "," in text:
        endpoints = text.split(",")
        if len(endpoints) != 2:
            raise ProcessingError(f"{label}“{text}”只允许两个逗号分隔的端点，例如 B3,D8 或 B3,B3。")
        first, second = point(endpoints[0]), point(endpoints[1])
        if first == second:
            return (first[0], first[1], first[0], first[1]), "cell"
        return (min(first[0], second[0]), min(first[1], second[1]),
                max(first[0], second[0]), max(first[1], second[1])), "rectangle"

    if ":" in text:
        endpoints = text.split(":")
        if len(endpoints) != 2:
            raise ProcessingError(f"{label}“{text}”格式无效；请使用两个端点。")
        first, second = point(endpoints[0]), point(endpoints[1])
        if first == second:
            return (first[0], first[1], first[0], first[1]), "cell"
        return (min(first[0], second[0]), min(first[1], second[1]),
                max(first[0], second[0]), max(first[1], second[1])), "rectangle"

    col, row = point(text)
    if col == 1 and row == 1:
        return (1, 1, 16_384, 1_048_576), "sheet"
    if col == 1:
        return (1, row, 16_384, row), "row"
    if row == 1:
        return (col, 1, col, 1_048_576), "column"
    return (col, row, col, row), "cell"


def _parse_selected_area(value: Any, label: str = "区域") -> tuple[int, int, int, int]:
    return _parse_selection_details(value, label)[0]


def _area_cell_count(bounds: tuple[int, int, int, int]) -> int:
    return (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1)


def _check_processing_area_size(bounds: tuple[int, int, int, int], label: str) -> None:
    count = _area_cell_count(bounds)
    if count > MAX_PROCESSING_CELLS:
        raise ProcessingError(f"{label}区域包含 {count:,} 个单元格，超过处理上限 {MAX_PROCESSING_CELLS:,}；请缩小范围。")


def _check_text_area_size(bounds: tuple[int, int, int, int], label: str) -> None:
    count = _area_cell_count(bounds)
    if count > MAX_TEXT_CELLS:
        raise ProcessingError(f"{label}包含 {count:,} 个单元格，超过文本预览/复制上限 {MAX_TEXT_CELLS:,}；请缩小范围。")


def _check_region_xlsx_area_size(bounds: tuple[int, int, int, int], label: str) -> None:
    cells = _area_cell_count(bounds)
    if cells > MAX_REGION_XLSX_CELLS:
        raise ProcessingError(
            f"{label}区域包含 {cells:,} 个单元格，超过独立 XLSX 区域输出上限 "
            f"{MAX_REGION_XLSX_CELLS:,}；请缩小范围。"
        )


def _area_address(bounds: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = bounds
    first = f"{get_column_letter(min_col)}{min_row}"
    last = f"{get_column_letter(max_col)}{max_row}"
    return first if first == last else f"{first}:{last}"


def _expand_bounds_for_merged_cells(worksheet, bounds: tuple[int, int, int, int],
                                    label: str) -> tuple[int, int, int, int]:
    """Expand a selected area until it completely contains every intersecting merged cell."""
    expanded = bounds
    while True:
        previous = expanded
        for merged in worksheet.merged_cells.ranges:
            merged_bounds = (merged.min_col, merged.min_row, merged.max_col, merged.max_row)
            if not _areas_overlap(expanded, merged_bounds):
                continue
            expanded = (min(expanded[0], merged_bounds[0]), min(expanded[1], merged_bounds[1]),
                        max(expanded[2], merged_bounds[2]), max(expanded[3], merged_bounds[3]))
            _check_region_xlsx_area_size(
                expanded, f"{label}完整包含合并单元格后的 {_area_address(expanded)}"
            )
        if expanded == previous:
            return expanded


def _areas_overlap(first: tuple[int, int, int, int], second: tuple[int, int, int, int]) -> bool:
    return (first[0] <= second[2] and second[0] <= first[2]
            and first[1] <= second[3] and second[1] <= first[3])


def _template_output_filename(template_path: Any) -> str:
    """Return the template basename in a platform-independent way, including its suffix."""
    raw = str(template_path or "").strip().replace("\\", "/")
    return raw.rsplit("/", 1)[-1] if raw else ""


def _filename_is_custom(raw_settings: dict[str, Any], filename: str) -> bool:
    marker = raw_settings.get("output_filename_custom")
    if marker is not None:
        if isinstance(marker, str):
            return marker.strip().lower() in ("1", "true", "yes", "on")
        return bool(marker)
    # Older configurations did not record whether the field was auto-generated.
    # Treat the former generic default and the current template basename as automatic.
    template_name = _template_output_filename(raw_settings.get("template_path"))
    return bool(filename and filename not in ("处理结果.xlsx", template_name))


def normalize_settings(raw: dict[str, Any]) -> dict[str, Any]:
    """Return the current schema; migrate the fixed three-source v1 schema in memory."""
    if not isinstance(raw, dict):
        raise ProcessingError("配置根节点必须是 JSON 对象。")
    settings = deepcopy(DEFAULT_SETTINGS)
    settings.update({k: deepcopy(v) for k, v in raw.items() if k not in ("jobs", "_migration_notices")})
    raw_jobs = raw.get("jobs", settings["jobs"])
    if not isinstance(raw_jobs, list):
        raise ProcessingError("配置中的 jobs 必须是数组。")
    old_sources = raw.get("sources")
    is_legacy = any(isinstance(job, dict) and "sources" not in job for job in raw_jobs)
    if is_legacy and isinstance(old_sources, list):
        migrated = []
        global_start = raw.get("source_start_row", 2)
        for idx, old in enumerate(raw_jobs):
            if not isinstance(old, dict):
                migrated.append(old)
                continue
            job = deepcopy(old)
            src_path = old_sources[idx] if idx < len(old_sources) else ""
            job["sources"] = [{"path": src_path, "source_sheet": job.pop("source_sheet", ""),
                               "copy_start_row": global_start, "copy_end_row": None}]
            result_col = job.pop("result_col", "")
            target_col = job.pop("target_col", "")
            job["mappings"] = [{"result_col": result_col, "target_col": target_col, "write_as": "value"}]
            migrated.append(job)
        raw_jobs = migrated
    normalized_jobs = []
    for original in raw_jobs:
        if not isinstance(original, dict):
            normalized_jobs.append(original)
            continue
        job = deepcopy(original)
        for legacy_key in LEGACY_FORMULA_ROW_KEYS:
            job.pop(legacy_key, None)
        if "mappings" not in job:
            result_col = job.pop("result_col", "")
            target_col = job.pop("target_col", "")
            job["mappings"] = [{"result_col": result_col, "target_col": target_col, "write_as": "value"}]
        source_items = job.get("sources", [])
        if isinstance(source_items, list):
            converted = []
            for source in source_items:
                if isinstance(source, str):
                    converted.append({"path": source, "source_sheet": "", "copy_start_row": raw.get("source_start_row", 2), "copy_end_row": None})
                elif isinstance(source, dict):
                    item = {"path": "", "source_sheet": "", "copy_start_row": raw.get("source_start_row", 2), "copy_end_row": None}
                    item.update(source)
                    converted.append(item)
                else:
                    converted.append(source)
            job["sources"] = converted
        normalized_jobs.append(job)
    settings["jobs"] = normalized_jobs
    legacy_open_exports = settings.pop("open_region_exports", None)
    if "export_region_pngs" in raw:
        settings["export_region_pngs"] = bool(settings.get("export_region_pngs", False))
    else:
        settings["export_region_pngs"] = bool(legacy_open_exports) if legacy_open_exports is not None else False
    for key in ("pre_mappings", "post_mappings", "image_outputs", "text_outputs"):
        if settings.get(key) is None:
            settings[key] = []
    for key in ("TextToColumns", "Deduplicate", "FormatConvert"):
        operation = settings.get(key)
        if operation is None:
            settings[key] = {"enabled": False, "rules": []}
        elif isinstance(operation, dict):
            settings[key] = {"enabled": operation.get("enabled", False),
                             "rules": deepcopy(operation.get("rules", []))}
    if settings.get("workflow") is None:
        settings["workflow"] = list(DEFAULT_WORKFLOW)
    if not isinstance(settings.get("workflow"), list):
        raise ProcessingError("配置中的 workflow 必须是模块 ID 数组。")
    try:
        prefix_length = int(settings.get("template_prefix_length", 5))
    except (TypeError, ValueError) as exc:
        raise ProcessingError("模板匹配前缀长度 N 必须是整数。") from exc
    if not 1 <= prefix_length <= 120:
        raise ProcessingError("模板匹配前缀长度 N 必须介于 1 和 120。")
    settings["template_prefix_length"] = prefix_length
    settings["config_name"] = str(settings.get("config_name") or "").strip()
    filename = str(settings.get("output_filename") or "").strip()
    is_custom_filename = _filename_is_custom(raw, filename)
    settings["output_filename_custom"] = is_custom_filename
    if not is_custom_filename:
        template_filename = _template_output_filename(settings.get("template_path"))
        if template_filename:
            settings["output_filename"] = template_filename
    settings.pop("sources", None)
    settings["config_version"] = CONFIG_VERSION
    return settings


def validate_static_settings(raw_settings: dict[str, Any]) -> None:
    settings = normalize_settings(raw_settings)
    try:
        validate_workflow(settings)
    except ValueError as exc:
        raise ProcessingError(str(exc)) from exc
    jobs = settings.get("jobs")
    if not isinstance(jobs, list) or not jobs:
        raise ProcessingError("至少需要配置一个目标工作表。")
    copy_start = column_number(settings.get("copy_start_col", "A"))
    copy_end_value = str(settings.get("copy_end_col", "")).strip()
    copy_range = None
    if copy_end_value:
        copy_end = column_number(copy_end_value)
        if copy_start > copy_end:
            raise ProcessingError("数据复制列范围的起始列不能在末列之后。")
        copy_range = (copy_start, copy_end)
    names: list[str] = []
    for idx, job in enumerate(jobs, start=1):
        if not isinstance(job, dict):
            raise ProcessingError(f"第 {idx} 个目标工作表配置必须是对象。")
        for key in ("template_sheet", "header_rows", "start_row", "formula_range", "sources", "mappings"):
            if key not in job:
                raise ProcessingError(f"第 {idx} 个目标工作表配置缺少“{key}”。")
        sheet_name = str(job["template_sheet"]).strip()
        if not sheet_name:
            raise ProcessingError(f"第 {idx} 个目标工作表名称不能为空。")
        if len(sheet_name) > 31 or any(char in sheet_name for char in "[]:*?/\\"):
            raise ProcessingError(f"第 {idx} 个目标工作表名称不符合 Excel 工作表命名规则。")
        names.append(sheet_name)
        header_rows = _positive_row(job["header_rows"], f"第 {idx} 项保留标题行数", allow_zero=True)
        start_row = _positive_row(job["start_row"], f"第 {idx} 项目标粘贴起始行")
        if start_row <= 1:
            raise ProcessingError(f"第 {idx} 项目标粘贴起始行必须至少为第 2 行；公式模板行自动取其上一行。")
        if start_row <= header_rows:
            raise ProcessingError(f"第 {idx} 项目标起始行必须大于保留标题行数。")
        formula_range = parse_column_range(job["formula_range"])
        sources = job["sources"]
        if not isinstance(sources, list) or not sources:
            raise ProcessingError(f"第 {idx} 个目标表至少需要关联一份源文件。")
        for source_idx, source in enumerate(sources, 1):
            if not isinstance(source, dict):
                raise ProcessingError(f"第 {idx} 个目标表的第 {source_idx} 份源文件配置无效。")
            if "path" not in source or "source_sheet" not in source:
                raise ProcessingError(f"第 {idx} 个目标表的第 {source_idx} 份源文件需要 path 和 source_sheet。")
            src_start = _positive_row(source.get("copy_start_row"), f"第 {idx} 个目标表第 {source_idx} 份源文件的复制起始行")
            src_end = source.get("copy_end_row")
            if src_end not in (None, ""):
                src_end = _positive_row(src_end, f"第 {idx} 个目标表第 {source_idx} 份源文件的复制结束行")
                if src_end < src_start:
                    raise ProcessingError(f"第 {idx} 个目标表第 {source_idx} 份源文件结束行不得早于起始行。")
        mappings = job["mappings"]
        if not isinstance(mappings, list) or not mappings:
            raise ProcessingError(f"第 {idx} 个目标表至少需要一条结果列映射。")
        targets: dict[int, int] = {}
        for map_idx, mapping in enumerate(mappings, 1):
            if not isinstance(mapping, dict):
                raise ProcessingError(f"第 {idx} 个目标表的第 {map_idx} 条映射无效。")
            try:
                result_col = column_number(mapping.get("result_col", ""))
                target_col = column_number(mapping.get("target_col", ""))
            except ProcessingError as exc:
                raise ProcessingError(f"第 {idx} 个目标表第 {map_idx} 条映射：{exc}") from exc
            if not (formula_range[0] <= result_col <= formula_range[1]):
                raise ProcessingError(f"第 {idx} 项结果列 {mapping['result_col']} 不在公式范围 {job['formula_range']} 内。")
            if formula_range[0] <= target_col <= formula_range[1]:
                raise ProcessingError(f"第 {idx} 个目标表第 {map_idx} 条映射的目标列 {mapping['target_col']} 位于公式范围 {job['formula_range']} 内；写回会覆盖必须保留的公式列，请改用公式范围外的目标列。")
            if target_col in targets:
                raise ProcessingError(f"第 {idx} 个目标表第 {targets[target_col]} 条和第 {map_idx} 条映射都写入目标列 {mapping['target_col']}，会互相覆盖；请为每条映射指定不同目标列。")
            targets[target_col] = map_idx
            mode = str(mapping.get("write_as", "value")).strip().lower()
            if mode not in ("value", "text"):
                raise ProcessingError(f"第 {idx} 项映射写入类型只能是 value（值）或 text（文本）。")
            if copy_range is not None:
                if ranges_overlap(copy_range, formula_range):
                    raise ProcessingError(f"第 {idx} 项数据复制列范围与公式区域 {job['formula_range']} 重叠。")
    if len(set(names)) != len(names):
        raise ProcessingError("目标工作表名称必须互不相同。")

    for key, label in (("pre_mappings", "前映射"), ("post_mappings", "后映射")):
        rules = settings.get(key, [])
        if not isinstance(rules, list):
            raise ProcessingError(f"配置中的 {key} 必须是数组。")
        destinations: dict[str, list[tuple[int, tuple[int, int, int, int]]]] = {}
        for rule_idx, rule in enumerate(rules, 1):
            if not isinstance(rule, dict):
                raise ProcessingError(f"第 {rule_idx} 条{label}规则必须是对象。")
            for field in ("source_sheet", "source_range", "target_sheet", "target_range"):
                if field not in rule:
                    raise ProcessingError(f"第 {rule_idx} 条{label}规则缺少“{field}”。")
            _validate_sheet_name(rule["source_sheet"], f"第 {rule_idx} 条{label}源")
            target_sheet = _validate_sheet_name(rule["target_sheet"], f"第 {rule_idx} 条{label}目标")
            source_bounds = _parse_area(rule["source_range"], f"第 {rule_idx} 条{label}源")
            target_bounds = _parse_area(rule["target_range"], f"第 {rule_idx} 条{label}目标")
            source_shape = (source_bounds[2] - source_bounds[0], source_bounds[3] - source_bounds[1])
            target_shape = (target_bounds[2] - target_bounds[0], target_bounds[3] - target_bounds[1])
            if source_shape != target_shape:
                raise ProcessingError(f"第 {rule_idx} 条{label}规则的源/目标区域行列数必须相同。")
            for previous_idx, previous in destinations.setdefault(target_sheet, []):
                if _areas_overlap(previous, target_bounds):
                    raise ProcessingError(f"第 {previous_idx} 条和第 {rule_idx} 条{label}规则的目标区域重叠：{target_sheet}!{_area_address(previous)} 与 {_area_address(target_bounds)}。")
            destinations[target_sheet].append((rule_idx, target_bounds))

    for key, label in (("image_outputs", "区域 XLSX 输出"), ("text_outputs", "文本输出")):
        outputs = settings.get(key, [])
        if not isinstance(outputs, list):
            raise ProcessingError(f"配置中的 {key} 必须是数组。")
        for output_idx, output in enumerate(outputs, 1):
            if not isinstance(output, dict):
                raise ProcessingError(f"第 {output_idx} 项{label}配置必须是对象。")
            if "sheet" not in output or "range" not in output:
                raise ProcessingError(f"第 {output_idx} 项{label}配置需要 sheet 和 range。")
            _validate_sheet_name(output["sheet"], f"第 {output_idx} 项{label}")
            bounds, selection_kind = _parse_selection_details(output["range"], f"第 {output_idx} 项{label}")
            if key == "image_outputs" and selection_kind not in ("sheet", "row", "column"):
                _check_region_xlsx_area_size(bounds, f"第 {output_idx} 项区域 XLSX 输出")
            elif key == "text_outputs" and selection_kind not in ("sheet", "row", "column"):
                _check_text_area_size(bounds, f"第 {output_idx} 项文本输出")

    for key, label in (("TextToColumns", "文本分列"), ("Deduplicate", "去重"), ("FormatConvert", "格式转换")):
        operation = settings.get(key)
        if not isinstance(operation, dict) or not isinstance(operation.get("enabled", False), bool):
            raise ProcessingError(f"{label}必须包含布尔 enabled 开关和 rules 数组。")
        rules = operation.get("rules", [])
        if not isinstance(rules, list):
            raise ProcessingError(f"{label}的 rules 必须是数组。")
        for idx, rule in enumerate(rules, 1):
            if not isinstance(rule, dict):
                raise ProcessingError(f"第 {idx} 条{label}规则必须是对象。")
            _validate_sheet_name(rule.get("sheet"), f"第 {idx} 条{label}")
            bounds = _parse_selected_area(rule.get("range"), f"第 {idx} 条{label}")
            _check_processing_area_size(bounds, f"第 {idx} 条{label}")
            if key == "TextToColumns":
                if bounds[0] != bounds[2]:
                    raise ProcessingError(f"第 {idx} 条文本分列规则的源范围必须是单列。")
                if not str(rule.get("delimiter", "")):
                    raise ProcessingError(f"第 {idx} 条文本分列规则必须指定非空分隔符。")
                column_number(rule.get("output_start_col", ""))
                formats = rule.get("column_formats", ["text"])
                if not isinstance(formats, list) or not formats:
                    raise ProcessingError(f"第 {idx} 条文本分列规则须配置每个结果列的格式。")
                if any(str(fmt).lower() not in ("text", "number", "date") for fmt in formats):
                    raise ProcessingError(f"第 {idx} 条文本分列规则的列格式只能是 text、number 或 date。")
            elif key == "Deduplicate":
                if not isinstance(rule.get("has_header", False), bool):
                    raise ProcessingError(f"第 {idx} 条去重规则的 has_header 必须是布尔值。")
                key_mode = str(rule.get("key_mode", "row")).lower()
                if key_mode not in ("row", "columns"):
                    raise ProcessingError(f"第 {idx} 条去重规则的 key_mode 只能是 row 或 columns。")
                keys = rule.get("key_columns", [])
                if key_mode == "columns":
                    if not isinstance(keys, list) or not keys:
                        raise ProcessingError(f"第 {idx} 条去重规则须指定至少一个关键列。")
                    seen = set()
                    for key_col in keys:
                        col = column_number(key_col)
                        if not bounds[0] <= col <= bounds[2]:
                            raise ProcessingError(f"第 {idx} 条去重规则关键列 {key_col} 不在所选区域内。")
                        if col in seen:
                            raise ProcessingError(f"第 {idx} 条去重规则重复指定关键列 {key_col}。")
                        seen.add(col)
            else:
                if str(rule.get("mode", "")).lower() not in ("text", "number", "date"):
                    raise ProcessingError(f"第 {idx} 条格式转换规则 mode 只能是 text、number 或 date。")

    _validate_region_module_dependencies(settings)


def _validate_region_module_dependencies(settings: dict[str, Any]) -> None:
    workflow = settings.get("workflow", DEFAULT_WORKFLOW)
    positions = {module_id: index for index, module_id in enumerate(workflow)}
    split_operation = settings.get("TextToColumns", {})
    if not split_operation.get("enabled", False) or "TextToColumns" not in positions:
        return
    for later_key in ("Deduplicate", "FormatConvert"):
        operation = settings.get(later_key, {})
        if not operation.get("enabled", False) or later_key not in positions:
            continue
        for split_idx, split_rule in enumerate(split_operation.get("rules", []), 1):
            source = _parse_selected_area(split_rule["range"], f"第 {split_idx} 条文本分列")
            formats = split_rule.get("column_formats", ["text"])
            start_col = column_number(split_rule["output_start_col"])
            target = (start_col, source[1], start_col + len(formats) - 1, source[3])
            for rule_idx, rule in enumerate(operation.get("rules", []), 1):
                if str(rule["sheet"]).strip() != str(split_rule["sheet"]).strip():
                    continue
                other = _parse_selected_area(rule["range"], f"第 {rule_idx} 条{MODULE_LABELS[later_key]}")
                if _areas_overlap(source, other) or _areas_overlap(target, other):
                    if positions["TextToColumns"] >= positions[later_key]:
                        raise ProcessingError(
                            f"工作流顺序无效：“文本分列”必须位于“{MODULE_LABELS[later_key]}”之前；"
                            f"第 {split_idx} 条分列规则的输入/输出区域与第 {rule_idx} 条{MODULE_LABELS[later_key]}规则重叠，"
                            "否则分列输入会先被修改，或分列目标会先被占用。"
                        )


def _resource_path(relative: str) -> Path:
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return root / relative


def load_settings(path: Path | str | None = None) -> dict[str, Any]:
    settings = deepcopy(DEFAULT_SETTINGS)
    source_path = Path(path) if path is not None else CONFIG_PATH
    if source_path.exists():
        try:
            saved = json.loads(source_path.read_text(encoding="utf-8"))
            settings = normalize_settings(saved)
            notices = configuration_migration_notes(saved)
            if notices:
                settings["_migration_notices"] = notices
        except Exception as exc:
            raise ProcessingError(f"读取配置失败：{source_path}\n{exc}") from exc
    return settings


def save_settings(settings: dict[str, Any], path: Path | str | None = None) -> None:
    normalized = normalize_settings(settings)
    validate_static_settings(normalized)
    target = Path(path) if path is not None else CONFIG_PATH
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")


def _check_xlsx_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise ProcessingError(f"{label}不存在或不是文件：\n{path}")
    if path.suffix.lower() != ".xlsx":
        raise ProcessingError(f"{label}必须是 .xlsx 文件（不接受含宏工作簿或其他格式）：\n{path}")
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = archive.namelist()
            if any(n.startswith("xl/externalLinks/") for n in names):
                raise ProcessingError(f"{label}包含外部工作簿链接；为避免访问外部内容，已拒绝处理：\n{path}")
            if "xl/vbaProject.bin" in names:
                raise ProcessingError(f"{label}包含 VBA 宏；程序只处理无宏 .xlsx 文件：\n{path}")
    except ProcessingError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise ProcessingError(f"{label}不是可读取的有效 XLSX 工作簿：\n{path}\n{exc}") from exc


def _zip_member_name_is_safe(name: str) -> bool:
    if not name or "\x00" in name or "\\" in name or name.startswith("/"):
        return False
    relative = name[:-1] if name.endswith("/") else name
    if not relative:
        return False
    path = PurePosixPath(relative)
    if path.is_absolute() or any(part in (".", "..") for part in path.parts):
        return False
    if path.parts and ":" in path.parts[0]:
        return False
    return True


def _xml_root(data: bytes, description: str):
    compact = data.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in compact or b"<!ENTITY" in compact:
        raise ValueError(f"{description}包含不支持的 DTD 或实体声明。")
    return ET.fromstring(data)


def _local_xml_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _read_bounded_zip_xml(archive: zipfile.ZipFile, name: str, limit: int) -> bytes:
    info = archive.getinfo(name)
    if info.file_size > limit:
        raise ValueError(f"工作簿 XML 部件 {name} 超过安全读取上限。")
    with archive.open(info, "r") as source:
        data = source.read(limit + 1)
    if len(data) > limit or len(data) != info.file_size:
        raise ValueError(f"工作簿 XML 部件 {name} 的实际大小异常。")
    return data


def _resolve_internal_relationship(base_part: str, target: str) -> str:
    split = urlsplit(target)
    if split.scheme or split.netloc or split.query or split.fragment:
        raise ValueError("工作表关系目标不是安全的包内路径。")
    decoded = unquote(split.path)
    if not decoded or "\\" in decoded or "\x00" in decoded:
        raise ValueError("工作表关系目标路径无效。")
    if decoded.startswith("/"):
        resolved = posixpath.normpath(decoded.lstrip("/"))
    else:
        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(base_part), decoded))
    if resolved in ("", ".", "..") or resolved.startswith("../") or not _zip_member_name_is_safe(resolved):
        raise ValueError("工作表关系目标越出了 XLSX 包目录。")
    return resolved


def _worksheet_parts(archive: zipfile.ZipFile, member_names: set[str]) -> dict[str, str]:
    workbook_part = "xl/workbook.xml"
    relationships_part = "xl/_rels/workbook.xml.rels"
    if workbook_part not in member_names or relationships_part not in member_names:
        raise ValueError("XLSX 缺少工作簿清单或工作表关系部件。")
    workbook_root = _xml_root(_read_bounded_zip_xml(archive, workbook_part, MAX_XLSX_METADATA_XML_BYTES), workbook_part)
    relationships_root = _xml_root(
        _read_bounded_zip_xml(archive, relationships_part, MAX_XLSX_METADATA_XML_BYTES), relationships_part
    )
    relationship_by_id = {}
    for relation in relationships_root.iter():
        if _local_xml_name(relation.tag) != "Relationship":
            continue
        relation_id = relation.attrib.get("Id")
        if relation_id:
            relationship_by_id[relation_id] = relation.attrib

    result: dict[str, str] = {}
    relationship_id_attributes = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
        "{http://purl.oclc.org/ooxml/officeDocument/relationships}id",
    )
    for sheet in workbook_root.iter():
        if _local_xml_name(sheet.tag) != "sheet":
            continue
        relation_id = next((sheet.attrib[key] for key in relationship_id_attributes if key in sheet.attrib), None)
        if not relation_id:
            continue
        relation = relationship_by_id.get(relation_id)
        if relation is None or not relation.get("Type", "").endswith("/worksheet"):
            continue
        if relation.get("TargetMode", "").lower() == "external":
            raise ValueError("工作表关系指向外部目标，无法安全检查。")
        part = _resolve_internal_relationship(workbook_part, relation.get("Target", ""))
        if part not in member_names:
            raise ValueError(f"工作表 XML 部件不存在：{part}")
        sheet_name = sheet.attrib.get("name", "(未命名工作表)")
        result[part] = sheet_name
    return result


def _valid_excel_a1_area(reference: str | None) -> bool:
    if not isinstance(reference, str) or not reference:
        return False
    cell = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}"
    column = r"\$?[A-Za-z]{1,3}"
    row = r"\$?[1-9][0-9]{0,6}"
    patterns = (rf"{cell}(?::{cell})?", rf"{column}:{column}", rf"{row}:{row}")
    if not any(re.fullmatch(pattern, reference) for pattern in patterns):
        return False
    try:
        min_col, min_row, max_col, max_row = range_boundaries(reference)
    except (TypeError, ValueError):
        return False
    return (
        (min_col is None or 1 <= min_col <= 16_384)
        and (max_col is None or 1 <= max_col <= 16_384)
        and (min_row is None or 1 <= min_row <= 1_048_576)
        and (max_row is None or 1 <= max_row <= 1_048_576)
    )


def _xml_start_tag_end(data: bytes, start: int) -> int:
    quote = None
    for index in range(start, len(data)):
        byte = data[index]
        if quote is not None:
            if byte == quote:
                quote = None
        elif byte in (ord("\""), ord("'")):
            quote = byte
        elif byte == ord(">"):
            return index + 1
    raise ValueError("无法定位 sortState XML 开始标签的结束位置。")


def _normalized_sort_state_ref(reference: str | None, sheet_name: str) -> str | None:
    if reference is None:
        raise ValueError(f"工作表“{sheet_name}”的 sortState 缺少 ref；为保护排序信息，未删除该排序定义。")
    row_range = re.fullmatch(r"([1-9][0-9]*):([1-9][0-9]*)", reference)
    if row_range:
        start_text, end_text = row_range.groups()
        if len(start_text) > 7 or len(end_text) > 7:
            raise ValueError(
                f"工作表“{sheet_name}”的 sortState ref {reference!r} 超出 Excel 行号范围；"
                "为保护排序信息，未删除该排序定义。"
            )
        start_row, end_row = int(start_text), int(end_text)
        if start_row > 1_048_576 or end_row > 1_048_576:
            raise ValueError(
                f"工作表“{sheet_name}”的 sortState ref {reference!r} 超出 Excel 行号范围；"
                "为保护排序信息，未删除该排序定义。"
            )
        if start_row > end_row:
            raise ValueError(
                f"工作表“{sheet_name}”的 sortState ref {reference!r} 起始行大于结束行；"
                "为保护排序信息，未删除该排序定义。"
            )
        return f"A{start_row}:XFD{end_row}"
    if _valid_excel_a1_area(reference):
        return None
    raise ValueError(
        f"工作表“{sheet_name}”的 sortState ref {reference!r} 不是有效 Excel 排序区域；"
        "无法安全规范化，为保护排序信息，未删除该排序定义。"
    )


def _preprocess_worksheet_xml(data: bytes, sheet_name: str) -> tuple[bytes, list[str], list[tuple[str, str]]]:
    if b"\x00" in data[:256]:
        raise ValueError("工作表 XML 不是可安全定位修改的 UTF-8 编码。")
    compact = data.replace(b"\x00", b"").upper()
    if b"<!DOCTYPE" in compact or b"<!ENTITY" in compact:
        raise ValueError("工作表 XML 包含不支持的 DTD 或实体声明。")

    parser = expat.ParserCreate(namespace_separator="}")
    stack: list[dict[str, Any]] = []
    invalid_filter_spans: list[tuple[int, int, str]] = []
    sort_ref_edits: list[tuple[int, int, bytes, str, str]] = []
    normalized_refs: list[tuple[str, str]] = []

    def start_element(name: str, attributes: dict[str, str]) -> None:
        depth = len(stack)
        node = {"name": name, "start": parser.CurrentByteIndex, "invalid_ref": None}
        if depth == 0:
            root_namespace = name.rsplit("}", 1)[0] if "}" in name else ""
            node["is_worksheet"] = _local_xml_name(name) == "worksheet"
            node["namespace"] = root_namespace
        elif stack[0].get("is_worksheet"):
            namespace = name.rsplit("}", 1)[0] if "}" in name else ""
            if namespace == stack[0].get("namespace") and _local_xml_name(name) == "autoFilter":
                reference = attributes.get("ref")
                if not _valid_excel_a1_area(reference):
                    node["invalid_ref"] = "<缺失>" if reference is None else reference
            if namespace == stack[0].get("namespace") and _local_xml_name(name) == "sortState":
                reference = attributes.get("ref")
                replacement = _normalized_sort_state_ref(reference, sheet_name)
                if replacement is not None:
                    start = int(node["start"])
                    tag_end = _xml_start_tag_end(data, start)
                    start_tag = data[start:tag_end]
                    matches = list(re.finditer(
                        rb"(?:^|[ \t\r\n])ref[ \t\r\n]*=[ \t\r\n]*(?:\"([^\"]*)\"|'([^']*)')",
                        start_tag,
                    ))
                    if len(matches) != 1:
                        raise ValueError(
                            f"工作表“{sheet_name}”的 sortState ref {reference!r} 无法在 XML 中安全定位；"
                            "为保护排序信息，未删除该排序定义。"
                        )
                    match = matches[0]
                    ref_group = 1 if match.group(1) is not None else 2
                    raw_ref = match.group(ref_group)
                    if raw_ref != reference.encode("ascii"):
                        raise ValueError(
                            f"工作表“{sheet_name}”的 sortState ref {reference!r} 使用了无法安全改写的 XML 编码；"
                            "为保护排序信息，未删除该排序定义。"
                        )
                    sort_ref_edits.append((
                        start + match.start(ref_group), start + match.end(ref_group),
                        replacement.encode("ascii"), reference, replacement,
                    ))
                    normalized_refs.append((reference, replacement))
        stack.append(node)

    def end_element(name: str) -> None:
        if not stack:
            raise ValueError("工作表 XML 元素结构异常。")
        node = stack.pop()
        if node["name"] != name or node["invalid_ref"] is None:
            return
        start = int(node["start"])
        event_position = parser.CurrentByteIndex
        if data[max(start, event_position - 2):event_position] == b"/>":
            end = event_position
        else:
            closing = data.find(b">", event_position)
            if closing < 0:
                raise ValueError("无法定位无效自动筛选定义的 XML 结束位置。")
            end = closing + 1
        invalid_filter_spans.append((start, end, str(node["invalid_ref"])))

    parser.StartElementHandler = start_element
    parser.EndElementHandler = end_element
    parser.ExternalEntityRefHandler = lambda *_args: 0
    parser.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
    parser.Parse(data, True)
    for edit_start, edit_end, _replacement, old_ref, _new_ref in sort_ref_edits:
        for filter_start, filter_end, _invalid_ref in invalid_filter_spans:
            if filter_start <= edit_start < edit_end <= filter_end:
                raise ValueError(
                    f"工作表“{sheet_name}”的 sortState ref {old_ref!r} 位于将被临时移除的无效 autoFilter 内；"
                    "为保护排序信息，已拒绝该文件且未删除排序定义。"
                )
    edits: list[tuple[int, int, bytes, str | None]] = [
        (start, end, b"", reference) for start, end, reference in invalid_filter_spans
    ]
    edits.extend((start, end, replacement, None) for start, end, replacement, _old, _new in sort_ref_edits)
    if not edits:
        return data, [], []
    edits.sort(key=lambda item: item[0])
    output: list[bytes] = []
    invalid_filter_refs: list[str] = []
    position = 0
    for start, end, replacement, reference in edits:
        if start < position or end < start or end > len(data):
            raise ValueError("工作表 XML 修改区域重叠或越界；未生成修改副本。")
        output.append(data[position:start])
        output.append(replacement)
        position = end
        if reference is not None:
            shown_reference = reference if len(reference) <= 200 else reference[:197] + "..."
            invalid_filter_refs.append(shown_reference)
    output.append(data[position:])
    return b"".join(output), invalid_filter_refs, normalized_refs


class _TemporaryWorkbookCopy:
    """Path-like temporary copy that owns every openpyxl workbook loaded from it."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.workbooks = []

    def __fspath__(self) -> str:
        return os.fspath(self.path)

    def load_workbook(self, *args, **kwargs):
        workbook = load_workbook(self.path, *args, **kwargs)
        self.workbooks.append(workbook)
        return workbook

    def close_workbooks(self) -> list[str]:
        errors = []
        for workbook in reversed(self.workbooks):
            try:
                workbook.close()
            except Exception as exc:
                errors.append(str(exc))
        self.workbooks.clear()
        return errors


_TEMP_COPY_CLEANUP_RETRY_DELAYS = (0.05, 0.1, 0.2)


def _remove_temporary_workbook_copy(path: Path) -> OSError | None:
    for attempt in range(len(_TEMP_COPY_CLEANUP_RETRY_DELAYS) + 1):
        try:
            path.unlink(missing_ok=True)
            return None
        except OSError as exc:
            transient_lock = isinstance(exc, PermissionError) or getattr(exc, "winerror", None) in (32, 33)
            if not transient_lock or attempt == len(_TEMP_COPY_CLEANUP_RETRY_DELAYS):
                return exc
            time.sleep(_TEMP_COPY_CLEANUP_RETRY_DELAYS[attempt])
    return None


def _log_temporary_copy_warning(log: Callable[[str], None], message: str) -> None:
    try:
        log(message)
    except Exception:
        print(message, file=sys.stderr)


@contextmanager
def _temporary_workbook_copy(path: Path, label: str, log: Callable[[str], None]):
    """Prepare a unique XLSX copy and close all tracked workbooks before deleting it."""
    try:
        fd, temp_name = tempfile.mkstemp(prefix="wps_sheet_autofilter_", suffix=".xlsx")
    except OSError as exc:
        raise ProcessingError(
            f"无法为{label}“{path.name}”创建安全的临时读取副本。请在 WPS 中打开该文件的副本，"
            f"移除或重建筛选后另存为 XLSX 再试；原文件未修改。\n原因：{exc}"
        ) from exc
    os.close(fd)
    temp_path = Path(temp_name)
    temporary_copy = _TemporaryWorkbookCopy(temp_path)
    operation_error = None
    try:
        try:
            with zipfile.ZipFile(path, "r") as source:
                infos = source.infolist()
                if len(infos) > MAX_XLSX_MEMBERS:
                    raise ValueError(f"ZIP 成员数超过 {MAX_XLSX_MEMBERS} 项安全上限。")
                total_declared = sum(info.file_size for info in infos)
                if total_declared > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ValueError("ZIP 解压总大小超过 512 MiB 安全上限。")
                names: set[str] = set()
                for info in infos:
                    if not _zip_member_name_is_safe(info.filename):
                        raise ValueError(f"ZIP 成员路径不安全：{info.filename!r}")
                    if info.filename in names:
                        raise ValueError(f"ZIP 中存在重名成员：{info.filename}")
                    names.add(info.filename)
                    if info.flag_bits & 0x1:
                        raise ValueError(f"ZIP 成员已加密，无法安全读取：{info.filename}")
                sheet_part_names = _worksheet_parts(source, names)
                changed_filters: list[tuple[str, str]] = []
                changed_sort_states: list[tuple[str, str, str]] = []
                actual_total = 0
                with zipfile.ZipFile(temp_path, "w") as destination:
                    destination.comment = source.comment
                    for info in infos:
                        output_info = copy(info)
                        if info.is_dir():
                            destination.writestr(output_info, b"")
                            continue
                        if info.filename in sheet_part_names:
                            if info.file_size > MAX_XLSX_WORKSHEET_XML_BYTES:
                                raise ValueError(f"工作表 XML 部件 {info.filename} 超过 64 MiB 安全上限。")
                            raw = _read_bounded_zip_xml(source, info.filename, MAX_XLSX_WORKSHEET_XML_BYTES)
                            actual_total += len(raw)
                            cleaned, invalid_refs, normalized_refs = _preprocess_worksheet_xml(
                                raw, sheet_part_names[info.filename]
                            )
                            destination.writestr(output_info, cleaned)
                            for invalid_ref in invalid_refs:
                                changed_filters.append((sheet_part_names[info.filename], invalid_ref))
                            for old_ref, new_ref in normalized_refs:
                                changed_sort_states.append((sheet_part_names[info.filename], old_ref, new_ref))
                        else:
                            member_total = 0
                            with source.open(info, "r") as member_in, destination.open(output_info, "w") as member_out:
                                while True:
                                    chunk = member_in.read(1024 * 1024)
                                    if not chunk:
                                        break
                                    member_total += len(chunk)
                                    actual_total += len(chunk)
                                    if actual_total > MAX_XLSX_UNCOMPRESSED_BYTES or member_total > info.file_size:
                                        raise ValueError("ZIP 实际解压数据超过安全上限或与目录大小不符。")
                                    member_out.write(chunk)
                            if member_total != info.file_size:
                                raise ValueError(f"ZIP 成员 {info.filename} 的实际大小与目录记录不符。")
                        if actual_total > MAX_XLSX_UNCOMPRESSED_BYTES:
                            raise ValueError("ZIP 实际解压总大小超过 512 MiB 安全上限。")
            for sheet_name, invalid_ref in changed_filters:
                log(
                    f"{label}“{path.name}”工作表“{sheet_name}”的自动筛选区域 {invalid_ref!r} 无效（文件：{path}）；"
                    "仅在临时读取副本中移除此筛选定义，原文件未更改。"
                )
            for sheet_name, old_ref, new_ref in changed_sort_states:
                log(
                    f"{label}“{path.name}”工作表“{sheet_name}”的 sortState ref {old_ref!r} 已在临时读取副本中"
                    f"规范为 {new_ref!r}（文件：{path}）；原文件未更改。"
                )
        except ProcessingError:
            raise
        except Exception as exc:
            raise ProcessingError(
                f"无法安全检查或准备{label}“{path.name}”的临时副本。请在 WPS 中打开该文件的副本，"
                f"检查并修正工作表排序/筛选区域引用后另存为 XLSX 再试；不要直接删除排序定义，原文件未修改。\n原因：{exc}"
            ) from exc
        yield temporary_copy
    except BaseException as exc:
        operation_error = exc
        raise
    finally:
        close_errors = temporary_copy.close_workbooks()
        if close_errors:
            _log_temporary_copy_warning(
                log,
                f"警告：{label}“{path.name}”临时读取副本中的 openpyxl 工作簿关闭时发生错误："
                f"{temp_path}；请确认该文件未被其他程序占用。原因：{'；'.join(close_errors)}",
            )
        cleanup_error = _remove_temporary_workbook_copy(temp_path)
        if cleanup_error is not None:
            state = "读取/处理已成功，但临时读取副本无法清理" if operation_error is None else "读取/处理本身已失败，且临时读取副本也无法清理"
            _log_temporary_copy_warning(
                log,
                f"警告：{state}：{temp_path}。请关闭占用该文件的程序后手动删除；原工作簿未修改。"
                f"清理原因：{cleanup_error}",
            )


def _find_external_formula(formula: str) -> bool:
    if re.search(r"\[[^\]]+\.(?:xlsx?|xlsm|xlsb)\]", formula, flags=re.IGNORECASE):
        return True
    return bool(re.search(r"\b(?:WEBSERVICE|FILTERXML|RTD|CALL|REGISTER\.ID|EXEC)\s*\(", formula, flags=re.IGNORECASE))


def _normalize_path(value: str, label: str) -> Path:
    raw = str(value or "").strip()
    if not raw:
        raise ProcessingError(f"请先选择{label}。")
    return Path(raw).expanduser().resolve()


def _first_missing_source_location(settings: dict[str, Any]) -> tuple[int, int] | None:
    """Return the zero-based job/source position of the first blank source path."""
    for job_idx, job in enumerate(settings.get("jobs", [])):
        for source_idx, source in enumerate(job.get("sources", [])):
            if not str(source.get("path") or "").strip():
                return job_idx, source_idx
    return None


def _missing_source_path_message(settings: dict[str, Any], location: tuple[int, int]) -> str:
    job_idx, source_idx = location
    sheet_name = settings["jobs"][job_idx].get("template_sheet", f"第 {job_idx + 1} 个目标表")
    return (f"目标表“{sheet_name}”的第 {source_idx + 1} 份源工作簿尚未配置文件路径。"
            "请在源文件列表中编辑该项并选择或填写 .xlsx 路径。")


def _get_formula_templates(workbook, job: dict[str, Any], index: int, start_row: int) -> tuple[int, int, dict[int, str]]:
    sheet_name = str(job["template_sheet"]).strip()
    if sheet_name not in workbook.sheetnames:
        raise ProcessingError(f"模板中找不到第 {index} 个目标工作表“{sheet_name}”。")
    ws = workbook[sheet_name]
    formula_row = _formula_template_row({**job, "start_row": start_row})
    col_start, col_end = parse_column_range(job["formula_range"])
    formulas: dict[int, str] = {}
    missing: list[str] = []
    for col in range(col_start, col_end + 1):
        value = ws.cell(row=formula_row, column=col).value
        if not (isinstance(value, str) and value.startswith("=")):
            missing.append(get_column_letter(col))
            continue
        if _find_external_formula(value):
            raise ProcessingError(f"模板公式 {sheet_name}!{get_column_letter(col)}{formula_row} 包含外部引用或外部访问函数，已拒绝重算。")
        formulas[col] = value
    if missing:
        columns = "、".join(missing)
        raise ProcessingError(
            f"模板“{sheet_name}”的公式模板行是第 {formula_row} 行（由粘贴起始行 {start_row} 自动确定），"
            f"以下公式列缺少公式：{columns}。请补齐该模板行中的公式。"
        )
    return col_start, col_end, formulas


def _cell_has_data(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or value.strip() != "")


def _last_nonempty_source_row(ws) -> tuple[int, int]:
    """Find the last row with content and the sheet's used-column boundary."""
    max_row = max(ws.max_row or 0, 1)
    max_col = max(ws.max_column or 0, 1)
    last_nonempty = 0
    for row_number, cells in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col), start=1
    ):
        if any(_cell_has_data(cell.value) for cell in cells):
            last_nonempty = row_number
    return last_nonempty, max_col


def _read_source_records(path: Path, sheet_name: str, start_row: int, col_start: int, col_end: int,
                         log: Callable[[str], None], end_row: int | None = None) -> list[list[Any]]:
    _check_xlsx_file(path, f"源工作簿“{path.name}”")
    with _temporary_workbook_copy(path, "源工作簿", log) as load_path:
        return _read_source_records_from_copy(load_path, path, sheet_name, start_row, col_start, col_end, log, end_row)


def _read_source_records_from_copy(load_path: _TemporaryWorkbookCopy, path: Path, sheet_name: str, start_row: int,
                                   col_start: int, col_end: int, log: Callable[[str], None],
                                   end_row: int | None = None) -> list[list[Any]]:
    try:
        values_book = load_path.load_workbook(read_only=True, data_only=True, keep_links=False)
        formula_book = load_path.load_workbook(read_only=True, data_only=False, keep_links=False)
        requested_sheet = str(sheet_name or "").strip()
        if not values_book.sheetnames or not formula_book.sheetnames:
            raise ProcessingError(f"源工作簿“{path.name}”没有可读取的工作表。")
        if not requested_sheet:
            resolved_sheet = formula_book.sheetnames[0]
            log(f"{path.name} 未填写源工作表，默认使用第一个工作表“{resolved_sheet}”。")
        else:
            resolved_sheet = requested_sheet
            if resolved_sheet not in values_book.sheetnames or resolved_sheet not in formula_book.sheetnames:
                raise ProcessingError(f"源工作簿“{path.name}”中找不到配置的工作表“{resolved_sheet}”。")
        val_ws, raw_ws = values_book[resolved_sheet], formula_book[resolved_sheet]
        actual_last_row, used_max_col = _last_nonempty_source_row(raw_ws)
        if end_row is not None and end_row < start_row:
            raise ProcessingError(f"源工作簿“{path.name}”的结束行早于复制起始行。")
        if actual_last_row < start_row:
            raise ProcessingError(f"源工作簿“{path.name}”工作表“{resolved_sheet}”在第 {start_row} 行及之后没有可复制数据。")

        if end_row is not None:
            last_row = min(end_row, actual_last_row)
            log(f"{path.name} 使用显式复制结束行 {end_row}；本次读取上限为第 {last_row} 行。")
        else:
            last_row = actual_last_row
            boundary_row = None
            for row_number, cells in enumerate(
                raw_ws.iter_rows(min_row=start_row, max_row=actual_last_row, min_col=1, max_col=used_max_col),
                start=start_row,
            ):
                if not any(_cell_has_data(cell.value) for cell in cells):
                    boundary_row = row_number
                    last_row = row_number - 1
                    break
            if boundary_row is not None:
                log(f"{path.name} 在第 {boundary_row} 行遇到整行空白边界（检查全部 {used_max_col} 个已使用列）；采用第 {last_row} 行为末行，边界后的数据不纳入此源。")
            else:
                log(f"{path.name} 从第 {start_row} 行至最后非空行 {actual_last_row} 未遇到整行空白边界；采用最后非空行作为末行。")
        if last_row < start_row:
            raise ProcessingError(f"源工作簿“{path.name}”工作表“{resolved_sheet}”从第 {start_row} 行起遇到空白边界，没有可复制数据。")

        records: list[list[Any]] = []
        blank_formula_cache_count = 0
        raw_rows = raw_ws.iter_rows(min_row=start_row, max_row=last_row, min_col=col_start, max_col=col_end)
        value_rows = val_ws.iter_rows(min_row=start_row, max_row=last_row, min_col=col_start, max_col=col_end)
        for raw_cells, value_cells in zip(raw_rows, value_rows):
            raw = [cell.value for cell in raw_cells]
            vals = [cell.value for cell in value_cells]
            if not any(_cell_has_data(value) for value in raw):
                continue
            for source_value, cached_value in zip(raw, vals):
                if isinstance(source_value, str) and source_value.startswith("=") and cached_value is None:
                    blank_formula_cache_count += 1
            records.append(vals)
        if not records:
            interval = f"第 {start_row} 至 {last_row} 行" if end_row is not None else f"第 {start_row} 行起至第 {last_row} 行"
            raise ProcessingError(f"源工作簿“{path.name}”工作表“{resolved_sheet}”的{interval}在复制列范围内没有数据。")
        if blank_formula_cache_count:
            log(f"提示：{path.name} 有 {blank_formula_cache_count} 个源公式单元格没有已保存的计算缓存；这些单元格将按空值复制，不会执行源公式。")
        return records
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"读取源工作簿“{path.name}”失败：{exc}") from exc


def _validate_inputs(raw_settings: dict[str, Any]) -> tuple[Path, Path]:
    settings = normalize_settings(raw_settings)
    validate_static_settings(settings)
    missing_source = _first_missing_source_location(settings)
    if missing_source is not None:
        raise ProcessingError(_missing_source_path_message(settings, missing_source))
    template_path = _normalize_path(settings.get("template_path", ""), "模板工作簿")
    _check_xlsx_file(template_path, "模板工作簿")
    all_inputs = {template_path}
    for job_idx, job in enumerate(settings["jobs"], 1):
        if not job["sources"]:
            raise ProcessingError(f"目标表“{job['template_sheet']}”至少需要一份源文件。")
        for source_idx, source in enumerate(job["sources"], 1):
            path = _normalize_path(source.get("path", ""), f"目标表“{job['template_sheet']}”第 {source_idx} 份源工作簿")
            # Blank source_sheet is intentionally resolved by the reader to the first sheet.
            if path == template_path:
                raise ProcessingError("模板工作簿不能同时作为源工作簿。")
            _check_xlsx_file(path, f"目标表“{job['template_sheet']}”第 {source_idx} 份源工作簿")
            source["path"] = str(path)
            all_inputs.add(path)
    output_dir = _normalize_path(settings.get("output_dir", ""), "输出目录")
    if not output_dir.is_dir():
        raise ProcessingError(f"输出目录不存在：\n{output_dir}")
    filename = str(settings.get("output_filename", "")).strip()
    if not filename:
        raise ProcessingError("输出文件名不能为空。")
    if Path(filename).name != filename or "/" in filename or "\\" in filename:
        raise ProcessingError("输出文件名只能填写文件名，不能包含目录路径。")
    if Path(filename).suffix and Path(filename).suffix.lower() != ".xlsx":
        raise ProcessingError("输出文件只支持 .xlsx 扩展名；请勿使用 .xlsm 或其他格式。")
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    final_path = (output_dir / filename).resolve()
    if os.path.normcase(str(final_path)) == os.path.normcase(str(template_path)):
        raise ProcessingError(
            f"输出文件名与模板同名且位于同一目录，不能覆盖模板：\n{template_path}\n"
            "请改名或选择其他输出目录。"
        )
    if final_path.exists():
        raise ProcessingError(f"输出文件已存在，为防止覆盖任何文件已停止：\n{final_path}\n请更改输出文件名或选择其他输出目录。")
    if final_path in all_inputs:
        raise ProcessingError("输出路径不能与模板或任何源工作簿相同。")
    if not str(settings.get("copy_end_col", "")).strip():
        raise ProcessingError("请填写数据复制末列；程序不会猜测需要复制的源列范围。")
    column_number(settings["copy_start_col"])
    column_number(settings["copy_end_col"])
    return template_path, final_path


def _validate_template_area_sheets(workbook, settings: dict[str, Any]) -> None:
    for key, label in (("pre_mappings", "前映射"), ("post_mappings", "后映射")):
        for idx, rule in enumerate(settings.get(key, []), 1):
            for field in ("source_sheet", "target_sheet"):
                sheet_name = str(rule[field]).strip()
                if sheet_name not in workbook.sheetnames:
                    raise ProcessingError(f"第 {idx} 条{label}规则引用了模板中不存在的工作表“{sheet_name}”。")
    for key, label in (("image_outputs", "区域 XLSX 输出"), ("text_outputs", "文本输出")):
        for idx, output in enumerate(settings.get(key, []), 1):
            sheet_name = str(output["sheet"]).strip()
            if sheet_name not in workbook.sheetnames:
                raise ProcessingError(f"第 {idx} 项{label}配置引用了模板中不存在的工作表“{sheet_name}”。")
    for key, label in (("TextToColumns", "文本分列"), ("Deduplicate", "去重"), ("FormatConvert", "格式转换")):
        for idx, rule in enumerate(settings.get(key, {}).get("rules", []), 1):
            sheet_name = str(rule["sheet"]).strip()
            if sheet_name not in workbook.sheetnames:
                raise ProcessingError(f"第 {idx} 条{label}规则引用了模板中不存在的工作表“{sheet_name}”。")


def _capture_pre_mapping_values(workbook, settings: dict[str, Any]):
    """Read all pre-mapping sources from the untouched template, before any target cleanup."""
    _validate_template_area_sheets(workbook, settings)
    captured = []
    for idx, rule in enumerate(settings.get("pre_mappings", []), 1):
        bounds = _parse_area(rule["source_range"], f"第 {idx} 条前映射源")
        min_col, min_row, max_col, max_row = bounds
        worksheet = workbook[str(rule["source_sheet"]).strip()]
        values = [[worksheet.cell(row=row, column=col).value
                   for col in range(min_col, max_col + 1)]
                  for row in range(min_row, max_row + 1)]
        captured.append((deepcopy(rule), bounds, values))
    return captured


def _apply_pre_mapping_values(workbook, captured, log: Callable[[str], None]) -> None:
    for idx, (rule, source_bounds, values) in enumerate(captured, 1):
        target_bounds = _parse_area(rule["target_range"], f"第 {idx} 条前映射目标")
        source_min_col, source_min_row, _source_max_col, _source_max_row = source_bounds
        target_min_col, target_min_row, _target_max_col, _target_max_row = target_bounds
        source_sheet = workbook[str(rule["source_sheet"]).strip()]
        target_sheet = workbook[str(rule["target_sheet"]).strip()]
        for row_offset, row_values in enumerate(values):
            for col_offset, value in enumerate(row_values):
                target = target_sheet.cell(row=target_min_row + row_offset, column=target_min_col + col_offset)
                if isinstance(target, MergedCell):
                    raise ProcessingError(f"前映射目标 {rule['target_sheet']}!{target.coordinate} 是合并单元格的非左上角，不能写入。")
                if isinstance(value, str) and value.startswith("="):
                    source_coord = f"{get_column_letter(source_min_col + col_offset)}{source_min_row + row_offset}"
                    try:
                        value = Translator(value, origin=source_coord).translate_formula(target.coordinate)
                    except Exception as exc:
                        raise ProcessingError(f"无法将前映射公式 {rule['source_sheet']}!{source_coord} 平移到 {rule['target_sheet']}!{target.coordinate}：{exc}") from exc
                target.value = value
        log(f"已应用前映射 {rule['source_sheet']}!{_area_address(source_bounds)} → {rule['target_sheet']}!{_area_address(target_bounds)}。")


def _validate_pre_mapping_conflicts(settings: dict[str, Any], source_records: list[list[list[Any]]]) -> None:
    """Reject pre-map targets that overlap any main-data/formula/post-map write area."""
    copy_start = column_number(settings["copy_start_col"])
    copy_end = column_number(settings["copy_end_col"])
    protected: dict[str, list[tuple[str, tuple[int, int, int, int]]]] = {}
    for job, records in zip(settings["jobs"], source_records):
        start_row = int(job["start_row"])
        end_row = start_row + len(records) - 1
        sheet = str(job["template_sheet"]).strip()
        if end_row < start_row:
            continue
        protected.setdefault(sheet, []).append(("主数据复制", (copy_start, start_row, copy_end, end_row)))
        formula_start, formula_end = parse_column_range(job["formula_range"])
        protected[sheet].append(("公式填充", (formula_start, start_row, formula_end, end_row)))
        template_row = _formula_template_row(job)
        protected[sheet].append(("公式模板", (1, template_row, 16_384, template_row)))
        for mapping in job["mappings"]:
            col = column_number(mapping["target_col"])
            protected[sheet].append(("结果映射", (col, start_row, col, end_row)))
    for rule in settings.get("post_mappings", []):
        bounds = _parse_area(rule["target_range"], "后映射目标")
        protected.setdefault(str(rule["target_sheet"]).strip(), []).append(("后映射", bounds))

    for idx, rule in enumerate(settings.get("pre_mappings", []), 1):
        target_sheet = str(rule["target_sheet"]).strip()
        target_bounds = _parse_area(rule["target_range"], f"第 {idx} 条前映射目标")
        for write_kind, write_bounds in protected.get(target_sheet, []):
            if _areas_overlap(target_bounds, write_bounds):
                raise ProcessingError(
                    f"第 {idx} 条前映射目标 {target_sheet}!{_area_address(target_bounds)} 与{write_kind}区域 "
                    f"{_area_address(write_bounds)} 冲突；已在修改临时工作簿前停止，避免覆盖。请调整前映射目标区域。"
                )


def _create_staging_workbook(template_path: Path, raw_settings: dict[str, Any], source_records: list[list[list[Any]]],
                             staging_path: Path, log: Callable[[str], None]) -> list[tuple[str, int, int, list[dict[str, str]]]]:
    _check_xlsx_file(template_path, "模板工作簿")
    with _temporary_workbook_copy(template_path, "模板工作簿", log) as load_path:
        return _create_staging_workbook_from_copy(template_path, load_path, raw_settings, source_records, staging_path, log)


def _create_staging_workbook_from_copy(template_path: Path, load_path: _TemporaryWorkbookCopy, raw_settings: dict[str, Any],
                                       source_records: list[list[list[Any]]], staging_path: Path,
                                       log: Callable[[str], None]) -> list[tuple[str, int, int, list[dict[str, str]]]]:
    settings = normalize_settings(raw_settings)
    try:
        wb = load_path.load_workbook(keep_links=False, data_only=False)
        captured_pre_mappings = _capture_pre_mapping_values(wb, settings)
        for ws in wb.worksheets:
            for cell in ws._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith("=") and _find_external_formula(cell.value):
                    raise ProcessingError(f"模板“{ws.title}”单元格 {cell.coordinate} 含外部引用或外部访问函数；为避免外部访问已停止。")
        copy_start = column_number(settings["copy_start_col"])
        prepared = []
        for idx, job in enumerate(settings["jobs"], 1):
            start_row = int(job["start_row"])
            c1, c2, formulas = _get_formula_templates(wb, job, idx, start_row)
            prepared.append((job, wb[str(job["template_sheet"]).strip()], start_row, int(job["header_rows"]), c1, c2, formulas))
        if len(source_records) != len(prepared):
            raise ProcessingError("内部数据与目标工作表数量不一致。")
        job_ranges = []
        for idx, ((job, ws, start_row, header_rows, _c1, _c2, formulas), records) in enumerate(zip(prepared, source_records), 1):
            if not records:
                raise ProcessingError(f"目标表“{job['template_sheet']}”没有可写入的数据。")
            end_row = start_row + len(records) - 1
            if end_row > 1_048_576:
                raise ProcessingError(f"第 {idx} 个工作表目标末行超出 Excel 最大行数。")
            job_ranges.append((job, ws, start_row, end_row, header_rows, formulas, records))

        # Cleanup is deliberately a separate pass so every pre-map source was captured first.
        for job, ws, start_row, _end_row, header_rows, formulas, _records in job_ranges:
            template_row = _formula_template_row(job)
            cleared = 0
            for cell in list(ws._cells.values()):
                if isinstance(cell, MergedCell) or cell.row <= header_rows:
                    continue
                if cell.row == template_row:
                    continue
                if cell.value is not None or cell.hyperlink is not None or cell.comment is not None:
                    cell.value = None
                    cell.hyperlink = None
                    cell.comment = None
                    cleared += 1
            log(f"已清理“{job['template_sheet']}”第 {header_rows + 1} 行起的旧内容（{cleared} 个单元格），保留第 {template_row} 行公式模板及其行内容。")

        _apply_pre_mapping_values(wb, captured_pre_mappings, log)

        formula_jobs: list[tuple[str, int, int, list[dict[str, str]]]] = []
        for job, ws, start_row, end_row, _header_rows, formulas, records in job_ranges:
            for offset, values in enumerate(records):
                target_row = start_row + offset
                for col_offset, value in enumerate(values):
                    ws.cell(row=target_row, column=copy_start + col_offset).value = value
            for col, formula in formulas.items():
                origin = f"{get_column_letter(col)}{_formula_template_row(job)}"
                for target_row in range(start_row, end_row + 1):
                    target = f"{get_column_letter(col)}{target_row}"
                    try:
                        ws[target] = Translator(formula, origin=origin).translate_formula(target)
                    except Exception as exc:
                        raise ProcessingError(f"无法将模板公式 {job['template_sheet']}!{origin} 填充到 {target}：{exc}") from exc
            mappings = [{"result_col": str(m["result_col"]).strip().upper(),
                         "target_col": str(m["target_col"]).strip().upper(),
                         "write_as": str(m.get("write_as", "value")).strip().lower()} for m in job["mappings"]]
            formula_jobs.append((str(job["template_sheet"]).strip(), start_row, end_row, mappings))
            log(f"已准备“{job['template_sheet']}”：按配置顺序追加 {len(records)} 行，填充第 {start_row} 至 {end_row} 行公式。")
        wb.save(staging_path)
        return formula_jobs
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"创建中间工作簿失败：{exc}") from exc


def _write_mapped_values(destination, values: list[Any], write_as: str) -> None:
    """Write one column to a COM-like range; text mode forces string values and text format."""
    mode = str(write_as).strip().lower()
    if mode == "text":
        values = [None if value is None else str(value) for value in values]
        try:
            destination.NumberFormat = "@"
        except Exception:
            pass
    destination.Value2 = tuple((value,) for value in values)


def _com_value_matrix(value, rows: int, cols: int) -> list[list[Any]]:
    if rows == 1 and cols == 1:
        if isinstance(value, (tuple, list)):
            first = value[0]
            return [list(first) if isinstance(first, (tuple, list)) else [first]]
        return [[value]]
    if rows == 1:
        if value and isinstance(value[0], (tuple, list)):
            return [list(value[0])]
        return [list(value)]
    if cols == 1:
        return [[row[0]] if isinstance(row, (tuple, list)) else [row] for row in value]
    return [list(row) for row in value]


def _recalculate_and_wait(app, log: Callable[[str], None], phase: str) -> None:
    calculate_errors = []
    for method_name in ("CalculateFullRebuild", "CalculateFull", "Calculate"):
        try:
            getattr(app, method_name)()
            calculate_errors = []
            log(f"WPS 已执行 {method_name}（{phase}）。")
            break
        except Exception as exc:
            calculate_errors.append(f"{method_name}: {exc}")
    if calculate_errors:
        raise ProcessingError("WPS 无法触发公式重算：\n" + "\n".join(calculate_errors))
    deadline = time.time() + 180
    while time.time() < deadline:
        try:
            if int(app.CalculationState) == 0:
                return
        except Exception:
            return
        time.sleep(0.25)
    raise ProcessingError("等待 WPS 完成公式重算超过 180 秒，已中止以避免输出不完整文件。")


def _read_com_area(com_range, bounds: tuple[int, int, int, int]) -> list[list[Any]]:
    rows = bounds[3] - bounds[1] + 1
    cols = bounds[2] - bounds[0] + 1
    return _com_value_matrix(com_range.Value2, rows, cols)


def _write_com_area(com_range, matrix: list[list[Any]]) -> None:
    com_range.Value2 = tuple(tuple(row) for row in matrix)


def _flatten_com_values(value):
    if isinstance(value, (tuple, list)):
        for item in value:
            yield from _flatten_com_values(item)
    else:
        yield value


def _assert_no_formulas(com_range, label: str) -> None:
    try:
        formulas = com_range.Formula
    except Exception as exc:
        raise ProcessingError(f"无法检查{label}是否包含公式；为避免覆盖公式，已拒绝此区域处理。") from exc
    if any(isinstance(value, str) and value.startswith("=") for value in _flatten_com_values(formulas)):
        raise ProcessingError(f"{label}包含公式单元格；为避免删除公式，已拒绝此区域处理。")


def _assert_target_area_empty(com_range, bounds: tuple[int, int, int, int], label: str) -> None:
    _assert_no_formulas(com_range, label)
    matrix = _read_com_area(com_range, bounds)
    for row_offset, row in enumerate(matrix):
        for col_offset, value in enumerate(row):
            if _cell_has_data(value):
                cell = f"{get_column_letter(bounds[0] + col_offset)}{bounds[1] + row_offset}"
                raise ProcessingError(f"{label}将写入的目标单元格 {cell} 已有数据；不会静默覆盖，请调整目标列或先清空该区域。")


def _excel_number_format(mode: str) -> str:
    return {"text": "@", "number": "0.###############", "date": "yyyy-mm-dd"}[mode]


def _apply_region_processing(workbook, settings: dict[str, Any], formula_jobs, log: Callable[[str], None]) -> bool:
    formula_areas: dict[str, list[tuple[int, int, int, int]]] = {}
    for job, formula_job in zip(settings.get("jobs", []), formula_jobs):
        sheet_name, start_row, end_row, _mappings = formula_job
        col_start, col_end = parse_column_range(job["formula_range"])
        formula_areas.setdefault(sheet_name, []).append((col_start, start_row, col_end, end_row))
        template_row = _formula_template_row(job)
        formula_areas[sheet_name].append((1, template_row, 16_384, template_row))

    applied_any = False
    operation_labels = {"TextToColumns": "文本分列", "Deduplicate": "去重", "FormatConvert": "格式转换"}
    configured_order = settings.get("workflow", DEFAULT_WORKFLOW)
    for operation_key in (module_id for module_id in configured_order if module_id in operation_labels):
        label = operation_labels[operation_key]
        operation = settings.get(operation_key, {"enabled": False, "rules": []})
        if not operation.get("enabled", False):
            continue
        for rule_idx, rule in enumerate(operation.get("rules", []), 1):
            sheet_name = str(rule["sheet"]).strip()
            bounds = _parse_selected_area(rule["range"], f"第 {rule_idx} 条{label}")
            _check_processing_area_size(bounds, f"第 {rule_idx} 条{label}")
            if any(_areas_overlap(bounds, protected) for protected in formula_areas.get(sheet_name, [])):
                raise ProcessingError(f"第 {rule_idx} 条{label}规则区域 {sheet_name}!{_area_address(bounds)} 与本次填充的公式区域重叠；为避免删除保留公式，已停止。")
            worksheet = workbook.Worksheets(sheet_name)
            source_range = worksheet.Range(_area_address(bounds))

            if operation_key == "TextToColumns":
                values_matrix = _read_com_area(source_range, bounds)
                values = [row[0] for row in values_matrix]
                formats = [str(value).lower() for value in rule.get("column_formats", ["text"])]
                split_matrix, result_columns = split_text_column(values, str(rule["delimiter"]), formats)
                start_col = column_number(rule["output_start_col"])
                target_bounds = (start_col, bounds[1], start_col + result_columns - 1, bounds[3])
                if _areas_overlap(bounds, target_bounds):
                    raise ProcessingError(f"第 {rule_idx} 条文本分列规则的输出区域 {_area_address(target_bounds)} 与源区域 {_area_address(bounds)} 重叠；请更改输出起始列。")
                if target_bounds[2] > 16_384:
                    raise ProcessingError(f"第 {rule_idx} 条文本分列规则输出超出 XFD 列。")
                _check_processing_area_size(target_bounds, f"第 {rule_idx} 条文本分列目标")
                if any(_areas_overlap(target_bounds, protected) for protected in formula_areas.get(sheet_name, [])):
                    raise ProcessingError(f"第 {rule_idx} 条文本分列输出区域 {sheet_name}!{_area_address(target_bounds)} 会覆盖公式区域；已停止。")
                target_range = worksheet.Range(_area_address(target_bounds))
                _assert_target_area_empty(target_range, target_bounds, f"第 {rule_idx} 条文本分列")
                _write_com_area(target_range, split_matrix)
                for col_offset in range(result_columns):
                    col_letter = get_column_letter(start_col + col_offset)
                    fmt = formats[col_offset]
                    worksheet.Range(f"{col_letter}{bounds[1]}:{col_letter}{bounds[3]}").NumberFormat = _excel_number_format(fmt)
                log(f"已执行文本分列规则 {rule_idx}：{sheet_name}!{_area_address(bounds)} → {sheet_name}!{_area_address(target_bounds)}，分隔符={rule['delimiter']!r}。")

            elif operation_key == "Deduplicate":
                _assert_no_formulas(source_range, f"第 {rule_idx} 条去重规则区域")
                matrix = _read_com_area(source_range, bounds)
                key_mode = str(rule.get("key_mode", "row")).lower()
                key_indexes = None
                if key_mode == "columns":
                    key_indexes = [column_number(value) - bounds[0] for value in rule["key_columns"]]
                result, removed = deduplicate_matrix(matrix, bool(rule.get("has_header", False)), key_mode, key_indexes)
                _write_com_area(source_range, result)
                log(f"已执行去重规则 {rule_idx}：{sheet_name}!{_area_address(bounds)}，按{'整行' if key_mode == 'row' else '关键列 ' + ','.join(rule['key_columns'])}判重，移除 {removed} 行，保留首次出现。")

            else:
                _assert_no_formulas(source_range, f"第 {rule_idx} 条格式转换规则区域")
                matrix = _read_com_area(source_range, bounds)
                mode = str(rule["mode"]).lower()
                converted = convert_matrix_values(matrix, mode)
                _write_com_area(source_range, converted)
                source_range.NumberFormat = _excel_number_format(mode)
                log(f"已执行格式转换规则 {rule_idx}：{sheet_name}!{_area_address(bounds)} → {mode}。")
            applied_any = True
    return applied_any


def _pywin32_import_error_message(exc: ImportError) -> str:
    return (
        "无法导入 Windows COM 依赖 pywin32（pythoncom/win32com）。\n"
        f"当前 Python 解释器：{sys.executable}\n"
        "请从本程序目录双击 run_app.bat，它会使用本项目 .venv 中的 Python。"
        "不要直接双击 main.py，也不要用系统 Python 运行。\n"
        "如果 run_app.bat 仍提示缺少依赖，请先运行 install_and_build.bat；"
        "安装失败诊断保存在 install_and_build.log。\n"
        f"导入错误详情：{type(exc).__name__}: {exc}"
    )


def _connect_wps(log: Callable[[str], None]):
    if sys.platform != "win32":
        raise ProcessingError("WPS 自动重算只能在安装了 Windows 版 WPS 表格的 Windows 环境运行；当前环境不是 Windows，未生成输出文件。")
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ProcessingError(_pywin32_import_error_message(exc)) from exc
    try:
        pythoncom.CoInitialize()
    except Exception as exc:
        raise ProcessingError(f"无法初始化 Windows COM，因而不能连接 WPS 表格：{type(exc).__name__}: {exc}") from exc
    progids = ("ket.Application", "ET.Application", "et.Application")
    failures = []
    for progid in progids:
        try:
            log(f"正在连接 WPS 表格 COM：{progid}")
            return win32com.client.DispatchEx(progid), pythoncom, progid
        except Exception as exc:
            failures.append(f"{progid}: {exc}")
    pythoncom.CoUninitialize()
    raise ProcessingError("无法连接 Windows 版 WPS 表格 COM/ProgID。已尝试 ket.Application、ET.Application、et.Application。\n"
                          "请确认 Windows 11 已安装并能正常启动 WPS 表格；修复/重新注册 WPS 表格自动化组件后重试。\n"
                          "本程序不会改用 Microsoft Excel，也不会在连接失败后留下最终输出文件。\n\n连接错误详情：\n" + "\n".join(failures))


def _calculate_in_wps(staging_path: Path, formula_jobs: list[tuple[str, int, int, list[dict[str, str]]]],
                      log: Callable[[str], None], post_mappings: list[dict[str, Any]] | None = None,
                      region_processing: dict[str, Any] | None = None) -> None:
    app = workbook = pythoncom = None
    com_initialized = False
    try:
        app, pythoncom, progid = _connect_wps(log)
        com_initialized = True
        for prop, value in (("Visible", False), ("DisplayAlerts", False), ("AskToUpdateLinks", False), ("EnableEvents", False), ("AutomationSecurity", 3)):
            try:
                setattr(app, prop, value)
            except Exception:
                pass
        try:
            app.Calculation = -4105
        except Exception:
            pass
        try:
            workbook = app.Workbooks.Open(str(staging_path), UpdateLinks=0, ReadOnly=False, AddToMru=False)
        except Exception as exc:
            raise ProcessingError(f"已连接 WPS（{progid}），但无法在禁用外部链接更新模式下打开中间工作簿：{exc}") from exc
        log("已在 WPS 中打开中间工作簿，开始完整重算公式。")
        _recalculate_and_wait(app, log, "映射前")
        log("映射前公式重算完成；现在写入结果列映射及配置的后映射区域。")
        for sheet_name, start_row, end_row, mappings in formula_jobs:
            ws = workbook.Worksheets(sheet_name)
            count = end_row - start_row + 1
            for mapping in mappings:
                result_col, target_col = mapping["result_col"], mapping["target_col"]
                src = ws.Range(f"{result_col}{start_row}:{result_col}{end_row}").Value2
                if count == 1:
                    values = [src[0][0] if isinstance(src, (tuple, list)) else src]
                else:
                    values = [row[0] for row in src]
                destination = ws.Range(f"{target_col}{start_row}:{target_col}{end_row}")
                _write_mapped_values(destination, values, mapping.get("write_as", "value"))
                log(f"已将 {sheet_name}!{result_col}{start_row}:{result_col}{end_row} 按{'文本' if mapping.get('write_as') == 'text' else '值'}写入 {target_col} 列。")
        for idx, rule in enumerate(post_mappings or [], 1):
            source_bounds = _parse_area(rule["source_range"], f"第 {idx} 条后映射源")
            target_bounds = _parse_area(rule["target_range"], f"第 {idx} 条后映射目标")
            rows = source_bounds[3] - source_bounds[1] + 1
            cols = source_bounds[2] - source_bounds[0] + 1
            source_address = _area_address(source_bounds)
            target_address = _area_address(target_bounds)
            source_range = workbook.Worksheets(str(rule["source_sheet"]).strip()).Range(source_address)
            source_values = _com_value_matrix(source_range.Value2, rows, cols)
            target_range = workbook.Worksheets(str(rule["target_sheet"]).strip()).Range(target_address)
            target_range.Value2 = tuple(tuple(row) for row in source_values)
            log(f"已应用后映射 {rule['source_sheet']}!{source_address} → {rule['target_sheet']}!{target_address}。")

        step_settings = dict(region_processing or {})
        step_settings.setdefault("TextToColumns", {"enabled": False, "rules": []})
        step_settings.setdefault("Deduplicate", {"enabled": False, "rules": []})
        step_settings.setdefault("FormatConvert", {"enabled": False, "rules": []})
        step_settings.setdefault("workflow", list(DEFAULT_WORKFLOW))
        applied_steps = _apply_region_processing(workbook, step_settings, formula_jobs, log)

        # Rebuild once after all post-mapping and region-processing writes so dependent formulas are current.
        recalc_phase = "后映射及区域处理后" if applied_steps else "后映射写入后"
        _recalculate_and_wait(app, log, recalc_phase)
        workbook.Save()
        log("WPS 已完成后映射、第二次完整重算并保存工作簿；接下来只导出区域 XLSX、按需转换 PNG 并读取文本预览，不再改动工作簿。")
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"WPS 自动化处理中失败：{type(exc).__name__}: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        if com_initialized and pythoncom is not None:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass


def _validate_post_mapping_conflicts(settings: dict[str, Any], source_records: list[list[list[Any]]]) -> None:
    formula_areas: dict[str, list[tuple[int, int, int, int]]] = {}
    column_map_areas: dict[str, list[tuple[int, int, int, int]]] = {}
    for job, records in zip(settings["jobs"], source_records):
        sheet = str(job["template_sheet"]).strip()
        start_row = int(job["start_row"])
        end_row = start_row + len(records) - 1
        c1, c2 = parse_column_range(job["formula_range"])
        formula_areas.setdefault(sheet, []).append((c1, start_row, c2, end_row))
        template_row = _formula_template_row(job)
        formula_areas[sheet].append((1, template_row, 16_384, template_row))
        for mapping in job["mappings"]:
            col = column_number(mapping["target_col"])
            column_map_areas.setdefault(sheet, []).append((col, start_row, col, end_row))

    for idx, rule in enumerate(settings.get("post_mappings", []), 1):
        sheet = str(rule["target_sheet"]).strip()
        target = _parse_area(rule["target_range"], f"第 {idx} 条后映射目标")
        for protected in formula_areas.get(sheet, []):
            if _areas_overlap(target, protected):
                raise ProcessingError(f"第 {idx} 条后映射目标 {sheet}!{_area_address(target)} 会覆盖保留公式区域 {_area_address(protected)}；已停止处理。")
        for protected in column_map_areas.get(sheet, []):
            if _areas_overlap(target, protected):
                raise ProcessingError(f"第 {idx} 条后映射目标 {sheet}!{_area_address(target)} 与结果映射目标 {_area_address(protected)} 重复；已停止处理。")


def _sanitize_filename_fragment(value: str) -> str:
    sanitized = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value).strip())
    sanitized = re.sub(r"\s+", "_", sanitized).strip("._ ")
    return sanitized[:80] or "区域"


def _region_xlsx_filename(output: dict[str, Any], index: int, collision: int = 1,
                          multiple_exports: bool = False) -> str:
    sheet = _sanitize_filename_fragment(str(output["sheet"]))
    export_suffix = str(index) if multiple_exports else ""
    collision_suffix = "" if collision == 1 else f"_{collision}"
    return f"{sheet}{export_suffix}{collision_suffix}.xlsx"


def _copy_file_exclusive(source: Path, destination: Path) -> None:
    """Copy a completed file without replacing any pre-existing destination."""
    descriptor = os.open(destination, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o666)
    try:
        with os.fdopen(descriptor, "wb") as target, source.open("rb") as origin:
            while chunk := origin.read(1024 * 1024):
                target.write(chunk)
    except Exception:
        destination.unlink(missing_ok=True)
        raise


def _publish_staging_workbook(staging_path: Path, final_path: Path) -> None:
    """Publish a completed workbook without replacing a pre-existing file."""
    try:
        os.link(staging_path, final_path)
    except FileExistsError as exc:
        raise ProcessingError(f"输出文件在处理过程中被创建；为避免覆盖已停止：\n{final_path}") from exc
    except OSError:
        try:
            _copy_file_exclusive(staging_path, final_path)
        except FileExistsError as exc:
            raise ProcessingError(f"输出文件在处理过程中被创建；为避免覆盖已停止：\n{final_path}") from exc


def _resolve_selected_output_bounds(worksheet, selected_bounds: tuple[int, int, int, int],
                                    selection_kind: str) -> tuple[int, int, int, int]:
    used_rows = max(1, min(1_048_576, int(worksheet.max_row or 1)))
    used_cols = max(1, min(16_384, int(worksheet.max_column or 1)))
    min_col, min_row, max_col, max_row = selected_bounds
    if selection_kind == "sheet":
        return 1, 1, used_cols, used_rows
    if selection_kind == "row":
        return 1, min_row, used_cols, min_row
    if selection_kind == "column":
        return min_col, 1, min_col, used_rows
    return selected_bounds


def _column_dimension_for_index(worksheet, column: int):
    letter = get_column_letter(column)
    exact = worksheet.column_dimensions.get(letter)
    if exact is not None and (exact.min or column) <= column <= (exact.max or exact.min or column):
        return exact
    for dimension in worksheet.column_dimensions.values():
        first = dimension.min
        if first is None:
            try:
                first = column_index_from_string(str(dimension.index))
            except ValueError:
                continue
        last = dimension.max or first
        if first <= column <= last:
            return dimension
    return None


def _column_is_hidden(worksheet, column: int) -> bool:
    dimension = _column_dimension_for_index(worksheet, column)
    return bool(dimension.hidden) if dimension is not None else False


def _split_excel_format_sections(format_code: str) -> list[str]:
    sections: list[str] = []
    current: list[str] = []
    quoted = False
    escaped = False
    bracket_depth = 0
    for char in str(format_code or "General"):
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            current.append(char)
            escaped = True
        elif char == '"' and bracket_depth == 0:
            current.append(char)
            quoted = not quoted
        elif not quoted and char == "[":
            bracket_depth += 1
            current.append(char)
        elif not quoted and char == "]" and bracket_depth:
            bracket_depth -= 1
            current.append(char)
        elif char == ";" and not quoted and bracket_depth == 0:
            sections.append("".join(current))
            current = []
        else:
            current.append(char)
    sections.append("".join(current))
    return sections


def _excel_format_condition(section: str) -> tuple[str, Decimal] | None:
    match = re.search(r"\[(<=|>=|<>|=|<|>)\s*(-?\d+(?:\.\d+)?)\]", section)
    if not match:
        return None
    try:
        return match.group(1), Decimal(match.group(2))
    except InvalidOperation:
        return None


def _format_condition_matches(value: Decimal, condition: tuple[str, Decimal]) -> bool:
    operator, threshold = condition
    return {
        "<": value < threshold,
        "<=": value <= threshold,
        ">": value > threshold,
        ">=": value >= threshold,
        "=": value == threshold,
        "<>": value != threshold,
    }[operator]


def _select_excel_format_section(value: Any, format_code: str) -> tuple[str, int, bool]:
    sections = _split_excel_format_sections(format_code)
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return sections[-1] if len(sections) >= 4 else sections[0], 0, False
    conditional = [_excel_format_condition(section) for section in sections]
    if any(condition is not None for condition in conditional):
        for index, condition in enumerate(conditional):
            if condition is not None and _format_condition_matches(number, condition):
                return sections[index], index, True
        for index, condition in enumerate(conditional):
            if condition is None:
                return sections[index], index, True
        return sections[0], 0, True
    if number < 0 and len(sections) >= 2:
        return sections[1], 1, False
    if number == 0 and len(sections) >= 3:
        return sections[2], 2, False
    return sections[0], 0, False


def _format_literal_affix(text: str) -> str:
    result: list[str] = []
    index = 0
    while index < len(text):
        char = text[index]
        if char == '"':
            end = text.find('"', index + 1)
            if end < 0:
                result.append(text[index + 1:])
                break
            result.append(text[index + 1:end])
            index = end + 1
        elif char == "\\" and index + 1 < len(text):
            result.append(text[index + 1])
            index += 2
        elif char == "_" and index + 1 < len(text):
            result.append(" ")
            index += 2
        elif char == "*" and index + 1 < len(text):
            index += 2
        elif char == "[":
            end = text.find("]", index + 1)
            if end < 0:
                index += 1
                continue
            directive = text[index + 1:end]
            currency = re.match(r"\$([^\]-]+)(?:-[^]]*)?$", directive)
            if currency and currency.group(1):
                result.append(currency.group(1))
            index = end + 1
        else:
            result.append(char)
            index += 1
    return "".join(result)


def _numeric_pattern_span(section: str) -> tuple[int, int, str] | None:
    visible: list[tuple[str, int, int]] = []
    index = 0
    while index < len(section):
        char = section[index]
        if char == '"':
            end = section.find('"', index + 1)
            index = len(section) if end < 0 else end + 1
        elif char in ("\\", "_", "*") and index + 1 < len(section):
            index += 2
        elif char == "[":
            end = section.find("]", index + 1)
            index = len(section) if end < 0 else end + 1
        else:
            visible.append((char, index, index + 1))
            index += 1
    for offset, (char, start, end) in enumerate(visible):
        if char not in "0#?":
            continue
        stop = offset + 1
        while stop < len(visible):
            next_char = visible[stop][0]
            if next_char in "0#?,.":
                stop += 1
                continue
            if next_char in "Ee" and stop + 1 < len(visible) and visible[stop + 1][0] in "+-":
                stop += 2
                while stop < len(visible) and visible[stop][0] in "0#?":
                    stop += 1
                continue
            break
        start_pos = start
        end_pos = visible[stop - 1][2]
        return start_pos, end_pos, section[start_pos:end_pos]
    return None


def _format_number_with_excel_mask(value: Any, section: str, section_index: int,
                                   conditional_format: bool) -> str:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return str(value)
    if not number.is_finite():
        return str(value)
    selected_negative_section = number < 0 and section_index == 1 and not conditional_format
    number_for_display = abs(number) if selected_negative_section else number
    pattern_info = _numeric_pattern_span(section)
    if pattern_info is None:
        literal = _format_literal_affix(section)
        return literal if literal else str(value) if section.strip().lower() == "general" else ""
    start, end, pattern = pattern_info
    if re.search(r"\d\s*/\s*[0#?]", pattern):
        return str(value)
    exponent_match = re.search(r"[Ee]([+-]?)([0#?]+)$", pattern)
    percent_count = section.count("%")
    scaled = number_for_display * (Decimal(100) ** percent_count)
    mantissa_pattern = pattern
    exponent_pattern = None
    if exponent_match:
        exponent_pattern = exponent_match.group(0)
        mantissa_pattern = pattern[:exponent_match.start()]
    integer_pattern, dot, decimal_pattern = mantissa_pattern.partition(".")
    trailing_scale_commas = len(integer_pattern) - len(integer_pattern.rstrip(","))
    integer_pattern = integer_pattern.rstrip(",")
    scaled /= Decimal(1000) ** trailing_scale_commas

    if exponent_pattern:
        decimals = len(decimal_pattern)
        try:
            rendered = f"{scaled:.{decimals}E}"
            mantissa, exponent_value = rendered.split("E")
            exponent_value_int = int(exponent_value)
            sign_char, exponent_digits = exponent_pattern[1], exponent_pattern[2:]
            exponent_sign = "+" if exponent_value_int >= 0 else "-"
            if sign_char == "-" and exponent_value_int >= 0:
                exponent_sign = ""
            formatted = (f"{mantissa}E{exponent_sign}"
                         f"{abs(exponent_value_int):0{max(1, len(exponent_digits))}d}")
        except (ValueError, InvalidOperation):
            return str(value)
    else:
        min_decimals = decimal_pattern.count("0")
        max_decimals = len(decimal_pattern)
        quantizer = Decimal(1).scaleb(-max_decimals)
        try:
            rounded = scaled.quantize(quantizer, rounding=ROUND_HALF_UP) if max_decimals else scaled.quantize(Decimal(1), rounding=ROUND_HALF_UP)
        except InvalidOperation:
            return str(value)
        fixed = f"{abs(rounded):.{max_decimals}f}" if rounded < 0 else f"{rounded:.{max_decimals}f}"
        integer_text, _, decimal_text = fixed.partition(".")
        required_integer_zeros = integer_pattern.count("0")
        if not integer_text and required_integer_zeros:
            integer_text = "0"
        integer_text = integer_text.zfill(required_integer_zeros)
        if "," in integer_pattern:
            integer_text = f"{int(integer_text or '0'):,}"
        if max_decimals > min_decimals:
            while len(decimal_text) > min_decimals and decimal_text.endswith("0"):
                decimal_text = decimal_text[:-1]
        decimal_text = decimal_text.ljust(min_decimals, "0")
        if decimal_text:
            formatted = f"{integer_text or '0'}.{decimal_text}"
        else:
            formatted = integer_text
        if not formatted and scaled == 0 and not required_integer_zeros and not decimal_text:
            formatted = ""

    prefix = _format_literal_affix(section[:start])
    suffix = _format_literal_affix(section[end:])
    if number < 0 and not selected_negative_section and not conditional_format:
        prefix = "-" + prefix
    return prefix + formatted + suffix


def _format_temporal_value(value: Any, format_code: str) -> str:
    if isinstance(value, timedelta):
        seconds = int(value.total_seconds())
        hours, remainder = divmod(abs(seconds), 3600)
        minutes, seconds = divmod(remainder, 60)
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    if isinstance(value, date) and not isinstance(value, datetime):
        value = datetime(value.year, value.month, value.day)
    if isinstance(value, datetime_time):
        value = datetime.combine(date(1899, 12, 30), value)
    if not isinstance(value, datetime):
        return str(value)
    section = _split_excel_format_sections(format_code)[0]
    tokens: list[tuple[str, str]] = []
    index = 0
    token_re = re.compile(r"AM/PM|A/P|yyyy|yy|mmmm|mmm|mm|m|dddd|ddd|dd|d|hh|h|ss|s", re.IGNORECASE)
    while index < len(section):
        char = section[index]
        if char == '"':
            end = section.find('"', index + 1)
            if end < 0:
                tokens.append(("literal", section[index + 1:]))
                break
            tokens.append(("literal", section[index + 1:end]))
            index = end + 1
        elif char == "\\" and index + 1 < len(section):
            tokens.append(("literal", section[index + 1]))
            index += 2
        elif char in "_*" and index + 1 < len(section):
            if char == "_":
                tokens.append(("literal", " "))
            index += 2
        elif char == "[":
            end = section.find("]", index + 1)
            if end < 0:
                index += 1
                continue
            directive = section[index + 1:end]
            if directive.lower() in ("h", "hh", "m", "mm", "s", "ss"):
                tokens.append(("elapsed", directive))
            else:
                currency = re.match(r"\$([^\]-]+)(?:-[^]]*)?$", directive)
                if currency and currency.group(1):
                    tokens.append(("literal", currency.group(1)))
            index = end + 1
        else:
            match = token_re.match(section, index)
            if match:
                tokens.append(("date", match.group(0)))
                index = match.end()
            else:
                tokens.append(("literal", char))
                index += 1

    date_tokens = [text.lower() for kind, text in tokens if kind == "date"]
    has_ampm = any(token in ("am/pm", "a/p") for token in date_tokens)
    rendered: list[str] = []
    prior_date_token = ""
    for token_index, (kind, token) in enumerate(tokens):
        if kind == "literal":
            rendered.append(token)
            continue
        if kind == "elapsed":
            rendered.append(token)
            continue
        lower = token.lower()
        if lower == "am/pm":
            rendered.append("AM" if value.hour < 12 else "PM")
        elif lower == "a/p":
            rendered.append("A" if value.hour < 12 else "P")
        elif lower.startswith("y"):
            rendered.append(f"{value.year % 100:02d}" if len(lower) <= 2 else f"{value.year:04d}")
        elif lower.startswith("d"):
            if len(lower) >= 4:
                rendered.append(("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")[value.weekday()])
            elif len(lower) == 3:
                rendered.append(("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")[value.weekday()])
            elif len(lower) == 2:
                rendered.append(f"{value.day:02d}")
            else:
                rendered.append(str(value.day))
        elif lower.startswith("h"):
            hour = value.hour % 12 or 12 if has_ampm else value.hour
            rendered.append(f"{hour:02d}" if len(lower) == 2 else str(hour))
        elif lower.startswith("s"):
            rendered.append(f"{value.second:02d}" if len(lower) == 2 else str(value.second))
        elif lower.startswith("m"):
            next_date_token = ""
            for later_kind, later_text in tokens[token_index + 1:]:
                if later_kind == "date":
                    next_date_token = later_text.lower()
                    break
                if later_kind == "literal" and later_text not in (":", ".", " "):
                    break
            minute = prior_date_token.startswith("h") or next_date_token.startswith("s")
            if minute:
                rendered.append(f"{value.minute:02d}" if len(lower) >= 2 else str(value.minute))
            elif len(lower) >= 4:
                rendered.append(("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")[value.month - 1])
            elif len(lower) == 3:
                rendered.append(("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")[value.month - 1])
            elif len(lower) == 2:
                rendered.append(f"{value.month:02d}")
            else:
                rendered.append(str(value.month))
        prior_date_token = lower
    return "".join(rendered)


def _value_as_export_text(value: Any, number_format: str = "General") -> str | None:
    """Return a non-formula text value approximating Excel's displayed cell text."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, datetime_time, timedelta)):
        if number_format and number_format.lower() not in ("general", "@") and is_date_format(number_format):
            return _format_temporal_value(value, number_format)
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, datetime_time):
            return value.isoformat()
        return str(value)
    if isinstance(value, (int, float, Decimal)):
        if not number_format or number_format.lower() in ("general", "@"):
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value)
        section, section_index, conditional = _select_excel_format_section(value, number_format)
        return _format_number_with_excel_mask(value, section, section_index, conditional)
    return str(value)


def _read_wps_display_text(safe_copy_path: Path, requested_cells: list[tuple[str, int, int]],
                           log: Callable[[str], None]) -> dict[tuple[str, int, int], str]:
    """Read numeric/date Range.Text from a sanitized temporary copy when WPS COM is available."""
    if sys.platform != "win32" or not requested_cells:
        return {}
    app = workbook = pythoncom = None
    com_initialized = False
    result: dict[tuple[str, int, int], str] = {}
    try:
        app, pythoncom, progid = _connect_wps(log)
        com_initialized = True
        for prop, value in (("Visible", False), ("DisplayAlerts", False), ("AskToUpdateLinks", False),
                            ("EnableEvents", False), ("AutomationSecurity", 3)):
            try:
                setattr(app, prop, value)
            except Exception:
                pass
        workbook = app.Workbooks.Open(str(safe_copy_path), UpdateLinks=0, ReadOnly=True, AddToMru=False)
        sheets: dict[str, Any] = {}
        for sheet_name, row, col in requested_cells:
            try:
                if sheet_name not in sheets:
                    sheets[sheet_name] = workbook.Worksheets(sheet_name)
                worksheet = sheets[sheet_name]
                text = worksheet.Cells(row, col).Text
                if text is None:
                    continue
                text = str(text)
                # Range.Text may return hashes solely because the source column is too narrow.
                if text and re.fullmatch(r"#+", text.strip()):
                    continue
                result[(sheet_name, row, col)] = text
            except Exception:
                continue
        if result:
            log(f"区域 XLSX 导出已通过 WPS {progid} 的单元格 Text 读取 {len(result):,} 个数值/日期显示文本。")
    except Exception as exc:
        log(f"提示：无法从 WPS COM 读取单元格显示文本，将按 number_format 转换；导出继续且原工作簿不变。原因：{type(exc).__name__}: {exc}")
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        if com_initialized and pythoncom is not None:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass
    return result


def _copy_dimension_properties(source_dimension, target_dimension, names: tuple[str, ...]) -> None:
    for name in names:
        try:
            setattr(target_dimension, name, getattr(source_dimension, name))
        except (AttributeError, TypeError, ValueError):
            pass
    if getattr(source_dimension, "has_style", False):
        for name in ("font", "fill", "border", "alignment", "protection", "number_format"):
            try:
                setattr(target_dimension, name, copy(getattr(source_dimension, name)))
            except (AttributeError, TypeError, ValueError):
                pass


def _export_table_areas(workbook_path: Path, area_outputs: list[dict[str, Any]], output_dir: Path,
                        log: Callable[[str], None], reserved_paths: set[Path] | None = None) -> list[Path]:
    """Export areas from a sanitized read-only copy, never directly from WPS's saved file."""
    if not area_outputs:
        return []
    output_dir.mkdir(parents=True, exist_ok=True)
    try:
        with _temporary_workbook_copy(workbook_path, "已保存工作簿", log) as safe_copy:
            source_workbook = safe_copy.load_workbook(data_only=True, keep_links=False)
            reserved = {Path(path).resolve() for path in (reserved_paths or set())}
            return _export_table_areas_loaded(
                source_workbook, area_outputs, output_dir, log,
                multiple_exports=len(area_outputs) > 1, reserved_paths=reserved,
                source_copy_path=safe_copy.path,
            )
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"独立区域 XLSX 导出失败：{type(exc).__name__}: {exc}") from exc


def _export_table_areas_loaded(source_workbook, area_outputs: list[dict[str, Any]], output_dir: Path,
                               log: Callable[[str], None], multiple_exports: bool,
                               reserved_paths: set[Path], source_copy_path: Path | None = None) -> list[Path]:
    """Create one styled, text-valued XLSX workbook for each configured source area."""
    generated: list[Path] = []
    try:
        export_specs = []
        requested_display_cells = set()
        for index, output in enumerate(area_outputs, 1):
            sheet_name = str(output["sheet"]).strip()
            if sheet_name not in source_workbook.sheetnames:
                raise ProcessingError(f"第 {index} 项区域 XLSX 输出工作表“{sheet_name}”在已保存工作簿中不存在。")
            worksheet = source_workbook[sheet_name]
            selected_bounds, selection_kind = _parse_selection_details(
                output["range"], f"第 {index} 项区域 XLSX 输出"
            )
            configured_bounds = _resolve_selected_output_bounds(worksheet, selected_bounds, selection_kind)
            _check_region_xlsx_area_size(configured_bounds, f"第 {index} 项区域 XLSX 输出")
            bounds = _expand_bounds_for_merged_cells(
                worksheet, configured_bounds, f"第 {index} 项区域 XLSX 输出"
            )
            _check_region_xlsx_area_size(bounds, f"第 {index} 项区域 XLSX 输出（合并区域闭包）")
            min_col, min_row, max_col, max_row = bounds
            source_rows = list(range(min_row, max_row + 1))
            source_cols = list(range(min_col, max_col + 1))
            export_specs.append({
                "index": index, "output": output, "sheet_name": sheet_name, "worksheet": worksheet,
                "bounds": bounds, "configured_bounds": configured_bounds,
                "min_col": min_col, "min_row": min_row,
                "source_rows": source_rows, "source_cols": source_cols,
            })
            for source_row in source_rows:
                for source_col in source_cols:
                    value = worksheet.cell(source_row, source_col).value
                    if value is not None and not isinstance(value, bool) and isinstance(
                            value, (int, float, Decimal, datetime, date, datetime_time, timedelta)):
                        requested_display_cells.add((sheet_name, source_row, source_col))

        display_texts = _read_wps_display_text(
            source_copy_path, sorted(requested_display_cells), log
        ) if source_copy_path is not None else {}

        for spec in export_specs:
            index = spec["index"]
            output = spec["output"]
            sheet_name = spec["sheet_name"]
            worksheet = spec["worksheet"]
            bounds = spec["bounds"]
            configured_bounds = spec["configured_bounds"]
            min_col, min_row, max_col, max_row = bounds
            source_rows = spec["source_rows"]
            source_cols = spec["source_cols"]
            row_map = {source_row: source_row - min_row + 1 for source_row in source_rows}
            col_map = {source_col: source_col - min_col + 1 for source_col in source_cols}

            export_workbook = Workbook()
            export_sheet = export_workbook.active
            export_sheet.title = worksheet.title
            export_sheet.sheet_view.showGridLines = worksheet.sheet_view.showGridLines
            export_sheet.sheet_view.zoomScale = worksheet.sheet_view.zoomScale
            for name in ("defaultRowHeight", "defaultColWidth", "baseColWidth", "zeroHeight"):
                try:
                    setattr(export_sheet.sheet_format, name, getattr(worksheet.sheet_format, name))
                except (AttributeError, TypeError, ValueError):
                    pass
            export_sheet.sheet_format.zeroHeight = worksheet.sheet_format.zeroHeight

            total_chars = 0
            try:
                for target_row, source_row in enumerate(source_rows, 1):
                    source_dimension = worksheet.row_dimensions.get(source_row)
                    if source_dimension is not None:
                        _copy_dimension_properties(
                            source_dimension, export_sheet.row_dimensions[target_row],
                            ("height", "outlineLevel", "collapsed", "hidden", "thickTop", "thickBot")
                        )
                    for target_col, source_col in enumerate(source_cols, 1):
                        source_cell = worksheet.cell(source_row, source_col)
                        target_cell = export_sheet.cell(target_row, target_col)
                        if source_cell.has_style:
                            for style_name in ("font", "fill", "border", "alignment", "protection"):
                                setattr(target_cell, style_name, copy(getattr(source_cell, style_name)))
                            target_cell.number_format = source_cell.number_format
                        display_key = (sheet_name, source_row, source_col)
                        text = display_texts.get(display_key)
                        if text is None:
                            text = _value_as_export_text(source_cell.value, source_cell.number_format)
                        if text is not None:
                            total_chars += len(text)
                            if total_chars > MAX_REGION_XLSX_CHARS:
                                raise ProcessingError(
                                    f"第 {index} 项区域 XLSX 输出超过 {MAX_REGION_XLSX_CHARS:,} 个字符上限；请缩小范围。"
                                )
                            target_cell.value = text
                            target_cell.data_type = "s"
                for target_col, source_col in enumerate(source_cols, 1):
                    source_dimension = _column_dimension_for_index(worksheet, source_col)
                    target_dimension = export_sheet.column_dimensions[get_column_letter(target_col)]
                    if source_dimension is not None:
                        _copy_dimension_properties(
                            source_dimension, target_dimension,
                            ("width", "bestFit", "outlineLevel", "collapsed", "hidden")
                        )

                for merged in worksheet.merged_cells.ranges:
                    if not _areas_overlap(bounds, (merged.min_col, merged.min_row, merged.max_col, merged.max_row)):
                        continue
                    if (merged.min_row in row_map and merged.max_row in row_map
                            and merged.min_col in col_map and merged.max_col in col_map):
                        export_sheet.merge_cells(
                            start_row=row_map[merged.min_row], start_column=col_map[merged.min_col],
                            end_row=row_map[merged.max_row], end_column=col_map[merged.max_col]
                        )

                temp_handle, temp_name = tempfile.mkstemp(
                    prefix=".wps_region_export_", suffix=".xlsx", dir=str(output_dir)
                )
                os.close(temp_handle)
                temp_path = Path(temp_name)
                try:
                    export_workbook.save(temp_path)
                    if not temp_path.is_file() or temp_path.stat().st_size == 0:
                        raise ProcessingError(f"第 {index} 项区域 XLSX 输出未生成有效文件。")
                    collision = 1
                    while True:
                        target = output_dir / _region_xlsx_filename(output, index, collision, multiple_exports)
                        if target.resolve() in reserved_paths:
                            collision += 1
                            continue
                        try:
                            os.link(temp_path, target)
                            break
                        except FileExistsError:
                            collision += 1
                        except OSError:
                            if target.exists():
                                collision += 1
                                continue
                            try:
                                _copy_file_exclusive(temp_path, target)
                                break
                            except FileExistsError:
                                collision += 1
                    generated.append(target)
                finally:
                    temp_path.unlink(missing_ok=True)
            finally:
                export_workbook.close()

            if bounds != configured_bounds:
                log(f"区域 {sheet_name}!{_area_address(configured_bounds)} 与合并单元格相交，导出范围扩展为 {_area_address(bounds)}。")
            log(
                f"已导出区域 XLSX {sheet_name}!{_area_address(bounds)}：{target.name}；"
                f"选区内隐藏行列及其原始位置均保留；单元格值均为文本。"
            )
        return generated
    except Exception as exc:
        for path in generated:
            path.unlink(missing_ok=True)
        if isinstance(exc, ProcessingError):
            raise
        raise ProcessingError(f"独立区域 XLSX 导出失败：{type(exc).__name__}: {exc}") from exc

def _format_text_matrix(matrix: list[list[Any]]) -> str:
    return "\n".join("\t".join("" if value is None else str(value) for value in row) for row in matrix)


def convert_cell_value(value: Any, mode: str, epoch=WINDOWS_EPOCH) -> Any:
    """Pure value conversion used by both WPS range processing and Linux tests."""
    mode = str(mode).strip().lower()
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None if mode in ("number", "date") else value
    if mode == "text":
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else str(value)
    if mode == "number":
        if isinstance(value, bool):
            return int(value)
        try:
            number = Decimal(str(value).replace(",", "").strip())
        except (InvalidOperation, ValueError) as exc:
            raise ProcessingError(f"无法将 {value!r} 转换为数字。") from exc
        if not number.is_finite():
            raise ProcessingError(f"无法将 {value!r} 转换为有限数字。")
        return int(number) if number == number.to_integral_value() else float(number)
    if mode == "date":
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime(value.year, value.month, value.day)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            try:
                result = from_excel(value, epoch=epoch)
                if isinstance(result, datetime):
                    return result
                if isinstance(result, date):
                    return datetime(result.year, result.month, result.day)
            except (OverflowError, ValueError) as exc:
                raise ProcessingError(f"无法将 Excel 数值 {value!r} 转换为日期。") from exc
        text = str(value).strip()
        normalized = text.replace("/", "-")
        for parser in (datetime.fromisoformat,):
            try:
                return parser(normalized)
            except ValueError:
                pass
        for fmt in ("%Y%m%d", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                pass
        raise ProcessingError(f"无法将 {value!r} 转换为日期；支持日期单元格、Excel 日期序号或 YYYY-MM-DD/YYYYMMDD 文本。")
    raise ProcessingError(f"不支持的转换类型：{mode}")


def convert_matrix_values(matrix: list[list[Any]], mode: str, epoch=WINDOWS_EPOCH) -> list[list[Any]]:
    converted = []
    for row_index, row in enumerate(matrix, 1):
        converted_row = []
        for col_index, value in enumerate(row, 1):
            try:
                converted_row.append(convert_cell_value(value, mode, epoch))
            except ProcessingError as exc:
                raise ProcessingError(f"第 {row_index} 行第 {col_index} 列：{exc}") from exc
        converted.append(converted_row)
    return converted


def split_text_column(values: list[Any], delimiter: str, column_formats: list[str]) -> tuple[list[list[Any]], int]:
    rows = []
    maximum_parts = 1
    for value in values:
        parts = [None] if value is None else str(value).split(delimiter)
        rows.append(parts)
        maximum_parts = max(maximum_parts, len(parts))
    if maximum_parts > len(column_formats):
        raise ProcessingError(f"分列结果最多有 {maximum_parts} 列，但仅配置了 {len(column_formats)} 个结果列格式。")
    result = []
    for row in rows:
        padded = row + [None] * (maximum_parts - len(row))
        converted = []
        for col_index, value in enumerate(padded):
            try:
                converted.append(convert_cell_value(value, column_formats[col_index]))
            except ProcessingError as exc:
                raise ProcessingError(f"分列结果第 {col_index + 1} 列：{exc}") from exc
        result.append(converted)
    return result, maximum_parts


def deduplicate_matrix(matrix: list[list[Any]], has_header: bool = False, key_mode: str = "row",
                       key_indexes: list[int] | None = None) -> tuple[list[list[Any]], int]:
    if not matrix:
        return [], 0
    width = max((len(row) for row in matrix), default=0)
    rows = [list(row) + [None] * (width - len(row)) for row in matrix]
    result = [rows[0]] if has_header else []
    seen = set()
    removed = 0
    data_rows = rows[1:] if has_header else rows
    for row in data_rows:
        if not any(_cell_has_data(value) for value in row):
            result.append(row)
            continue
        key = tuple(row) if key_mode == "row" else tuple(row[index] for index in (key_indexes or []))
        try:
            already_seen = key in seen
        except TypeError:
            key = tuple(repr(value) for value in key)
            already_seen = key in seen
        if already_seen:
            removed += 1
        else:
            seen.add(key)
            result.append(row)
    result.extend([[None] * width for _ in range(len(rows) - len(result))])
    return result, removed


def _read_text_output_areas(workbook_path: Path, text_outputs: list[dict[str, Any]],
                            log: Callable[[str], None] = lambda _message: None) -> list[dict[str, Any]]:
    if not text_outputs:
        return []
    try:
        with _temporary_workbook_copy(workbook_path, "已保存工作簿", log) as safe_copy:
            workbook = safe_copy.load_workbook(data_only=True, keep_links=False, read_only=True)
            return _read_text_output_areas_loaded(workbook, text_outputs)
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"读取文本输出区域失败：{exc}") from exc


def _read_text_output_areas_loaded(workbook, text_outputs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    try:
        previews = []
        for index, output in enumerate(text_outputs, 1):
            sheet_name = str(output["sheet"]).strip()
            if sheet_name not in workbook.sheetnames:
                raise ProcessingError(f"第 {index} 项文本输出工作表“{sheet_name}”在已保存工作簿中不存在。")
            selected_bounds, selection_kind = _parse_selection_details(output["range"], f"第 {index} 项文本输出")
            worksheet = workbook[sheet_name]
            used_rows = max(1, min(1_048_576, int(worksheet.max_row or 1)))
            used_cols = max(1, min(16_384, int(worksheet.max_column or 1)))
            min_col, min_row, max_col, max_row = selected_bounds
            if selection_kind == "sheet":
                bounds = (1, 1, used_cols, used_rows)
            elif selection_kind == "row":
                bounds = (1, min_row, used_cols, min_row)
            elif selection_kind == "column":
                bounds = (min_col, 1, min_col, used_rows)
            else:
                bounds = selected_bounds
            _check_text_area_size(bounds, f"第 {index} 项文本输出（实际范围 {_area_address(bounds)}）")
            min_col, min_row, max_col, max_row = bounds
            chunks = []
            full_text_parts = []
            total_chars = 0
            for chunk_start in range(min_row, max_row + 1, TEXT_COPY_CHUNK_ROWS):
                chunk_end = min(max_row, chunk_start + TEXT_COPY_CHUNK_ROWS - 1)
                lines = []
                for row in worksheet.iter_rows(min_row=chunk_start, max_row=chunk_end,
                                               min_col=min_col, max_col=max_col, values_only=True):
                    lines.append("\t".join("" if value is None else str(value) for value in row))
                chunk_text = "\n".join(lines)
                separator_length = 1 if full_text_parts else 0
                total_chars += separator_length + len(chunk_text)
                if total_chars > MAX_TEXT_CHARS:
                    raise ProcessingError(f"第 {index} 项文本输出超过 {MAX_TEXT_CHARS:,} 个字符的剪贴板上限；请缩小区域。")
                full_text_parts.append(chunk_text)
                chunks.append({"start_row": chunk_start, "end_row": chunk_end, "text": chunk_text})
            previews.append({"sheet": sheet_name, "range": str(output["range"]),
                             "resolved_range": _area_address(bounds),
                             "text": "\n".join(full_text_parts), "chunks": chunks})
        return previews
    except ProcessingError:
        raise
    except Exception as exc:
        raise ProcessingError(f"读取文本输出区域失败：{exc}") from exc


def process_workbooks(raw_settings: dict[str, Any], log: Callable[[str], None] = print,
                      on_text_previews: Callable[[list[dict[str, str]]], None] | None = None,
                      on_region_outputs: Callable[[list[Path]], None] | None = None) -> Path:
    """Process and save the XLSX, then publish it after terminal outputs."""
    settings = normalize_settings(raw_settings)
    template_path, final_path = _validate_inputs(settings)
    copy_start, copy_end = column_number(settings["copy_start_col"]), column_number(settings["copy_end_col"])
    all_records: list[list[list[Any]]] = []
    for job_idx, job in enumerate(settings["jobs"], 1):
        job_records: list[list[Any]] = []
        for source_idx, source in enumerate(job["sources"], 1):
            path = Path(source["path"]).expanduser().resolve()
            start_row = int(source["copy_start_row"])
            end_value = source.get("copy_end_row")
            end_row = int(end_value) if end_value not in (None, "") else None
            display_sheet = str(source.get("source_sheet") or "").strip() or "首个工作表（默认）"
            end_label = str(end_row) if end_row is not None else "自动空白边界/末行"
            log(f"读取目标表 {job_idx}/{len(settings['jobs'])} 的源文件 {source_idx}/{len(job['sources'])}：{path.name}（{display_sheet}，第 {start_row} 行起，结束：{end_label}）")
            records = _read_source_records(path, str(source["source_sheet"]).strip(), start_row, copy_start, copy_end, log, end_row)
            job_records.extend(records)
        all_records.append(job_records)
    _validate_pre_mapping_conflicts(settings, all_records)
    _validate_post_mapping_conflicts(settings, all_records)
    fd, temp_name = tempfile.mkstemp(prefix=f".{final_path.stem}_", suffix=".staging.xlsx", dir=str(final_path.parent))
    os.close(fd)
    staging_path = Path(temp_name)
    generated_region_files: list[Path] = []
    published_final = False
    preserve_final_after_region_failure = False
    preserve_region_files_after_image_failure = False
    try:
        log("校验模板公式和目标子表后，在输出目录创建临时副本。")
        formula_jobs = _create_staging_workbook(template_path, settings, all_records, staging_path, log)
        _calculate_in_wps(staging_path, formula_jobs, log, settings.get("post_mappings", []), settings)
        if not staging_path.is_file() or staging_path.stat().st_size == 0:
            raise ProcessingError("WPS 处理结束但临时工作簿不存在或为空；没有发布输出文件。")
        if final_path.exists():
            raise ProcessingError(f"输出文件在处理过程中被创建；为避免覆盖已停止：\n{final_path}")

        # WPS has saved all cell changes. These are terminal read/export steps; the workbook is never modified again.
        log("工作簿已保存；开始最后的区域 XLSX 导出与文本区域读取（后续不再修改工作簿）。")
        text_previews = []
        for module_id in settings.get("workflow", DEFAULT_WORKFLOW):
            if module_id == "image_output":
                region_xlsx_files: list[Path] = []
                try:
                    region_xlsx_files = _export_table_areas(
                        staging_path, settings.get("image_outputs", []), final_path.parent, log,
                        reserved_paths={final_path},
                    )
                    if settings.get("export_region_pngs"):
                        try:
                            generated_region_files = _export_region_xlsx_as_pngs(region_xlsx_files, log)
                        except Exception:
                            preserve_region_files_after_image_failure = True
                            raise
                        for region_path in region_xlsx_files:
                            try:
                                region_path.unlink(missing_ok=True)
                                if region_path.exists():
                                    log(f"PNG 已导出，但未能删除独立区域 XLSX（文件仍保留）：{region_path}")
                                else:
                                    log(f"PNG 已成功导出，临时独立区域 XLSX 已删除：{region_path}")
                            except OSError as delete_exc:
                                log(f"PNG 已导出，但删除独立区域 XLSX 失败；该表已保留：{region_path}；{delete_exc}")
                    else:
                        generated_region_files = region_xlsx_files
                except Exception as image_exc:
                    try:
                        _publish_staging_workbook(staging_path, final_path)
                    except Exception as publish_exc:
                        raise ProcessingError(
                            f"区域输出失败：{image_exc}\n"
                            f"WPS 已完成计算并成功保存工作簿，但无法安全发布到目标路径：\n"
                            f"{final_path}\n发布原因：{publish_exc}"
                        ) from image_exc
                    published_final = True
                    preserve_final_after_region_failure = True
                    output_kind = "区域 PNG 图片导出" if settings.get("export_region_pngs") else "区域 XLSX 输出"
                    fallback_note = ""
                    if preserve_region_files_after_image_failure and region_xlsx_files:
                        fallback_paths = "\n".join(str(path) for path in region_xlsx_files)
                        fallback_note = f"\n图片导出未完成；独立区域 XLSX 已保留：\n{fallback_paths}"
                    raise ProcessingError(
                        f"{output_kind}失败：{image_exc}{fallback_note}\n"
                        f"WPS 已完成计算并成功保存 XLSX；该文件已保留：\n{final_path}"
                    ) from image_exc
            elif module_id == "text_output":
                text_previews = _read_text_output_areas(staging_path, settings.get("text_outputs", []), log)
        for preview in text_previews:
            log(f"已读取文本区域 {preview['sheet']}!{preview['range']}，内容将在处理完成后显示并可复制。")

        _publish_staging_workbook(staging_path, final_path)
        published_final = True
        staging_path.unlink(missing_ok=True)
        log(f"处理成功，已另存为：{final_path}")
        if on_text_previews is not None:
            on_text_previews(text_previews)
        if on_region_outputs is not None:
            on_region_outputs(list(generated_region_files))
        return final_path
    except Exception as exc:
        try:
            staging_path.unlink(missing_ok=True)
            if not preserve_region_files_after_image_failure:
                for region_path in generated_region_files:
                    region_path.unlink(missing_ok=True)
            if published_final and not preserve_final_after_region_failure:
                final_path.unlink(missing_ok=True)
        except Exception as cleanup_exc:
            raise ProcessingError(f"处理失败：{exc}\n同时无法清理临时工作簿：{staging_path}\n清理错误：{cleanup_exc}") from exc
        raise


def _copy_region_png_exclusive(source: Path, region_xlsx: Path) -> Path:
    """Publish one generated PNG beside its XLSX without replacing existing files."""
    collision = 1
    while True:
        suffix = "" if collision == 1 else f"_{collision}"
        destination = region_xlsx.parent / f"{region_xlsx.stem}{suffix}.png"
        try:
            _copy_file_exclusive(source, destination)
            return destination
        except FileExistsError:
            collision += 1
        except OSError:
            if destination.exists():
                collision += 1
                continue
            raise


def _export_region_xlsx_as_pngs(paths: list[Path], log: Callable[[str], None]) -> list[Path]:
    """Render each already-exported region workbook through WPS and publish PNGs safely."""
    if not paths:
        return []
    app = workbook = pythoncom = None
    com_initialized = False
    generated: list[Path] = []
    try:
        app, pythoncom, progid = _connect_wps(log)
        com_initialized = True
        for prop, value in (("Visible", False), ("DisplayAlerts", False),
                            ("AskToUpdateLinks", False), ("EnableEvents", False),
                            ("AutomationSecurity", 3)):
            try:
                setattr(app, prop, value)
            except Exception:
                pass
        for region_path in paths:
            if not region_path.is_file():
                raise ProcessingError(f"待转 PNG 的独立区域 XLSX 不存在：{region_path}")
            log(f"正在从独立区域 XLSX 导出 PNG：{region_path.name}")
            with tempfile.TemporaryDirectory(prefix=".wps_region_png_", dir=str(region_path.parent)) as temp_dir:
                temp_png = Path(temp_dir) / "render.png"
                workbook = app.Workbooks.Open(
                    str(region_path), UpdateLinks=0, ReadOnly=False, AddToMru=False
                )
                worksheet = workbook.Worksheets(1)
                area = worksheet.UsedRange
                width = max(48.0, min(2000.0, float(area.Width)))
                height = max(24.0, min(12000.0, float(area.Height)))
                worksheet.Activate()
                chart_object = worksheet.ChartObjects().Add(
                    float(area.Left), float(area.Top), width, height
                )
                area.CopyPicture(Appearance=1, Format=-4147)
                chart_object.Activate()
                chart_object.Chart.Paste()
                exported = chart_object.Chart.Export(str(temp_png), "PNG")
                if not exported or not temp_png.is_file() or temp_png.stat().st_size == 0:
                    raise ProcessingError(f"WPS 未能从区域 XLSX 导出有效 PNG：{region_path}")
                workbook.Close(SaveChanges=False)
                workbook = None
                png_path = _copy_region_png_exclusive(temp_png, region_path)
                generated.append(png_path)
                log(f"已导出区域 PNG：{png_path}")
        return generated
    except Exception as exc:
        for path in generated:
            path.unlink(missing_ok=True)
        if isinstance(exc, ProcessingError):
            raise
        raise ProcessingError(f"WPS 将独立区域 XLSX 导出为 PNG 失败：{type(exc).__name__}: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
        if com_initialized and pythoncom is not None:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass


class App:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
        self.tk, self.filedialog, self.messagebox, self.ttk = tk, filedialog, messagebox, ttk
        self.root = root
        self.root.title(APP_NAME)
        try:
            screen_width = self.root.winfo_screenwidth()
            screen_height = self.root.winfo_screenheight()
        except Exception:
            screen_width = screen_height = None
        min_width, min_height, _max_width, _max_height = window_size_limits(screen_width, screen_height)
        width, height = load_window_size(screen_width=screen_width, screen_height=screen_height)
        self.root.resizable(True, True)
        self.root.geometry(f"{width}x{height}")
        self.root.minsize(min_width, min_height)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)
        self.events: queue.Queue[tuple[str, Any]] = queue.Queue()
        self.settings = load_settings()
        self.output_filename_custom = bool(self.settings.get("output_filename_custom", False))
        self.export_region_pngs_var = tk.BooleanVar(value=bool(self.settings.get("export_region_pngs", False)))
        self._export_pngs_for_run = False
        self._region_output_files: list[Path] = []
        self._suppress_output_filename_trace = False
        startup_migration_notices = self.settings.pop("_migration_notices", [])
        self.vars: dict[str, Any] = {}
        self.job_vars: dict[str, Any] = {}
        self.operation_enabled: dict[str, Any] = {}
        self.active_job_index: int | None = None
        self.active_profile_name = str(self.settings.get("config_name") or "")
        self._build()
        self._render_jobs()
        self._render_area_configs()
        self._render_workflow()
        self._refresh_profile_choices()
        if startup_migration_notices:
            self.root.after(200, self._show_configuration_migration_notices, startup_migration_notices)
        self.root.after(150, self._drain_events)

    def _on_close(self):
        try:
            save_window_size(
                self.root.winfo_width(), self.root.winfo_height(),
                screen_width=self.root.winfo_screenwidth(),
                screen_height=self.root.winfo_screenheight(),
            )
        except Exception:
            pass
        self.root.destroy()

    def _on_canvas_configure(self, event):
        requested_width = self.content_frame.winfo_reqwidth()
        requested_height = self.content_frame.winfo_reqheight()
        width = max(event.width, requested_width)
        height = max(event.height, requested_height)
        self.main_canvas.itemconfigure(self.content_window, width=width, height=height)
        self.main_canvas.configure(scrollregion=(0, 0, width, height))

    def _on_content_configure(self, _event=None):
        bounds = self.main_canvas.bbox("all")
        if bounds:
            self.main_canvas.configure(scrollregion=bounds)

    def _entry(self, parent, row, label, key, width=64, browse=None):
        self.ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(8, 8), pady=3)
        var = self.tk.StringVar(value=str(self.settings.get(key, "")))
        self.vars[key] = var
        if key == "output_filename":
            var.trace_add("write", self._on_output_filename_edited)
        self.ttk.Entry(parent, textvariable=var, width=width).grid(row=row, column=1, sticky="ew", padx=4, pady=3)
        if browse:
            self.ttk.Button(parent, text="浏览…", command=browse).grid(row=row, column=2, padx=4, pady=3)
        return var

    def _on_output_filename_edited(self, *_args):
        if not self._suppress_output_filename_trace:
            self.output_filename_custom = True

    def _set_output_filename(self, value: str, custom: bool):
        self._suppress_output_filename_trace = True
        try:
            self.vars["output_filename"].set(str(value))
        finally:
            self._suppress_output_filename_trace = False
        self.output_filename_custom = bool(custom)

    def _set_template_selected(self, path: str):
        self.vars["template_path"].set(str(path))
        if not self.output_filename_custom:
            default_name = _template_output_filename(path)
            if default_name:
                self._set_output_filename(default_name, custom=False)

    def _build(self):
        ttk = self.ttk
        viewport = ttk.Frame(self.root)
        viewport.grid(row=0, column=0, sticky="nsew")
        viewport.rowconfigure(0, weight=1)
        viewport.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)
        self.root.columnconfigure(0, weight=1)
        self.main_canvas = self.tk.Canvas(viewport, highlightthickness=0)
        vertical_scrollbar = ttk.Scrollbar(viewport, orient="vertical", command=self.main_canvas.yview)
        horizontal_scrollbar = ttk.Scrollbar(viewport, orient="horizontal", command=self.main_canvas.xview)
        self.main_canvas.configure(yscrollcommand=vertical_scrollbar.set,
                                   xscrollcommand=horizontal_scrollbar.set)
        self.main_canvas.grid(row=0, column=0, sticky="nsew")
        vertical_scrollbar.grid(row=0, column=1, sticky="ns")
        horizontal_scrollbar.grid(row=1, column=0, sticky="ew")
        outer = ttk.Frame(self.main_canvas, padding=9)
        self.content_frame = outer
        self.content_window = self.main_canvas.create_window((0, 0), window=outer, anchor="nw")
        self.main_canvas.bind("<Configure>", self._on_canvas_configure)
        outer.bind("<Configure>", self._on_content_configure)
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(2, weight=3)
        outer.rowconfigure(6, weight=1)
        files = ttk.LabelFrame(outer, text="模板、输出与源数据列", padding=7)
        files.grid(row=0, column=0, sticky="ew")
        files.columnconfigure(1, weight=1)
        self._entry(files, 0, "模板工作簿（无宏 .xlsx）", "template_path", browse=self._browse_template)
        self._entry(files, 1, "输出目录", "output_dir", browse=self._browse_output_dir)
        self._entry(files, 2, "输出文件名", "output_filename", width=35)
        row = 3
        for key, label in (("copy_start_col", "源数据复制起始列"), ("copy_end_col", "源数据复制末列（必填）")):
            ttk.Label(files, text=label).grid(row=row, column=0, sticky="w", padx=(8, 8), pady=3)
            var = self.tk.StringVar(value=str(self.settings.get(key, "")))
            self.vars[key] = var
            ttk.Entry(files, textvariable=var, width=12).grid(row=row, column=1, sticky="w", padx=4, pady=3)
            row += 1
        ttk.Label(files, text="模板前缀匹配长度 N（1–120）").grid(row=5, column=0, sticky="w", padx=(8, 8), pady=3)
        self.vars["template_prefix_length"] = self.tk.StringVar(value=str(self.settings.get("template_prefix_length", 5)))
        ttk.Entry(files, textvariable=self.vars["template_prefix_length"], width=8).grid(row=5, column=1, sticky="w", padx=4, pady=3)
        ttk.Label(files, text="选模板时按文件名（不含扩展名）前 N 个字符匹配已保存配置；无匹配时不会自动套用其他配置。",
                  foreground="#4b5563", wraplength=1180, justify="left").grid(row=6, column=0, columnspan=3, sticky="w", padx=8, pady=3)

        target_box = ttk.LabelFrame(outer, text="目标子表配置（可添加多个；源文件按列表顺序追加）", padding=7)
        target_box.grid(row=1, column=0, sticky="nsew", pady=(8, 0))
        target_box.columnconfigure(0, weight=1)
        target_box.rowconfigure(1, weight=1)
        actions = ttk.Frame(target_box)
        actions.grid(row=0, column=0, sticky="ew")
        ttk.Button(actions, text="添加目标表", command=self._add_job).pack(side="left", padx=3)
        ttk.Button(actions, text="删除目标表", command=self._delete_job).pack(side="left", padx=3)
        self.job_tree = ttk.Treeview(target_box, columns=("sheet", "start", "formula", "sources", "maps"), show="headings", height=4, selectmode="browse")
        for col, title, width in (("sheet", "模板子表", 230), ("start", "粘贴起始行", 90), ("formula", "公式列范围", 120), ("sources", "源文件数（按顺序）", 140), ("maps", "结果映射（类型）", 290)):
            self.job_tree.heading(col, text=title)
            self.job_tree.column(col, width=width, anchor="w")
        self.job_tree.grid(row=1, column=0, sticky="nsew", padx=3, pady=4)
        self.job_tree.bind("<<TreeviewSelect>>", self._on_job_select)

        editor = ttk.LabelFrame(outer, text="选中目标子表的参数", padding=7)
        editor.grid(row=2, column=0, sticky="nsew", pady=(7, 0))
        editor.columnconfigure(0, weight=1)
        editor.columnconfigure(1, weight=1)
        top = ttk.Frame(editor)
        top.grid(row=0, column=0, columnspan=2, sticky="ew")
        self.job_vars = {key: self.tk.StringVar() for key in ("template_sheet", "header_rows", "start_row", "formula_range")}
        self.job_vars["formula_template_row"] = self.tk.StringVar()
        self.job_vars["start_row"].trace_add("write", self._update_formula_template_row_display)
        fields = (("template_sheet", "模板子表名", 25), ("header_rows", "保留标题行数", 10),
                  ("start_row", "粘贴起始行", 13), ("formula_template_row", "公式模板行（自动）", 13),
                  ("formula_range", "公式列范围", 13))
        for i, (key, label, width) in enumerate(fields):
            ttk.Label(top, text=label).grid(row=0, column=i * 2, sticky="w", padx=4)
            state = "readonly" if key == "formula_template_row" else "normal"
            ttk.Entry(top, textvariable=self.job_vars[key], width=width, state=state).grid(row=0, column=i * 2 + 1, sticky="w", padx=4)
        self.source_tree = ttk.Treeview(editor, columns=("path", "sheet", "from", "to"), show="headings", height=7)
        for col, title, width in (("path", "源文件路径", 390), ("sheet", "源工作表（空=首表）", 165), ("from", "复制起始行", 90), ("to", "复制结束行（空=空行边界）", 170)):
            self.source_tree.heading(col, text=title)
            self.source_tree.column(col, width=width, anchor="w")
        self.source_tree.grid(row=1, column=0, sticky="nsew", padx=3, pady=5)
        source_buttons = ttk.Frame(editor)
        source_buttons.grid(row=2, column=0, sticky="w", padx=3)
        for label, command in (("添加源文件…", self._add_source), ("编辑…", self._edit_source), ("删除", self._delete_source), ("上移", lambda: self._move_source(-1)), ("下移", lambda: self._move_source(1))):
            ttk.Button(source_buttons, text=label, command=command).pack(side="left", padx=2)
        self.mapping_tree = ttk.Treeview(editor, columns=("result", "target", "type"), show="headings", height=5)
        for col, title, width in (("result", "结果列", 90), ("target", "目标列", 90), ("type", "写入类型", 110)):
            self.mapping_tree.heading(col, text=title)
            self.mapping_tree.column(col, width=width, anchor="w")
        self.mapping_tree.grid(row=1, column=1, sticky="nsew", padx=3, pady=5)
        map_buttons = ttk.Frame(editor)
        map_buttons.grid(row=2, column=1, sticky="w", padx=3)
        for label, command in (("添加映射", self._add_mapping), ("编辑映射", self._edit_mapping), ("删除映射", self._delete_mapping)):
            ttk.Button(map_buttons, text=label, command=command).pack(side="left", padx=2)
        editor.rowconfigure(1, weight=1)
        note = "执行顺序：模板清理/前映射 → 多源追加 → 填公式与完整重算 → 后映射 → TextToColumns → Deduplicate → FormatConvert → 再完整重算并保存 → 最后导出区域 XLSX 与文本预览/复制（此后不改工作簿）。区域 XLSX 保留样式和隐藏行列，尽量按表格显示文本写入文本值；三个区域处理功能各有独立开关和规则列表。"
        ttk.Label(editor, text=note, foreground="#4b5563", wraplength=1180, justify="left").grid(row=3, column=0, columnspan=2, sticky="w", padx=5, pady=(5, 0))

        area_box = ttk.LabelFrame(outer, text="模板内部前/后映射及终端输出区域（与每份源文件无关）", padding=5)
        area_box.grid(row=3, column=0, sticky="nsew", pady=(7, 0))
        area_box.columnconfigure(0, weight=1)
        area_box.rowconfigure(0, weight=1)
        area_tabs = ttk.Notebook(area_box)
        area_tabs.grid(row=0, column=0, sticky="nsew")
        self.pre_mapping_tree = self._create_area_tab(area_tabs, "前映射", "pre_mappings",
            (("source_sheet", "源工作表", 180), ("source_range", "源范围", 110), ("target_sheet", "目标工作表", 180), ("target_range", "目标范围", 110)))
        self.image_tree = self._create_area_tab(area_tabs, "区域 XLSX 输出", "image_outputs",
            (("sheet", "模板工作表", 220), ("range", "单元格范围", 180)))
        self.text_tree = self._create_area_tab(area_tabs, "文本预览/复制", "text_outputs",
            (("sheet", "模板工作表", 220), ("range", "单元格范围", 180)))
        self.post_mapping_tree = self._create_area_tab(area_tabs, "后映射", "post_mappings",
            (("source_sheet", "源工作表", 180), ("source_range", "源范围", 110), ("target_sheet", "目标工作表", 180), ("target_range", "目标范围", 110)))
        self.text_to_columns_tree = self._create_operation_tab(area_tabs, "TextToColumns", "TextToColumns",
            (("sheet", "工作表", 150), ("range", "源范围", 110), ("delimiter", "分隔符", 80), ("output_start_col", "输出起始列", 90), ("column_formats", "结果列格式", 190)))
        self.deduplicate_tree = self._create_operation_tab(area_tabs, "Deduplicate", "Deduplicate",
            (("sheet", "工作表", 150), ("range", "处理范围", 120), ("has_header", "含表头", 80), ("key_mode", "判重方式", 100), ("key_columns", "关键列", 130)))
        self.format_convert_tree = self._create_operation_tab(area_tabs, "FormatConvert", "FormatConvert",
            (("sheet", "工作表", 150), ("range", "处理范围", 140), ("mode", "目标类型", 100)))
        self.workflow_tree = self._create_workflow_tab(area_tabs)
        selector_help = ("范围选择（映射区保持普通矩形语法）：单格 A1=整表，A5=整行5，D1=整列D；"
                         "B3,B3=单格；两个不同端点如 D8,B3=自动规范化的矩形。Excel 上限 XFD / 1048576 行。"
                         f"区域 XLSX 整行/整列/整表按已用范围解析，最多 {MAX_REGION_XLSX_CELLS:,} 格、{MAX_REGION_XLSX_CHARS:,} 字符；隐藏行列保留原位置并继续隐藏。文本整行/整列/整表按工作表已用范围解析，"
                         f"最多 {MAX_TEXT_CELLS:,} 格、{MAX_TEXT_CHARS:,} 字符，并按 {TEXT_COPY_CHUNK_ROWS} 行分块预览/复制；"
                         f"区域处理单条规则最多 {MAX_PROCESSING_CELLS:,} 格。")
        ttk.Label(area_box, text=selector_help, foreground="#4b5563", wraplength=1180,
                  justify="left").grid(row=1, column=0, sticky="w", padx=4, pady=(4, 1))
        ttk.Checkbutton(area_box,
                        text="处理成功后从单独导出的区域 XLSX 生成 PNG，并删除成功转换的 XLSX（默认关闭；图片失败时保留 XLSX）",
                        variable=self.export_region_pngs_var).grid(row=2, column=0, sticky="w", padx=4, pady=(3, 1))

        controls = ttk.Frame(outer)
        controls.grid(row=4, column=0, sticky="ew", pady=7)
        self.run_button = ttk.Button(controls, text="开始处理", command=self._run)
        self.run_button.pack(side="left")
        ttk.Button(controls, text="保存配置…", command=self._export_config).pack(side="left", padx=5)
        ttk.Button(controls, text="加载配置…", command=self._import_config).pack(side="left", padx=5)
        ttk.Button(controls, text="保存为默认配置", command=self._save_default).pack(side="left", padx=5)
        ttk.Button(controls, text="恢复默认", command=self._reset_config).pack(side="left", padx=5)
        ttk.Label(controls, text="命名配置").pack(side="left", padx=(12, 3))
        self.profile_choice = self.tk.StringVar(value=self.active_profile_name)
        self.profile_combo = ttk.Combobox(controls, textvariable=self.profile_choice, state="readonly", width=18)
        self.profile_combo.pack(side="left", padx=3)
        ttk.Button(controls, text="新建并保存", command=self._new_profile).pack(side="left", padx=3)
        ttk.Button(controls, text="保存当前", command=self._save_current_profile).pack(side="left", padx=3)
        ttk.Button(controls, text="载入所选", command=self._load_selected_profile).pack(side="left", padx=3)
        self.progress = ttk.Progressbar(controls, mode="indeterminate", length=160)
        self.progress.pack(side="right", padx=5)

        preview_frame = ttk.LabelFrame(outer, text="文本区域预览（按行换行、按列 Tab 分隔；复制按钮写入 Windows 剪贴板）", padding=5)
        preview_frame.grid(row=5, column=0, sticky="nsew", pady=(0, 6))
        self.preview_empty_label = ttk.Label(preview_frame, text="处理成功后，配置的每个文本区域会在此显示并提供独立复制按钮。", foreground="#4b5563")
        self.preview_empty_label.pack(anchor="w", padx=4, pady=4)
        self.preview_notebook = ttk.Notebook(preview_frame)
        self.preview_notebook.pack(fill="both", expand=True)
        self.preview_notebook.pack_forget()

        log_frame = ttk.LabelFrame(outer, text="处理日志", padding=5)
        log_frame.grid(row=6, column=0, sticky="nsew")
        self.log_text = self.tk.Text(log_frame, height=8, wrap="word", state="disabled")
        scrollbar = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scrollbar.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._log("为每个目标子表维护有序源文件列表；源工作表名留空默认首表，下一份源会紧接已复制记录追加。")

    def _create_workflow_tab(self, notebook):
        ttk = self.ttk
        page = ttk.Frame(notebook, padding=6)
        notebook.add(page, text="模块工作流")
        page.columnconfigure(0, weight=1)
        page.rowconfigure(1, weight=1)
        ttk.Label(page, text="添加模块并用上移/下移定制执行顺序。核心依赖不可颠倒；分列、去重、格式转换可互换，区域 XLSX 与文本输出可互换。非法顺序会在保存或开始处理时给出具体原因。",
                  wraplength=1080, justify="left").grid(row=0, column=0, sticky="w", pady=(0, 5))
        self.workflow_tree = ttk.Treeview(page, columns=("module", "label", "role"), show="headings", height=7, selectmode="browse")
        for column, title, width in (("module", "模块 ID", 190), ("label", "模块名称", 190), ("role", "配置/依赖说明", 470)):
            self.workflow_tree.heading(column, text=title)
            self.workflow_tree.column(column, width=width, anchor="w")
        self.workflow_tree.grid(row=1, column=0, sticky="nsew")
        controls = ttk.Frame(page)
        controls.grid(row=2, column=0, sticky="w", pady=(5, 0))
        self.module_choice = self.tk.StringVar()
        self.module_combo = ttk.Combobox(controls, textvariable=self.module_choice, state="readonly", width=34)
        self.module_combo.pack(side="left", padx=2)
        ttk.Button(controls, text="添加模块", command=self._add_workflow_module).pack(side="left", padx=2)
        ttk.Button(controls, text="移除所选", command=self._remove_workflow_module).pack(side="left", padx=2)
        ttk.Button(controls, text="上移", command=lambda: self._move_workflow_module(-1)).pack(side="left", padx=2)
        ttk.Button(controls, text="下移", command=lambda: self._move_workflow_module(1)).pack(side="left", padx=2)
        return self.workflow_tree

    def _render_workflow(self):
        for item in self.workflow_tree.get_children():
            self.workflow_tree.delete(item)
        workflow = self.settings.get("workflow", DEFAULT_WORKFLOW)
        descriptions = {
            "template_cleanup": "保留模板格式/公式种子，清理旧数据",
            "pre_mapping": "读取模板原值并在源数据前写入",
            "source_copy": "按目标表及源文件列表追加数据",
            "formula_fill": "按复制数据行数填充公式",
            "formula_recalc": "首次完整重算供结果映射读取",
            "result_mapping": "把公式结果写到配置目标列",
            "post_mapping": "公式结果映射后执行模板内部写入",
            "TextToColumns": "区域处理阶段，可与另两项区域操作互换",
            "Deduplicate": "区域处理阶段，可与另两项区域操作互换",
            "FormatConvert": "区域处理阶段，可与另两项区域操作互换",
            "final_recalc_save": "所有工作簿写入后最终完整重算并保存",
            "image_output": "只读末端区域 XLSX 导出，必须在最终保存后",
            "text_output": "只读末端输出，必须在最终保存后",
        }
        for index, module_id in enumerate(workflow):
            self.workflow_tree.insert("", "end", iid=str(index),
                                      values=(module_id, MODULE_LABELS.get(module_id, module_id), descriptions.get(module_id, "")))
        available = [f"{key} — {label}" for key, label in MODULE_LABELS.items() if key not in workflow]
        self.module_combo.configure(values=available)
        if self.module_choice.get() not in available:
            self.module_choice.set(available[0] if available else "")

    def _change_workflow(self, candidate):
        previous = list(self.settings.get("workflow", DEFAULT_WORKFLOW))
        self.settings["workflow"] = candidate
        try:
            validate_workflow(self.settings)
            _validate_region_module_dependencies(self.settings)
        except (ValueError, ProcessingError) as exc:
            self.settings["workflow"] = previous
            self.messagebox.showerror(APP_NAME, str(exc))
            return False
        self._render_workflow()
        return True

    def _add_workflow_module(self):
        selected = self.module_choice.get().split(" — ", 1)[0]
        if selected in MODULE_LABELS and selected not in self.settings.get("workflow", []):
            candidate = add_workflow_module(list(self.settings.get("workflow", [])), selected)
            self._change_workflow(candidate)

    def _remove_workflow_module(self):
        selected = self.workflow_tree.selection()
        if not selected:
            return
        index = int(selected[0])
        candidate = remove_workflow_module(list(self.settings.get("workflow", DEFAULT_WORKFLOW)), index)
        self._change_workflow(candidate)

    def _move_workflow_module(self, direction):
        selected = self.workflow_tree.selection()
        if not selected:
            return
        index = int(selected[0])
        destination = index + direction
        previous = list(self.settings.get("workflow", DEFAULT_WORKFLOW))
        if not 0 <= destination < len(previous):
            return
        candidate = move_workflow_module(previous, index, direction)
        if self._change_workflow(candidate):
            self.workflow_tree.selection_set(str(destination))
            self.workflow_tree.focus(str(destination))

    def _refresh_profile_choices(self):
        names = list_profiles(PROFILE_DIR)
        self.profile_combo.configure(values=names)
        if self.active_profile_name in names:
            self.profile_choice.set(self.active_profile_name)
        elif self.profile_choice.get() not in names:
            self.profile_choice.set("")

    def _create_area_tab(self, notebook, title, key, columns):
        ttk = self.ttk
        page = ttk.Frame(notebook, padding=4)
        notebook.add(page, text=title)
        page.columnconfigure(0, weight=1)
        page.rowconfigure(0, weight=1)
        identifiers = tuple(column[0] for column in columns)
        tree = ttk.Treeview(page, columns=identifiers, show="headings", height=4, selectmode="browse")
        for identifier, heading, width in columns:
            tree.heading(identifier, text=heading)
            tree.column(identifier, width=width, anchor="w")
        tree.grid(row=0, column=0, sticky="nsew")
        buttons = ttk.Frame(page)
        buttons.grid(row=1, column=0, sticky="w", pady=(4, 0))
        for action, label in (("add", "添加"), ("edit", "编辑"), ("delete", "删除")):
            ttk.Button(buttons, text=label, command=lambda k=key, t=tree, a=action: self._edit_area_config(k, t, a)).pack(side="left", padx=2)
        return tree

    def _create_operation_tab(self, notebook, title, key, columns):
        ttk = self.ttk
        page = ttk.Frame(notebook, padding=4)
        notebook.add(page, text=title)
        page.columnconfigure(0, weight=1)
        page.rowconfigure(1, weight=1)
        enabled = self.tk.BooleanVar(value=self.settings.get(key, {}).get("enabled", False))
        self.operation_enabled[key] = enabled
        ttk.Checkbutton(page, text=f"启用 {title}", variable=enabled).grid(row=0, column=0, sticky="w", pady=(0, 3))
        identifiers = tuple(column[0] for column in columns)
        tree = ttk.Treeview(page, columns=identifiers, show="headings", height=3, selectmode="browse")
        for identifier, heading, width in columns:
            tree.heading(identifier, text=heading)
            tree.column(identifier, width=width, anchor="w")
        tree.grid(row=1, column=0, sticky="nsew")
        buttons = ttk.Frame(page)
        buttons.grid(row=2, column=0, sticky="w", pady=(4, 0))
        for action, label in (("add", "添加"), ("edit", "编辑"), ("delete", "删除"), ("up", "上移"), ("down", "下移")):
            ttk.Button(buttons, text=label, command=lambda k=key, t=tree, a=action: self._manage_operation_rule(k, t, a)).pack(side="left", padx=2)
        return tree

    def _render_area_configs(self):
        for key, tree in (("pre_mappings", self.pre_mapping_tree), ("post_mappings", self.post_mapping_tree),
                          ("image_outputs", self.image_tree), ("text_outputs", self.text_tree)):
            for item in tree.get_children():
                tree.delete(item)
            columns = tree["columns"]
            for idx, rule in enumerate(self.settings.get(key, [])):
                tree.insert("", "end", iid=str(idx), values=tuple(rule.get(column, "") for column in columns))
        operation_trees = (("TextToColumns", self.text_to_columns_tree, ("sheet", "range", "delimiter", "output_start_col", "column_formats")),
                           ("Deduplicate", self.deduplicate_tree, ("sheet", "range", "has_header", "key_mode", "key_columns")),
                           ("FormatConvert", self.format_convert_tree, ("sheet", "range", "mode")))
        for key, tree, columns in operation_trees:
            operation = self.settings.get(key, {"enabled": False, "rules": []})
            self.operation_enabled[key].set(bool(operation.get("enabled", False)))
            for item in tree.get_children():
                tree.delete(item)
            for idx, rule in enumerate(operation.get("rules", [])):
                if key == "TextToColumns":
                    values = (rule.get("sheet", ""), rule.get("range", ""), repr(rule.get("delimiter", "")),
                              rule.get("output_start_col", ""), ",".join(rule.get("column_formats", ["text"])))
                elif key == "Deduplicate":
                    values = (rule.get("sheet", ""), rule.get("range", ""), "是" if rule.get("has_header") else "否",
                              "整行" if rule.get("key_mode", "row") == "row" else "指定列",
                              ",".join(rule.get("key_columns", [])))
                else:
                    values = tuple(rule.get(column, "") for column in columns)
                tree.insert("", "end", iid=str(idx), values=values)

    def _edit_area_config(self, key: str, tree, action: str):
        entries = self.settings.setdefault(key, [])
        selected = tree.selection()
        if action == "edit" and not selected:
            return
        if action == "delete":
            if selected:
                del entries[int(selected[0])]
                self._render_area_configs()
            return
        index = int(selected[0]) if action == "edit" else None
        existing = entries[index] if index is not None else {}
        is_mapping = key in ("pre_mappings", "post_mappings")
        title = {"pre_mappings": "前映射规则", "post_mappings": "后映射规则",
                 "image_outputs": "区域 XLSX 输出", "text_outputs": "文本输出区域"}[key]
        fields = (("source_sheet", "源工作表"), ("source_range", "源区域"),
                  ("target_sheet", "目标工作表"), ("target_range", "目标区域")) if is_mapping else (("sheet", "模板工作表"), ("range", "单元格区域"))
        dialog = self.tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.grab_set()
        variables = {}
        for row, (field, label) in enumerate(fields):
            variables[field] = self.tk.StringVar(value=str(existing.get(field, "")))
            self.ttk.Label(dialog, text=label).grid(row=row, column=0, sticky="w", padx=8, pady=5)
            self.ttk.Entry(dialog, textvariable=variables[field], width=34).grid(row=row, column=1, sticky="ew", padx=5, pady=5)
        result = {"value": None}

        def accept():
            value = {field: variables[field].get().strip() for field, _label in fields}
            try:
                if is_mapping:
                    _validate_sheet_name(value["source_sheet"], "映射源")
                    _validate_sheet_name(value["target_sheet"], "映射目标")
                    source_bounds = _parse_area(value["source_range"], "映射源")
                    target_bounds = _parse_area(value["target_range"], "映射目标")
                    source_shape = (source_bounds[2] - source_bounds[0], source_bounds[3] - source_bounds[1])
                    target_shape = (target_bounds[2] - target_bounds[0], target_bounds[3] - target_bounds[1])
                    if source_shape != target_shape:
                        raise ProcessingError("源区域与目标区域的行数和列数必须一致。")
                    value["source_range"] = _area_address(source_bounds)
                    value["target_range"] = _area_address(target_bounds)
                else:
                    _validate_sheet_name(value["sheet"], "输出")
                    bounds, selection_kind = _parse_selection_details(value["range"], "输出")
                    if key == "image_outputs" and selection_kind not in ("sheet", "row", "column"):
                        _check_region_xlsx_area_size(bounds, "区域 XLSX 输出")
                    elif key == "text_outputs" and selection_kind not in ("sheet", "row", "column"):
                        _check_text_area_size(bounds, "文本输出")
            except ProcessingError as exc:
                self.messagebox.showerror(APP_NAME, str(exc), parent=dialog)
                return
            result["value"] = value
            dialog.destroy()

        buttons = self.ttk.Frame(dialog)
        buttons.grid(row=len(fields), column=0, columnspan=2, pady=8)
        self.ttk.Button(buttons, text="确定", command=accept).pack(side="left", padx=5)
        self.ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="left", padx=5)
        dialog.wait_window()
        if result["value"] is not None:
            if index is None:
                entries.append(result["value"])
                selected_index = len(entries) - 1
            else:
                entries[index] = result["value"]
                selected_index = index
            self._render_area_configs()
            if str(selected_index) in tree.get_children():
                tree.selection_set(str(selected_index))

    def _operation_rule_dialog(self, key: str, existing: dict[str, Any] | None = None):
        title = {"TextToColumns": "TextToColumns 规则", "Deduplicate": "Deduplicate 规则", "FormatConvert": "FormatConvert 规则"}[key]
        if key == "TextToColumns":
            fields = (("sheet", "工作表名", "entry"), ("range", "单列源范围", "entry"),
                      ("delimiter", r"分隔符（可填 \\t / \\n）", "entry"), ("output_start_col", "输出起始列", "entry"),
                      ("column_formats", "结果列格式（逗号分隔）", "entry"))
            initial = {"sheet": "", "range": "", "delimiter": ",", "output_start_col": "A", "column_formats": ["text"]}
        elif key == "Deduplicate":
            fields = (("sheet", "工作表名", "entry"), ("range", "整行/选列数据范围", "entry"),
                      ("has_header", "表头", "header"), ("key_mode", "判重方式", "keymode"),
                      ("key_columns", "关键列（如 A,C）", "entry"))
            initial = {"sheet": "", "range": "", "has_header": False, "key_mode": "row", "key_columns": []}
        else:
            fields = (("sheet", "工作表名", "entry"), ("range", "转换范围", "entry"), ("mode", "目标类型", "mode"))
            initial = {"sheet": "", "range": "", "mode": "text"}
        data = {**initial, **(existing or {})}
        dialog = self.tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.grab_set()
        variables = {}
        for row, (field, label, kind) in enumerate(fields):
            self.ttk.Label(dialog, text=label).grid(row=row, column=0, sticky="w", padx=8, pady=5)
            if kind == "header":
                variable = self.tk.StringVar(value="有表头" if data.get(field, False) else "无表头")
                widget = self.ttk.Combobox(dialog, textvariable=variable, values=("有表头", "无表头"), state="readonly", width=20)
            elif kind == "keymode":
                variable = self.tk.StringVar(value="整行" if data.get(field, "row") == "row" else "指定关键列")
                widget = self.ttk.Combobox(dialog, textvariable=variable, values=("整行", "指定关键列"), state="readonly", width=20)
            elif kind == "mode":
                variable = self.tk.StringVar(value=str(data.get(field, "text")))
                widget = self.ttk.Combobox(dialog, textvariable=variable, values=("text", "number", "date"), state="readonly", width=20)
            else:
                value = data.get(field, "")
                if field == "column_formats":
                    value = ",".join(value)
                elif field == "key_columns":
                    value = ",".join(value)
                variable = self.tk.StringVar(value=str(value))
                widget = self.ttk.Entry(dialog, textvariable=variable, width=32)
            widget.grid(row=row, column=1, sticky="ew", padx=5, pady=5)
            variables[field] = variable
        result = {"value": None}

        def accept():
            rule = {}
            for field, _label, kind in fields:
                value = variables[field].get().strip()
                if field == "has_header":
                    rule[field] = value == "有表头"
                elif field == "key_mode":
                    rule[field] = "row" if value == "整行" else "columns"
                elif field in ("column_formats", "key_columns"):
                    rule[field] = [part.strip().lower() if field == "column_formats" else part.strip().upper()
                                   for part in value.split(",") if part.strip()]
                elif field == "delimiter":
                    rule[field] = value.replace(r"\t", "\t").replace(r"\n", "\n")
                else:
                    rule[field] = value.lower() if field == "mode" else value
            if key == "Deduplicate" and rule["key_mode"] == "row":
                rule["key_columns"] = []
            try:
                trial = normalize_settings(self.settings)
                trial[key]["rules"] = [*trial[key].get("rules", []), rule]
                validate_static_settings(trial)
            except ProcessingError as exc:
                self.messagebox.showerror(APP_NAME, str(exc), parent=dialog)
                return
            result["value"] = rule
            dialog.destroy()

        buttons = self.ttk.Frame(dialog)
        buttons.grid(row=len(fields), column=0, columnspan=2, pady=8)
        self.ttk.Button(buttons, text="确定", command=accept).pack(side="left", padx=5)
        self.ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="left", padx=5)
        dialog.wait_window()
        return result["value"]

    def _manage_operation_rule(self, key: str, tree, action: str):
        rules = self.settings.setdefault(key, {"enabled": False, "rules": []}).setdefault("rules", [])
        selected = tree.selection()
        if action in ("edit", "delete", "up", "down") and not selected:
            return
        if action == "delete":
            del rules[int(selected[0])]
        elif action in ("up", "down"):
            index = int(selected[0])
            target = index + (-1 if action == "up" else 1)
            if 0 <= target < len(rules):
                rules[index], rules[target] = rules[target], rules[index]
                self._render_area_configs()
                tree.selection_set(str(target))
                return
        else:
            index = int(selected[0]) if action == "edit" else None
            value = self._operation_rule_dialog(key, rules[index] if index is not None else None)
            if value is None:
                return
            if index is None:
                rules.append(value)
                index = len(rules) - 1
            else:
                rules[index] = value
        self._render_area_configs()
        if rules and action != "delete":
            tree.selection_set(str(index))

    def _render_jobs(self, select_index: int | None = None):
        for item in self.job_tree.get_children():
            self.job_tree.delete(item)
        for idx, job in enumerate(self.settings.get("jobs", [])):
            map_summary = ", ".join(f"{m.get('result_col')}→{m.get('target_col')}({m.get('write_as', 'value')})" for m in job.get("mappings", []))
            self.job_tree.insert("", "end", iid=str(idx), values=(job.get("template_sheet", ""), job.get("start_row", ""), job.get("formula_range", ""), len(job.get("sources", [])), map_summary))
        if self.settings.get("jobs"):
            index = 0 if select_index is None else max(0, min(select_index, len(self.settings["jobs"]) - 1))
            self.job_tree.selection_set(str(index))
            self.job_tree.focus(str(index))
            self._show_job(index)

    def _current_index(self) -> int | None:
        selected = self.job_tree.selection()
        if not selected:
            return None
        return int(selected[0])

    def _focus_source_location(self, job_index: int, source_index: int):
        job_id = str(job_index)
        self.job_tree.selection_set(job_id)
        self.job_tree.focus(job_id)
        self.job_tree.see(job_id)
        self._show_job(job_index)
        source_id = str(source_index)
        if source_id in self.source_tree.get_children():
            self.source_tree.selection_set(source_id)
            self.source_tree.focus(source_id)
            self.source_tree.see(source_id)

    def _commit_job_editor(self, show_errors=True):
        index = self.active_job_index
        if index is None:
            return True
        if index >= len(self.settings.get("jobs", [])):
            return True
        job = self.settings["jobs"][index]
        job["template_sheet"] = self.job_vars["template_sheet"].get().strip()
        job["formula_range"] = self.job_vars["formula_range"].get().strip()
        try:
            job["header_rows"] = int(self.job_vars["header_rows"].get())
            job["start_row"] = int(self.job_vars["start_row"].get())
        except ValueError:
            if show_errors:
                self.messagebox.showerror(APP_NAME, "保留标题行数和粘贴起始行必须为整数。")
            return False
        return True

    def _show_job(self, index: int):
        if not (0 <= index < len(self.settings.get("jobs", []))):
            return
        self.active_job_index = index
        job = self.settings["jobs"][index]
        for key in ("template_sheet", "header_rows", "start_row", "formula_range"):
            self.job_vars[key].set(str(job.get(key, "")))
        self._update_formula_template_row_display()
        for tree in (self.source_tree, self.mapping_tree):
            for item in tree.get_children():
                tree.delete(item)
        for src_idx, source in enumerate(job.get("sources", [])):
            sheet_display = str(source.get("source_sheet") or "").strip() or "（首个工作表）"
            self.source_tree.insert("", "end", iid=str(src_idx), values=(source.get("path", ""), sheet_display, source.get("copy_start_row", ""), source.get("copy_end_row") or ""))
        for map_idx, mapping in enumerate(job.get("mappings", [])):
            self.mapping_tree.insert("", "end", iid=str(map_idx), values=(mapping.get("result_col", ""), mapping.get("target_col", ""), mapping.get("write_as", "value")))

    def _on_job_select(self, _event=None):
        if not self._commit_job_editor(show_errors=False):
            if self.active_job_index is not None:
                self.job_tree.selection_set(str(self.active_job_index))
            return
        index = self._current_index()
        if index is not None:
            previous_index = self.active_job_index
            if previous_index is not None:
                self._refresh_job_row(previous_index)
            self._show_job(index)
            self._refresh_job_row(index)

    def _update_formula_template_row_display(self, *_args):
        if not self.job_vars or "formula_template_row" not in self.job_vars:
            return
        try:
            start_row = int(self.job_vars["start_row"].get())
            value = str(start_row - 1) if start_row >= 2 else "需从第 2 行开始"
        except (TypeError, ValueError):
            value = "请输入有效行号"
        self.job_vars["formula_template_row"].set(value)

    def _show_configuration_migration_notices(self, notices):
        self.messagebox.showinfo(APP_NAME, "\n\n".join(notices))

    def _refresh_job_row(self, index):
        if not (0 <= index < len(self.settings["jobs"])):
            return
        job = self.settings["jobs"][index]
        maps = ", ".join(f"{m.get('result_col')}→{m.get('target_col')}({m.get('write_as', 'value')})" for m in job.get("mappings", []))
        self.job_tree.item(str(index), values=(job.get("template_sheet", ""), job.get("start_row", ""), job.get("formula_range", ""), len(job.get("sources", [])), maps))

    def _add_job(self):
        if not self._commit_job_editor():
            return
        self.settings["jobs"].append({"template_sheet": "新目标表", "header_rows": 0, "start_row": 2, "formula_range": "A:A",
                                      "sources": [{"path": "", "source_sheet": "", "copy_start_row": 2, "copy_end_row": None}],
                                      "mappings": [{"result_col": "A", "target_col": "B", "write_as": "value"}]})
        self._render_jobs(len(self.settings["jobs"]) - 1)

    def _delete_job(self):
        index = self._current_index()
        if index is None:
            return
        del self.settings["jobs"][index]
        self._render_jobs(min(index, len(self.settings["jobs"]) - 1) if self.settings["jobs"] else None)

    def _selected_job(self):
        if not self._commit_job_editor():
            return None
        index = self._current_index()
        return self.settings["jobs"][index] if index is not None else None

    def _source_dialog(self, existing=None):
        tk, ttk = self.tk, self.ttk
        dialog = tk.Toplevel(self.root)
        dialog.title("源文件设置")
        dialog.transient(self.root)
        dialog.grab_set()
        vals = {"path": tk.StringVar(value=(existing or {}).get("path", "")),
                "source_sheet": tk.StringVar(value=(existing or {}).get("source_sheet", "")),
                "copy_start_row": tk.StringVar(value=str((existing or {}).get("copy_start_row", 2))),
                "copy_end_row": tk.StringVar(value=str((existing or {}).get("copy_end_row") or ""))}
        ttk.Label(dialog, text="源工作簿（.xlsx）").grid(row=0, column=0, sticky="w", padx=8, pady=5)
        path_entry = ttk.Entry(dialog, textvariable=vals["path"], width=58)
        path_entry.grid(row=0, column=1, padx=5, pady=5)
        def browse():
            path = self.filedialog.askopenfilename(filetypes=[("Excel 工作簿（无宏）", "*.xlsx")])
            if path:
                vals["path"].set(path)
        ttk.Button(dialog, text="浏览…", command=browse).grid(row=0, column=2, padx=5)
        for row, key, label in ((1, "source_sheet", "源工作表名（留空=首个工作表）"), (2, "copy_start_row", "复制起始行"), (3, "copy_end_row", "复制结束行（留空=遇整行空白）")):
            ttk.Label(dialog, text=label).grid(row=row, column=0, sticky="w", padx=8, pady=5)
            ttk.Entry(dialog, textvariable=vals[key], width=26).grid(row=row, column=1, sticky="w", padx=5, pady=5)
        result = {"value": None}
        def accept():
            try:
                src = {key: vals[key].get().strip() for key in vals}
                if not src["path"]:
                    self.messagebox.showerror(APP_NAME, "请选择或填写源工作簿 .xlsx 路径。", parent=dialog)
                    path_entry.focus_set()
                    return
                src["copy_start_row"] = int(src["copy_start_row"])
                src["copy_end_row"] = int(src["copy_end_row"]) if src["copy_end_row"] else None
                result["value"] = src
                dialog.destroy()
            except ValueError:
                self.messagebox.showerror(APP_NAME, "复制起止行必须填写整数；结束行可以留空。", parent=dialog)
        btns = ttk.Frame(dialog)
        btns.grid(row=4, column=0, columnspan=3, pady=8)
        ttk.Button(btns, text="确定", command=accept).pack(side="left", padx=5)
        ttk.Button(btns, text="取消", command=dialog.destroy).pack(side="left", padx=5)
        dialog.wait_window()
        return result["value"]

    def _add_source(self):
        job = self._selected_job()
        if job is None:
            return
        source = self._source_dialog()
        if source:
            job["sources"].append(source)
            self._show_job(self._current_index())
            self._refresh_job_row(self._current_index())

    def _edit_source(self):
        job = self._selected_job()
        selected = self.source_tree.selection()
        if job is None or not selected:
            return
        idx = int(selected[0])
        source = self._source_dialog(job["sources"][idx])
        if source:
            job["sources"][idx] = source
            self._show_job(self._current_index())
            self.source_tree.selection_set(str(idx))
            self._refresh_job_row(self._current_index())

    def _delete_source(self):
        job = self._selected_job()
        selected = self.source_tree.selection()
        if job is not None and selected:
            del job["sources"][int(selected[0])]
            self._show_job(self._current_index())
            self._refresh_job_row(self._current_index())

    def _move_source(self, direction):
        job = self._selected_job()
        selected = self.source_tree.selection()
        if job is None or not selected:
            return
        idx, dest = int(selected[0]), int(selected[0]) + direction
        if 0 <= dest < len(job["sources"]):
            job["sources"][idx], job["sources"][dest] = job["sources"][dest], job["sources"][idx]
            self._show_job(self._current_index())
            self.source_tree.selection_set(str(dest))

    def _mapping_dialog(self, existing=None):
        tk, ttk = self.tk, self.ttk
        dialog = tk.Toplevel(self.root)
        dialog.title("结果列映射")
        dialog.transient(self.root)
        dialog.grab_set()
        vals = {"result_col": tk.StringVar(value=(existing or {}).get("result_col", "")),
                "target_col": tk.StringVar(value=(existing or {}).get("target_col", "")),
                "write_as": tk.StringVar(value=(existing or {}).get("write_as", "value"))}
        for row, key, label in ((0, "result_col", "公式结果列"), (1, "target_col", "写入目标列")):
            ttk.Label(dialog, text=label).grid(row=row, column=0, sticky="w", padx=8, pady=5)
            ttk.Entry(dialog, textvariable=vals[key], width=15).grid(row=row, column=1, sticky="w", padx=5, pady=5)
        ttk.Label(dialog, text="目标列写入为").grid(row=2, column=0, sticky="w", padx=8, pady=5)
        ttk.Combobox(dialog, textvariable=vals["write_as"], values=("value", "text"), state="readonly", width=12).grid(row=2, column=1, sticky="w", padx=5, pady=5)
        result = {"value": None}
        def accept():
            mapping = {key: vals[key].get().strip() for key in vals}
            try:
                column_number(mapping["result_col"])
                column_number(mapping["target_col"])
            except ProcessingError as exc:
                self.messagebox.showerror(APP_NAME, str(exc), parent=dialog)
                return
            result["value"] = mapping
            dialog.destroy()
        buttons = ttk.Frame(dialog)
        buttons.grid(row=3, column=0, columnspan=2, pady=8)
        ttk.Button(buttons, text="确定", command=accept).pack(side="left", padx=5)
        ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="left", padx=5)
        dialog.wait_window()
        return result["value"]

    def _add_mapping(self):
        job = self._selected_job()
        if job is None:
            return
        mapping = self._mapping_dialog()
        if mapping:
            job["mappings"].append(mapping)
            self._show_job(self._current_index())
            self._refresh_job_row(self._current_index())

    def _edit_mapping(self):
        job = self._selected_job()
        selected = self.mapping_tree.selection()
        if job is None or not selected:
            return
        idx = int(selected[0])
        mapping = self._mapping_dialog(job["mappings"][idx])
        if mapping:
            job["mappings"][idx] = mapping
            self._show_job(self._current_index())
            self.mapping_tree.selection_set(str(idx))
            self._refresh_job_row(self._current_index())

    def _delete_mapping(self):
        job = self._selected_job()
        selected = self.mapping_tree.selection()
        if job is not None and selected:
            del job["mappings"][int(selected[0])]
            self._show_job(self._current_index())
            self._refresh_job_row(self._current_index())

    def _browse_template(self):
        path = self.filedialog.askopenfilename(filetypes=[("Excel 工作簿（无宏）", "*.xlsx")])
        if path:
            self._set_template_selected(path)
            try:
                prefix_length = int(self.vars["template_prefix_length"].get())
                matches = matching_profiles(path, prefix_length, PROFILE_DIR)
            except Exception as exc:
                self.messagebox.showerror(APP_NAME, str(exc))
                return
            if not matches:
                self.messagebox.showinfo(APP_NAME, f"没有已保存配置与模板名“{Path(path).stem}”前 {prefix_length} 个字符匹配。\n未自动载入其他配置；当前界面配置保持不变。")
            elif len(matches) == 1:
                name = matches[0]
                if self.messagebox.askyesno(APP_NAME, f"找到唯一匹配配置“{name}”。是否载入？"):
                    try:
                        self._load_profile_named(name, selected_template_path=path)
                    except Exception as exc:
                        self.messagebox.showerror(APP_NAME, f"载入匹配配置失败：\n{exc}")
            else:
                name = self._choose_profile_dialog(matches)
                if name:
                    try:
                        self._load_profile_named(name, selected_template_path=path)
                    except Exception as exc:
                        self.messagebox.showerror(APP_NAME, f"载入匹配配置失败：\n{exc}")

    def _browse_output_dir(self):
        path = self.filedialog.askdirectory()
        if path:
            self.vars["output_dir"].set(path)

    def _collect(self):
        if not self._commit_job_editor(show_errors=False):
            raise ProcessingError("当前目标表的保留标题行数和粘贴起始行必须是整数。")
        settings = {key: var.get().strip() for key, var in self.vars.items()}
        settings["output_filename_custom"] = bool(self.output_filename_custom)
        settings["export_region_pngs"] = bool(self.export_region_pngs_var.get())
        settings["config_version"] = CONFIG_VERSION
        settings["config_name"] = self.active_profile_name
        settings["jobs"] = deepcopy(self.settings["jobs"])
        settings["workflow"] = deepcopy(self.settings.get("workflow", DEFAULT_WORKFLOW))
        for key in ("pre_mappings", "post_mappings", "image_outputs", "text_outputs"):
            settings[key] = deepcopy(self.settings.get(key, []))
        for key in ("TextToColumns", "Deduplicate", "FormatConvert"):
            operation = deepcopy(self.settings.get(key, {"enabled": False, "rules": []}))
            operation["enabled"] = bool(self.operation_enabled[key].get())
            settings[key] = operation
        return normalize_settings(settings)

    def _save_default(self):
        try:
            settings = self._collect()
            save_settings(settings)
            self.settings = settings
            self.messagebox.showinfo(APP_NAME, f"默认配置已保存：\n{CONFIG_PATH}")
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))

    def _new_profile(self):
        from tkinter import simpledialog
        name = simpledialog.askstring(APP_NAME, "新配置名称：", parent=self.root)
        if name is not None and name.strip():
            self._save_profile_named(name.strip())

    def _save_current_profile(self):
        name = self.active_profile_name or self.profile_choice.get().strip()
        if not name:
            self._new_profile()
            return
        self._save_profile_named(name)

    def _save_profile_named(self, name):
        try:
            settings = self._collect()
            settings["config_name"] = name
            validate_static_settings(settings)
            if name in list_profiles(PROFILE_DIR) and not self.messagebox.askyesno(APP_NAME, f"配置“{name}”已存在，是否覆盖保存？"):
                return
            path = save_profile(settings, name, PROFILE_DIR)
            self.settings = normalize_settings(settings)
            self.active_profile_name = name
            self.profile_choice.set(name)
            self._refresh_profile_choices()
            self.messagebox.showinfo(APP_NAME, f"已保存命名配置“{name}”：\n{path}")
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))

    def _apply_settings_to_gui(self, settings, active_name=None):
        migration_notices = settings.get("_migration_notices") or configuration_migration_notes(settings)
        normalized = normalize_settings(settings)
        self.settings = normalized
        self.active_profile_name = str(active_name if active_name is not None else normalized.get("config_name", ""))
        self.output_filename_custom = bool(normalized.get("output_filename_custom", False))
        self._suppress_output_filename_trace = True
        try:
            for key, var in self.vars.items():
                if key in normalized:
                    var.set(str(normalized.get(key, "")))
        finally:
            self._suppress_output_filename_trace = False
        for key, var in self.operation_enabled.items():
            var.set(bool(normalized.get(key, {}).get("enabled", False)))
        self.export_region_pngs_var.set(bool(normalized.get("export_region_pngs", False)))
        self.profile_choice.set(self.active_profile_name)
        self._render_jobs()
        self._render_area_configs()
        self._render_workflow()
        self._refresh_profile_choices()
        if migration_notices:
            self._show_configuration_migration_notices(migration_notices)

    def _load_profile_named(self, name, selected_template_path=None):
        settings = load_profile(name, PROFILE_DIR)
        if selected_template_path:
            settings["template_path"] = selected_template_path
        settings = normalize_settings(settings)
        validate_static_settings(settings)
        self._apply_settings_to_gui(settings, name)

    def _load_selected_profile(self):
        name = self.profile_choice.get().strip()
        if not name:
            self.messagebox.showinfo(APP_NAME, "当前没有可载入的命名配置。")
            return
        try:
            self._load_profile_named(name)
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, f"载入配置失败：\n{exc}")

    def _choose_profile_dialog(self, names):
        dialog = self.tk.Toplevel(self.root)
        dialog.title("选择匹配配置")
        dialog.transient(self.root)
        dialog.grab_set()
        self.ttk.Label(dialog, text="模板前缀匹配到多份配置，请选择：").pack(anchor="w", padx=10, pady=(10, 4))
        choice = self.tk.StringVar(value=names[0])
        combo = self.ttk.Combobox(dialog, textvariable=choice, values=names, state="readonly", width=32)
        combo.pack(padx=10, pady=4)
        result = {"name": None}
        buttons = self.ttk.Frame(dialog)
        buttons.pack(pady=8)
        def accept():
            result["name"] = choice.get()
            dialog.destroy()
        self.ttk.Button(buttons, text="载入所选", command=accept).pack(side="left", padx=4)
        self.ttk.Button(buttons, text="取消", command=dialog.destroy).pack(side="left", padx=4)
        dialog.wait_window()
        return result["name"]

    def _export_config(self):
        try:
            settings = self._collect()
            validate_static_settings(settings)
            path = self.filedialog.asksaveasfilename(defaultextension=".json", filetypes=[("JSON 配置", "*.json")], initialfile="wps_sheet_config.json")
            if path:
                save_settings(settings, path)
                self.settings = settings
                self._render_jobs(self._current_index())
                self.messagebox.showinfo(APP_NAME, f"配置已导出：\n{path}")
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))

    def _import_config(self):
        path = self.filedialog.askopenfilename(filetypes=[("JSON 配置", "*.json")])
        if not path:
            return
        try:
            settings = load_settings(path)
            validate_static_settings(settings)
            self._apply_settings_to_gui(settings)
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, f"加载配置失败：\n{exc}")

    def _reset_config(self):
        if not self.messagebox.askyesno(APP_NAME, "恢复程序内置示例配置？不会删除已保存或导出的配置文件。"):
            return
        self._apply_settings_to_gui(deepcopy(DEFAULT_SETTINGS), active_name="")

    def _log(self, message: str):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", message + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _copy_preview_to_clipboard(self, text: str):
        try:
            self.root.clipboard_clear()
            self.root.clipboard_append(text)
            self.root.update_idletasks()
            self._log("已将文本区域复制到系统剪贴板。")
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, f"写入 Windows 剪贴板失败：{exc}")

    def _render_text_previews(self, previews: list[dict[str, Any]]):
        self.text_previews = deepcopy(previews)
        for tab in self.preview_notebook.tabs():
            self.preview_notebook.forget(tab)
        if not previews:
            self.preview_notebook.pack_forget()
            self.preview_empty_label.pack(anchor="w", padx=4, pady=4)
            return
        self.preview_empty_label.pack_forget()
        self.preview_notebook.pack(fill="both", expand=True)
        for idx, preview in enumerate(previews, 1):
            page = self.ttk.Frame(self.preview_notebook, padding=5)
            page.columnconfigure(0, weight=1)
            page.rowconfigure(2, weight=1)
            resolved = preview.get("resolved_range", preview["range"])
            range_text = f"{preview['range']}（实际读取 {resolved}）" if resolved != preview["range"] else preview["range"]
            self.ttk.Label(page, text=f"{preview['sheet']}!{range_text}").grid(row=0, column=0, sticky="w", padx=4, pady=3)
            chunks = preview.get("chunks") or [{"start_row": 1, "end_row": 1, "text": preview["text"]}]
            chunk_box = None
            if len(chunks) > 1:
                labels = [f"第 {n} 块（行 {chunk['start_row']}-{chunk['end_row']}）" for n, chunk in enumerate(chunks, 1)]
                chunk_var = self.tk.StringVar(value=labels[0])
                chunk_box = self.ttk.Combobox(page, textvariable=chunk_var, values=labels, state="readonly", width=32)
                chunk_box.grid(row=1, column=0, sticky="w", padx=4, pady=3)
            text_widget = self.tk.Text(page, height=5, wrap="none")
            text_widget.insert("1.0", chunks[0]["text"])
            text_widget.configure(state="disabled")
            text_widget.grid(row=2, column=0, sticky="nsew", padx=4, pady=3)
            if chunk_box is not None:
                def show_chunk(_event=None, box=chunk_box, items=chunks, widget=text_widget):
                    selected_index = max(0, box.current())
                    widget.configure(state="normal")
                    widget.delete("1.0", "end")
                    widget.insert("1.0", items[selected_index]["text"])
                    widget.configure(state="disabled")
                chunk_box.bind("<<ComboboxSelected>>", show_chunk)
            def copy_current_chunk(items=chunks, box=chunk_box):
                selected_index = max(0, box.current()) if box is not None else 0
                self._copy_preview_to_clipboard(items[selected_index]["text"])
            copy_label = f"复制第 {idx} 个区域的当前分块" if len(chunks) > 1 else f"复制第 {idx} 个区域"
            self.ttk.Button(page, text=copy_label, command=copy_current_chunk).grid(row=3, column=0, sticky="w", padx=4, pady=4)
            self.preview_notebook.add(page, text=f"{idx}. {preview['sheet']}!{preview['range']}")

    def _run(self):
        try:
            settings = self._collect()
            validate_static_settings(settings)
            missing_source = _first_missing_source_location(settings)
            if missing_source is not None:
                self._focus_source_location(*missing_source)
                raise ProcessingError(_missing_source_path_message(settings, missing_source))
            save_settings(settings)
            self.settings = settings
        except Exception as exc:
            self.messagebox.showerror(APP_NAME, str(exc))
            return
        self.run_button.configure(state="disabled")
        self.progress.start(10)
        self._render_text_previews([])
        self._export_pngs_for_run = bool(settings.get("export_region_pngs", False))
        self._region_output_files = []
        self._log("开始处理……")
        def work():
            try:
                path = process_workbooks(settings, lambda text: self.events.put(("log", text)),
                                         lambda previews: self.events.put(("text_previews", previews)),
                                         lambda paths: self.events.put(("region_outputs", paths)))
                self.events.put(("success", str(path)))
            except Exception as exc:
                self.events.put(("error", f"{exc}\n\n详细诊断：\n{traceback.format_exc()}"))
        threading.Thread(target=work, daemon=True).start()

    def _drain_events(self):
        try:
            while True:
                event, payload = self.events.get_nowait()
                if event == "log":
                    self._log(str(payload))
                elif event == "text_previews":
                    self._render_text_previews(payload)
                elif event == "region_outputs":
                    self._region_output_files = [Path(path) for path in payload]
                elif event in ("success", "error"):
                    self.progress.stop()
                    self.run_button.configure(state="normal")
                    if event == "success":
                        self._log(f"完成：{payload}")
                        output_note = ""
                        if self._export_pngs_for_run:
                            output_note = f"\n已导出区域 PNG：{len(self._region_output_files)} 张；图片成功后会删除对应区域 XLSX。"
                        self.messagebox.showinfo(APP_NAME, f"处理成功，已保存新文件：\n{payload}{output_note}")
                    else:
                        self._log("处理未全部完成；请查看详细诊断中的已保存文件路径及清理信息。")
                        self.messagebox.showerror(APP_NAME, str(payload))
        except queue.Empty:
            pass
        self.root.after(150, self._drain_events)


def self_check() -> int:
    try:
        validate_static_settings(DEFAULT_SETTINGS)
        expected = {"EB:EN": (132, 144), "EB:EU": (132, 151), "EB:FV": (132, 178)}
        for text, answer in expected.items():
            if parse_column_range(text) != answer:
                raise AssertionError(f"列范围换算错误：{text}")
        legacy = {"sources": ["1.xlsx", "2.xlsx", "3.xlsx"], "source_start_row": 2,
                  "jobs": [{"template_sheet": "一", "source_sheet": "源", "header_rows": 2, "start_row": 3, "formula_range": "EB:EN", "result_col": "EJ", "target_col": "DG"},
                           {"template_sheet": "二", "source_sheet": "源", "header_rows": 2, "start_row": 3, "formula_range": "EB:EU", "result_col": "EQ", "target_col": "DG"},
                           {"template_sheet": "三", "source_sheet": "源", "header_rows": 3, "start_row": 4, "formula_range": "EB:FV", "result_col": "FQ", "target_col": "DG"}]}
        if len(normalize_settings(legacy)["jobs"]) != 3:
            raise AssertionError("v1 配置迁移失败")
        print(f"静态检查通过：v{CONFIG_VERSION} 默认配置有效；公式范围/映射、两阶段映射、三个独立区域处理配置和首版配置迁移均通过。")
        print("未连接 WPS COM；未读取真实业务文件；未验证 Windows GUI 或 EXE。")
        return 0
    except Exception as exc:
        print(f"静态检查失败：{exc}", file=sys.stderr)
        return 1


def main() -> int:
    if "--self-check" in sys.argv:
        return self_check()
    if sys.platform != "win32":
        print("此程序的公式重算与最终保存要求 Windows 版 WPS 表格 COM。请在 Windows 11 + WPS 表格环境运行。")
        print("当前执行环境不是 Windows；没有生成或验证 EXE。可运行 python main.py --self-check 和假数据单元测试。")
        return 2
    try:
        import tkinter as tk
        root = tk.Tk()
        App(root)
        root.mainloop()
        return 0
    except Exception as exc:
        print(f"启动失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
